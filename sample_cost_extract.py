import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from gartman_connection import get_connection


START_DATE = date(2024, 1, 1)
END_DATE = date(2026, 8, 26)
OUTPUT_DIR = Path(__file__).parent


def write_rows(path, headers, rows):
    with path.open("w", newline="", encoding="utf-8-sig") as output:
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)

def write_workbook(path, order_headers, order_rows, carrier_headers, carrier_rows):
    workbook = Workbook()
    order_sheet = workbook.active
    order_sheet.title = "Sample Order Candidates"
    carrier_sheet = workbook.create_sheet("Carrier AP")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet, headers, rows in (
        (order_sheet, order_headers, order_rows),
        (carrier_sheet, carrier_headers, carrier_rows),
    ):
        sheet.append(headers)
        for row in rows:
            sheet.append(list(row))
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
            sheet.column_dimensions[column[0].column_letter].width = width

    workbook.save(path)


conn = get_connection()
cur = conn.cursor()

cur.execute(
    """
    select H.SHIDAT, H.SHORD#, H.SHINV#, trim(H.SHCUST), trim(L.SLITEM),
           trim(coalesce(L.SLDESC, '')), trim(coalesce(L.SLCDESC, '')),
           trim(coalesce(M.IMDESC, '')), M.IMDIV, L.SLQORD, L.SLBLUO,
           L.SLCOST, 'SHIPPED'
    from GSFL2K.SHLINE L
    join GSFL2K.SHHEAD H on H.SHCO=L.SLCO and H.SHLOC=L.SLLOC
      and H.SHORD#=L.SLORD# and H.SHINV#=L.SLINV#
    left join GSFL2K.ITEMMAST M on trim(M.IMITEM)=trim(L.SLITEM)
    where H.SHIDAT between ? and ?
      and (M.IMDIV=5 or upper(coalesce(L.SLDESC,'')) like '%SAMPLE%'
        or upper(coalesce(L.SLCDESC,'')) like '%SAMPLE%'
        or upper(coalesce(M.IMDESC,'')) like '%SAMPLE%'
        or trim(L.SLITEM)='NONSTOCK')
    order by H.SHIDAT desc, H.SHORD#, L.SLITEM
    """,
    START_DATE,
    END_DATE,
)
shipped = cur.fetchall()

cur.execute(
    """
    select H.OHDATE, H.OHORD#, '', trim(H.OHCUST), trim(L.OLITEM),
           trim(coalesce(L.OLDESC, '')), trim(coalesce(L.OLCDESC, '')),
           trim(coalesce(M.IMDESC, '')), M.IMDIV, L.OLQORD, L.OLBLUO,
           L.OLCOST, 'OPEN'
    from GSFL2K.OOLINE L
    join GSFL2K.OOHEAD H on H.OHCO=L.OLCO and H.OHLOC=L.OLLOC
      and H.OHORD#=L.OLORD#
    left join GSFL2K.ITEMMAST M on trim(M.IMITEM)=trim(L.OLITEM)
    where H.OHDATE between ? and ?
      and H.OHOTYP not like '%RA%'
      and trim(H.OHCUST) not like '%TRANSFER%'
      and (M.IMDIV=5 or upper(coalesce(L.OLDESC,'')) like '%SAMPLE%'
        or upper(coalesce(L.OLCDESC,'')) like '%SAMPLE%'
        or upper(coalesce(M.IMDESC,'')) like '%SAMPLE%'
        or trim(L.OLITEM)='NONSTOCK')
    order by H.OHDATE desc, H.OHORD#, L.OLITEM
    """,
    START_DATE,
    END_DATE,
)
open_orders = cur.fetchall()

write_rows(
    OUTPUT_DIR / "sample_order_candidates.csv",
    [
        "Order Date", "Order Number", "Invoice Number", "Account Number",
        "Item Number", "Line Description", "Line Comment", "Item Description",
        "Division", "Order Qty", "Billing Qty", "Item Cost", "Order Status",
    ],
    shipped + open_orders,
)

cur.execute(
    """
    select D.APDIDT, D.APDVEN, trim(V.VMNAME), trim(D.APDINV), D.APDAMT,
           D.APDGL#, trim(D.APDDES), trim(D.APDARCUST), trim(D.APDARINV#)
    from GSFL2K.APDETAIL D
    join GSFL2K.VENDMAST V on V.VMVEND=D.APDVEN
    where D.APDIDT between ? and ?
      and (upper(V.VMNAME) like '%UPS%' or upper(V.VMNAME) like '%FEDEX%'
        or upper(V.VMNAME) like '%FED EX%')
    order by D.APDIDT, D.APDVEN, D.APDINV
    """,
    START_DATE,
    END_DATE,
)
carrier_rows = cur.fetchall()
carrier_headers = [
  "Invoice Date", "Vendor Number", "Vendor Name", "Invoice Number",
  "Amount", "GL Account", "Description", "AR Customer", "AR Invoice",
]
write_rows(OUTPUT_DIR / "sample_carrier_ap.csv", carrier_headers, carrier_rows)
write_workbook(
  OUTPUT_DIR / "sample_cost_review.xlsx",
  [
    "Order Date", "Order Number", "Invoice Number", "Account Number",
    "Item Number", "Line Description", "Line Comment", "Item Description",
    "Division", "Order Qty", "Billing Qty", "Item Cost", "Order Status",
  ],
  shipped + open_orders,
  carrier_headers,
  carrier_rows,
)

print(f"Sample order candidates: {len(shipped) + len(open_orders)}")
print(f"Carrier AP rows: {len(carrier_rows)}")
print(f"Created: {OUTPUT_DIR / 'sample_order_candidates.csv'}")
print(f"Created: {OUTPUT_DIR / 'sample_carrier_ap.csv'}")
print(f"Created: {OUTPUT_DIR / 'sample_cost_review.xlsx'}")
conn.close()