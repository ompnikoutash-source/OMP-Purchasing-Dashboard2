import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from gartman_connection import get_connection


START_DATE = date(2024, 1, 1)
END_DATE = date.today()
OUTPUT_DIR = Path(__file__).parent


def write_csv(path, headers, rows):
    with path.open("w", newline="", encoding="utf-8-sig") as output:
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)


def write_sheet(sheet, headers, rows):
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
        sheet.column_dimensions[column[0].column_letter].width = width


conn = get_connection()
cur = conn.cursor()

work_headers = [
    "Work Source", "Code", "Work Order", "Release", "Sequence", "Date",
    "Item Number", "Description", "Quantity", "Unit Cost", "Cost Extension",
    "User", "Item Division", "Item Master Description", "Product/Form Code",
]
work_rows = []
for table in ("ITEMWORK", "TRANWORK"):
    cur.execute(
        f"""
        select ITCODE, ITORD#, ITREL#, ITSEQ#, ITDATE, trim(ITITEM),
             trim(ITDESC), ITQTY, ITCOST, ITECST, trim(ITUSER),
             ITPRCD, ITFMCD, M.IMDIV, trim(M.IMDESC)
        from GSFL2K.{table}
         left join GSFL2K.ITEMMAST M on trim(M.IMITEM)=trim(ITITEM)
        order by ITDATE, ITORD#, ITSEQ#
        """,
    )
    for row in cur.fetchall():
        code, order_number, release, sequence, when, item, description, quantity, unit_cost, extension, user, product_code, form_code, division, master_description = row
        work_rows.append(
            [table, code, order_number, release, sequence, when, str(item or "").strip(),
             str(description or "").strip(), quantity, unit_cost, extension,
             str(user or "").strip(), division or "", master_description or "",
             f"{product_code}/{form_code}"],
        )

history_headers = [
    "Transaction Date", "Item Number", "Description", "Quantity", "Unit Cost",
    "Material Cost", "Order Number", "Source", "User", "Conversion Code",
    "Item Division", "Item Master Description", "Inventory GL", "Freight GL",
]
cur.execute(
    """
    select IRDATE, trim(IRITEM), trim(IRDESC), IRQTY, IRCOST, IRMATCOST,
            IRORD#, IRSRC, trim(IRUSER), IRCONV, IRINGL, IRFRGL,
            M.IMDIV, trim(M.IMDESC)
        from GSFL2K.ITEMRCHSOX R
        left join GSFL2K.ITEMMAST M on trim(M.IMITEM)=trim(R.IRITEM)
    where IRDATE between ? and ?
      and (upper(trim(IRITEM))='NONSTOCK'
        or upper(IRDESC) like '%SAMPLE%'
        or upper(IRCOMT) like '%SAMPLE%')
    order by IRDATE, IRITEM
    """,
    START_DATE,
    END_DATE,
)
history_rows = []
for row in cur.fetchall():
    when, item, description, quantity, unit_cost, material_cost, order_number, source, user, conversion, inventory_gl, freight_gl, division, master_description = row
    history_rows.append(
        [when, str(item or "").strip(), str(description or "").strip(), quantity,
         unit_cost, material_cost, order_number, str(source or "").strip(),
         str(user or "").strip(), str(conversion or "").strip(), division or "",
         master_description or "", inventory_gl, freight_gl],
    )

division1_headers = history_headers + ["Link Basis"]
division1_rows = []
nonstock_orders = {
    row[6] for row in history_rows
    if str(row[1]).upper() == "NONSTOCK" and row[6] not in (None, 0, "0", "")
}
if nonstock_orders:
    placeholders = ",".join("?" for _ in nonstock_orders)
    cur.execute(
        f"""
        select R.IRDATE, trim(R.IRITEM), trim(R.IRDESC), R.IRQTY, R.IRCOST,
               R.IRMATCOST, R.IRORD#, R.IRSRC, trim(R.IRUSER), R.IRCONV,
               R.IRINGL, R.IRFRGL, M.IMDIV, trim(M.IMDESC)
        from GSFL2K.ITEMRCHSOX R
        join GSFL2K.ITEMMAST M on trim(M.IMITEM)=trim(R.IRITEM) and M.IMDIV=1
        where R.IRDATE between ? and ? and R.IRORD# in ({placeholders})
        order by R.IRDATE, R.IRORD#, R.IRITEM
        """,
        START_DATE,
        END_DATE,
        *nonstock_orders,
    )
    for row in cur.fetchall():
        when, item, description, quantity, unit_cost, material_cost, order_number, source, user, conversion, inventory_gl, freight_gl, division, master_description = row
        division1_rows.append(
            [when, str(item or "").strip(), str(description or "").strip(), quantity,
             unit_cost, material_cost, order_number, str(source or "").strip(),
             str(user or "").strip(), str(conversion or "").strip(), division or "",
             master_description or "", inventory_gl, freight_gl,
             "Shared nonzero inventory order reference"],
        )

write_csv(OUTPUT_DIR / "sample_chip_work_orders.csv", work_headers, work_rows)
write_csv(OUTPUT_DIR / "sample_chip_inventory_history.csv", history_headers, history_rows)
write_csv(OUTPUT_DIR / "sample_chip_division1_links.csv", division1_headers, division1_rows)

workbook = Workbook()
write_sheet(workbook.active, work_headers, work_rows)
workbook.active.title = "Current Work Orders"
history_sheet = workbook.create_sheet("Historical Sample Movements")
write_sheet(history_sheet, history_headers, history_rows)
division1_sheet = workbook.create_sheet("Division 1 Links")
write_sheet(division1_sheet, division1_headers, division1_rows)
workbook.save(OUTPUT_DIR / "sample_chip_conversion_review.xlsx")

print(f"Current work-order rows: {len(work_rows)}")
print(f"Historical sample/NONSTOCK rows: {len(history_rows)}")
print(f"Division 1 linked rows: {len(division1_rows)}")
print(f"Created: {OUTPUT_DIR / 'sample_chip_conversion_review.xlsx'}")
conn.close()