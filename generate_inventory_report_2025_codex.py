r"""
Generate the 2025 inventory reports modeled after the 2024 Accounting files.

Inputs used as references:
  H:\2025\MISC Reports\For Accounting\INVENTORY REPORT 2024 (01-01-2024 to 12-31-2024) - JB.xlsx
  H:\2025\MISC Reports\For Accounting\Inventory report 01.01.24-12.31.24.pdf

Outputs:
  H:\2025\MISC Reports\For Accounting\INVENTORY REPORT 2025 Codex (01-01-2025 to 12-31-2025) - JB.xlsx
  H:\2025\MISC Reports\For Accounting\Inventory report Codex 01.01.25-12.31.25.pdf
  H:\2025\MISC Reports\For Accounting\Inventory report Codex 01.01.25-12.31.25 IMVALUTOT source.xlsx
  H:\2025\MISC Reports\For Accounting\Inventory_Report_2025_Codex_Diagnostics.xlsx
  H:\2025\MISC Reports\For Accounting\inventory_report_2025_codex_diagnostic_queries.sql

The Jacob workbook basis was validated against the 2024 file:
  GSFL2K.ITEMBA_EOD at 2024-12-31, company 1, location <> 90,
  INV_QTY = IBDBOH, BILL_INV = IBDBOH * IMFACT, BILL_COST = BILL_INV * IBLCST.

The COO-provided 2025 QPRINT confirms the legacy IMVALUTOT PDF basis is the
ITEMBA_EOD ending-value field (IBDEBVL) grouped by item family. For 2025, the
QPRINT itself is used as the authoritative PDF/source value because retained
Gartman rows can change slightly after the year-end spool is produced.
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import subprocess
import sys
import warnings
from copy import copy
from pathlib import Path

REPO_ROOT = Path(r"H:\2025\NewForecastingModel\OMPforecasting5")
sys.path.insert(0, str(REPO_ROOT))

import pandas as pd
import pyodbc
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

from OMP_secrets import GARTMAN_PWD, GARTMAN_UID

warnings.filterwarnings(
    "ignore",
    message="pandas only supports SQLAlchemy connectable",
    category=UserWarning,
)


BASE_DIR = Path(r"H:\2025\MISC Reports\For Accounting")
TEMPLATE_XLSX = BASE_DIR / "INVENTORY REPORT 2024 (01-01-2024 to 12-31-2024) - JB.xlsx"
REFERENCE_PDF = BASE_DIR / "Inventory report 01.01.24-12.31.24.pdf"
COO_QPRINT_2025 = BASE_DIR / "QPRINT_ENDDAY_ENDDAY_980431_75.pdf"

START_DATE = dt.date(2025, 1, 1)
END_DATE = dt.date(2025, 12, 31)
AS_OF_DATE = dt.date(2025, 12, 31)

OUTPUT_XLSX = BASE_DIR / "INVENTORY REPORT 2025 Codex (01-01-2025 to 12-31-2025) - JB.xlsx"
OUTPUT_IMVALUTOT_XLSX = BASE_DIR / "Inventory report Codex 01.01.25-12.31.25 IMVALUTOT source.xlsx"
OUTPUT_IMVALUTOT_PDF = BASE_DIR / "Inventory report Codex 01.01.25-12.31.25.pdf"
OUTPUT_DIAGNOSTICS = BASE_DIR / "Inventory_Report_2025_Codex_Diagnostics.xlsx"
OUTPUT_SQL = BASE_DIR / "inventory_report_2025_codex_diagnostic_queries.sql"

JACOB_2024_TOTAL = 13_767_955.99693362
JACOB_2024_DETAIL_ROWS = 68_158
JACOB_2024_ITEM_ROWS = 15_246
IMVALUTOT_2024_PDF_TOTAL = 15_336_431.72
COO_QPRINT_2025_NE90_TOTAL = 20_048_139.55
COO_QPRINT_2025_ALL_LOC_TOTAL = 20_238_507.18

ACCOUNTING_CURRENCY_FMT = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'


def connect() -> pyodbc.Connection:
    return pyodbc.connect(f"DSN=Gartman;UID={GARTMAN_UID};PWD={GARTMAN_PWD}", autocommit=True, timeout=30)


def date_sql(value: dt.date) -> str:
    return f"DATE('{value:%Y-%m-%d}')"


def normalize_numeric(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    df = df.copy()
    for col in columns:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
    return df


def jacob_item_sort(df: pd.DataFrame, extra_cols: list[str] | None = None) -> pd.DataFrame:
    """Approximate the DB2/Jacob workbook order: alphabetic item codes, then numeric-leading codes."""
    extra_cols = extra_cols or []
    work = df.copy()
    work["_ITEM_DIGIT_SORT"] = work["ITEM"].astype(str).str.match(r"^\d").astype(int)
    work = work.sort_values(["_ITEM_DIGIT_SORT", "ITEM", *extra_cols]).drop(columns=["_ITEM_DIGIT_SORT"])
    return work.reset_index(drop=True)


def money(value: float | int | None) -> str:
    return f"${float(value or 0):,.2f}"


def load_jacob_inventory_by_loc(
    conn: pyodbc.Connection,
    as_of_date: dt.date,
    *,
    apply_2024_zero_row_exception: bool = False,
) -> pd.DataFrame:
    sql = f"""
        SELECT
            IB.IBLOC AS LOCATION,
            TRIM(COALESCE(LC.LCRNAM, '')) AS LOCATION_NAME,
            TRIM(IB.IBITEM) AS ITEM,
            TRIM(IM.IMDESC) AS DESCRIPTION,
            TRIM(FM.FMDESC) AS FAMILY,
            TRIM(IM.IMUM1) AS INV_UNIT,
            SUM(IB.IBDBOH) AS INV_QTY,
            IM.IMFACT AS FACTOR,
            TRIM(IM.IMUM2) AS BILL_UNIT,
            CASE
                WHEN IM.IMFACT = 0 THEN SUM(IB.IBDBOH)
                ELSE SUM(IB.IBDBOH * IM.IMFACT)
            END AS BILL_INV,
            CASE
                WHEN IM.IMFACT = 0 THEN SUM(IB.IBDBOH * IB.IBLCST)
                ELSE SUM(IB.IBDBOH * IM.IMFACT * IB.IBLCST)
            END AS BILL_COST,
            CASE
                WHEN (CASE WHEN IM.IMFACT = 0 THEN SUM(IB.IBDBOH) ELSE SUM(IB.IBDBOH * IM.IMFACT) END) = 0
                THEN 0
                ELSE
                    (CASE
                        WHEN IM.IMFACT = 0 THEN SUM(IB.IBDBOH * IB.IBLCST)
                        ELSE SUM(IB.IBDBOH * IM.IMFACT * IB.IBLCST)
                     END)
                    /
                    (CASE
                        WHEN IM.IMFACT = 0 THEN SUM(IB.IBDBOH)
                        ELSE SUM(IB.IBDBOH * IM.IMFACT)
                     END)
            END AS AVG_COST
        FROM GSFL2K.ITEMBA_EOD IB
        JOIN GSFL2K.ITEMMAST IM
          ON TRIM(IB.IBITEM) = TRIM(IM.IMITEM)
        JOIN GSFL2K.FAMILY FM
          ON TRIM(IM.IMFMCD) = TRIM(FM.FMFMCD)
        LEFT JOIN GSFL2K.LOCATION LC
          ON LC.LCCO = IB.IBCO
         AND LC.LCLOC = IB.IBLOC
        WHERE IB.IBEODDATE = {date_sql(as_of_date)}
          AND IB.IBCO = 1
          AND IB.IBLOC <> 90
        GROUP BY
            IB.IBLOC,
            TRIM(COALESCE(LC.LCRNAM, '')),
            TRIM(IB.IBITEM),
            TRIM(IM.IMDESC),
            TRIM(FM.FMDESC),
            IM.IMFACT,
            TRIM(IM.IMUM1),
            TRIM(IM.IMUM2)
        ORDER BY
            ITEM,
            LOCATION
    """
    df = pd.read_sql(sql, conn)
    df = normalize_numeric(df, ["LOCATION", "INV_QTY", "FACTOR", "BILL_INV", "BILL_COST", "AVG_COST"])
    df["LOCATION"] = df["LOCATION"].astype(int)

    # Current ITEMMAST now joins four historical zero rows for RQ2142RS that
    # were not present in Jacob's 2024 workbook. They have no value impact.
    if apply_2024_zero_row_exception:
        mask = (
            df["ITEM"].eq("RQ2142RS")
            & df["LOCATION"].isin([1, 2, 3, 4])
            & df["INV_QTY"].eq(0)
            & df["BILL_COST"].eq(0)
        )
        df = df.loc[~mask].copy()

    return jacob_item_sort(df, ["LOCATION"])


def build_jacob_summaries(inv_loc: pd.DataFrame) -> dict[str, pd.DataFrame]:
    inv_by_item = (
        inv_loc.groupby(["ITEM", "DESCRIPTION", "FAMILY", "INV_UNIT", "FACTOR", "BILL_UNIT"], dropna=False)
        .agg(INV_QTY=("INV_QTY", "sum"), BILL_INV=("BILL_INV", "sum"), BILL_COST=("BILL_COST", "sum"))
        .reset_index()
    )
    inv_by_item = jacob_item_sort(inv_by_item)
    inv_by_item["AVG_COST"] = inv_by_item.apply(
        lambda row: row["BILL_COST"] / row["BILL_INV"] if row["BILL_INV"] else 0.0,
        axis=1,
    )
    inv_by_item = inv_by_item[
        ["ITEM", "DESCRIPTION", "FAMILY", "INV_UNIT", "INV_QTY", "FACTOR", "BILL_UNIT", "BILL_INV", "BILL_COST", "AVG_COST"]
    ]

    loc_sum = (
        inv_loc.groupby(["LOCATION", "LOCATION_NAME"], dropna=False)
        .agg(INV_QTY=("INV_QTY", "sum"), BILL_INV=("BILL_INV", "sum"), BILL_COST=("BILL_COST", "sum"))
        .reset_index()
        .query("BILL_COST != 0")
        .sort_values("LOCATION")
    )
    loc_sum["AVG_COST"] = loc_sum.apply(
        lambda row: row["BILL_COST"] / row["BILL_INV"] if row["BILL_INV"] else 0.0,
        axis=1,
    )
    return {"by_item": inv_by_item, "by_loc": inv_loc, "loc_sum": loc_sum}


def load_imvalutot_family(
    conn: pyodbc.Connection,
    as_of_date: dt.date,
    location_filter_sql: str = "",
) -> pd.DataFrame:
    sql = f"""
        SELECT
            TRIM(IM.IMFMCD) AS FAMILY_CODE,
            TRIM(FM.FMDESC) AS FAMILY,
            SUM(IB.IBDEBVL) AS TOTAL
        FROM GSFL2K.ITEMBA_EOD IB
        JOIN GSFL2K.ITEMMAST IM
          ON TRIM(IB.IBITEM) = TRIM(IM.IMITEM)
        JOIN GSFL2K.FAMILY FM
          ON TRIM(IM.IMFMCD) = TRIM(FM.FMFMCD)
        WHERE IB.IBEODDATE = {date_sql(as_of_date)}
          AND IB.IBCO = 1
          {location_filter_sql}
        GROUP BY
            TRIM(IM.IMFMCD),
            TRIM(FM.FMDESC)
        HAVING SUM(IB.IBDEBVL) <> 0
        ORDER BY
            TRIM(IM.IMFMCD)
    """
    df = pd.read_sql(sql, conn)
    return normalize_numeric(df, ["TOTAL"])


def copy_cell_format(src, dst) -> None:
    if src.has_style:
        dst._style = copy(src._style)
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)


def clear_below_header(ws) -> None:
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def write_table_from_template(ws, rows: list[list], total_row: list, total_formula_col: int | None) -> None:
    data_style = [copy(ws.cell(2, col)) for col in range(1, ws.max_column + 1)]
    total_style = [copy(ws.cell(ws.max_row, col)) for col in range(1, ws.max_column + 1)]
    clear_below_header(ws)

    for row_idx, row_values in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx <= len(data_style):
                copy_cell_format(data_style[col_idx - 1], cell)

    total_idx = len(rows) + 2
    for col_idx in range(1, ws.max_column + 1):
        value = total_row[col_idx - 1] if col_idx <= len(total_row) else None
        cell = ws.cell(row=total_idx, column=col_idx, value=value)
        if col_idx <= len(total_style):
            copy_cell_format(total_style[col_idx - 1], cell)

    if total_formula_col is not None and rows:
        letter = ws.cell(row=1, column=total_formula_col).column_letter
        ws.cell(row=total_idx, column=total_formula_col, value=f"=SUM({letter}2:{letter}{total_idx - 1})")


def write_jacob_workbook(summaries: dict[str, pd.DataFrame], output_path: Path) -> None:
    wb = load_workbook(TEMPLATE_XLSX)

    by_item_rows = summaries["by_item"][
        ["ITEM", "DESCRIPTION", "FAMILY", "INV_UNIT", "INV_QTY", "FACTOR", "BILL_UNIT", "BILL_INV", "BILL_COST", "AVG_COST"]
    ].values.tolist()
    write_table_from_template(
        wb["BY ITEM"],
        by_item_rows,
        ["Total", None, None, None, None, None, None, None, None, None],
        total_formula_col=9,
    )

    by_loc_rows = summaries["by_loc"][
        ["LOCATION", "ITEM", "DESCRIPTION", "FAMILY", "INV_UNIT", "INV_QTY", "FACTOR", "BILL_UNIT", "BILL_INV", "BILL_COST", "AVG_COST"]
    ].values.tolist()
    write_table_from_template(
        wb["BY ITEM & LOC"],
        by_loc_rows,
        ["Total", None, None, None, None, None, None, None, None, None, None],
        total_formula_col=10,
    )

    loc_rows = summaries["loc_sum"][["LOCATION", "LOCATION_NAME", "BILL_COST"]].values.tolist()
    loc_total = float(summaries["loc_sum"]["BILL_COST"].sum())
    write_table_from_template(
        wb["LOC SUM"],
        loc_rows,
        ["TOTAL", None, loc_total],
        total_formula_col=None,
    )

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def parse_imvalutot_pdf_sections(pdf_path: Path) -> list[dict[str, object]]:
    try:
        import pdfplumber
    except ImportError:
        return []

    sections: list[dict[str, object]] = []
    rows: list[dict[str, object]] = []
    pending: tuple[str, str] | None = None
    header = ""
    amount_re = re.compile(r"^-?\d[\d,]*\.\d{2}$")
    code_re = re.compile(r"^([A-Z0-9]{2})\s+(.+)$")

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            for raw_line in text.splitlines():
                line = raw_line.strip()
                if not line or line.startswith("FAMILY") or line.startswith("-------"):
                    continue
                if line.startswith("IMVALUTOT"):
                    header = line
                    continue
                if line == "*TOTALS*":
                    pending = ("*TOTALS*", "")
                    continue
                if pending and amount_re.match(line):
                    value = float(line.replace(",", ""))
                    if pending[0] == "*TOTALS*":
                        df = pd.DataFrame(rows)
                        calc_total = float(df["TOTAL"].sum()) if not df.empty else 0.0
                        sections.append(
                            {
                                "header": header,
                                "rows": df,
                                "reported_total": value,
                                "calculated_total": calc_total,
                            }
                        )
                        rows = []
                    else:
                        rows.append({"FAMILY_CODE": pending[0], "FAMILY": pending[1], "TOTAL": value})
                    pending = None
                    continue
                match = code_re.match(line)
                if match:
                    pending = (match.group(1), match.group(2).strip())

    return sections


def parse_reference_imvalutot_pdf(pdf_path: Path) -> pd.DataFrame:
    sections = parse_imvalutot_pdf_sections(pdf_path)
    if not sections:
        return pd.DataFrame(columns=["FAMILY_CODE", "FAMILY", "PDF_TOTAL"])

    section = sections[-1]
    df = section["rows"].rename(columns={"TOTAL": "PDF_TOTAL"}).copy()
    df.attrs["pdf_total"] = section["reported_total"]
    return df


def parse_coo_qprint_2025(pdf_path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    sections = parse_imvalutot_pdf_sections(pdf_path)
    if not sections:
        return pd.DataFrame(columns=["FAMILY_CODE", "FAMILY", "TOTAL"]), pd.DataFrame()

    summary_rows = []
    for idx, section in enumerate(sections, start=1):
        summary_rows.append(
            {
                "SECTION": idx,
                "HEADER": section["header"],
                "FAMILY_ROWS": len(section["rows"]),
                "CALCULATED_TOTAL_FROM_FAMILY_LINES": section["calculated_total"],
                "PRINTED_TOTAL_LINE": section["reported_total"],
                "PRINTED_MINUS_CALCULATED": section["reported_total"] - section["calculated_total"],
            }
        )
    summary = pd.DataFrame(summary_rows)

    # The COO QPRINT has two sections. Section 1 is the non-temporary-location
    # total. Section 2's family lines are the all-location total; its printed
    # total line is cumulative because section 1 was not reset before section 2.
    chosen = sections[-1]["rows"].copy()
    chosen.attrs["source"] = str(pdf_path)
    chosen.attrs["calculated_total"] = sections[-1]["calculated_total"]
    chosen.attrs["printed_total"] = sections[-1]["reported_total"]
    chosen.attrs["section"] = len(sections)
    return chosen, summary


def write_imvalutot_source_workbook(family: pd.DataFrame, output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "IMVALUTOT"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 94

    font = Font(name="Courier New", size=10)
    bold_font = Font(name="Courier New", size=10, bold=True)

    rows = family.copy()
    rows["TOTAL"] = pd.to_numeric(rows["TOTAL"], errors="coerce").fillna(0.0)
    total = float(rows["TOTAL"].sum())

    lines: list[tuple[str, bool]] = []
    per_page = 27
    pages = [rows.iloc[i : i + per_page] for i in range(0, len(rows), per_page)] or [rows]
    for page_no, chunk in enumerate(pages, start=1):
        lines.append((f"IMVALUTOT INVENTORY VALUE       {dt.datetime.now():%M:%S} {AS_OF_DATE:%m-%d-%Y} {page_no}", True))
        lines.append(("FAMILY                                              TOTAL", True))
        lines.append(("------- ---------- ---------- ---------- ---------- ---------- ---------- ---------- ---------- ---------- ---------- ----------", False))
        for _, row in chunk.iterrows():
            family_name = str(row["FAMILY"])[:42]
            amount = f"{float(row['TOTAL']):,.2f}"
            lines.append((f"{row['FAMILY_CODE']:<2} {family_name:<42} {amount:>15}", False))
        if page_no == len(pages):
            lines.append(("", False))
            lines.append((f"{'*TOTALS*':<45} {total:>15,.2f}", True))
        if page_no < len(pages):
            lines.append(("__PAGE_BREAK__", False))

    row_idx = 1
    for text, is_bold in lines:
        if text == "__PAGE_BREAK__":
            ws.row_breaks.append(Break(id=row_idx - 1))
            continue
        cell = ws.cell(row=row_idx, column=1, value=text)
        cell.font = bold_font if is_bold else font
        cell.alignment = Alignment(horizontal="left")
        row_idx += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def run_powershell(script: str) -> None:
    subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
        check=True,
        text=True,
    )


def recalc_workbook_with_excel(path: Path) -> None:
    escaped = str(path).replace("'", "''")
    ps = f"""
$ErrorActionPreference = 'Stop'
$excel = New-Object -ComObject Excel.Application
$excel.Visible = $false
$excel.DisplayAlerts = $false
try {{
    $wb = $excel.Workbooks.Open('{escaped}')
    $excel.CalculateFull()
    $wb.Save()
}} finally {{
    if ($wb) {{ $wb.Close($false) | Out-Null }}
    $excel.Quit()
    [System.Runtime.Interopservices.Marshal]::ReleaseComObject($excel) | Out-Null
}}
"""
    run_powershell(ps)


def export_pdf_with_excel(source_xlsx: Path, output_pdf: Path) -> None:
    src = str(source_xlsx).replace("'", "''")
    dst = str(output_pdf).replace("'", "''")
    ps = f"""
$ErrorActionPreference = 'Stop'
$excel = New-Object -ComObject Excel.Application
$excel.Visible = $false
$excel.DisplayAlerts = $false
try {{
    $wb = $excel.Workbooks.Open('{src}')
    $ws = $wb.Worksheets.Item(1)
    $ws.ExportAsFixedFormat(0, '{dst}')
    $wb.Close($false) | Out-Null
}} finally {{
    $excel.Quit()
    [System.Runtime.Interopservices.Marshal]::ReleaseComObject($excel) | Out-Null
}}
"""
    run_powershell(ps)


def build_diagnostics(
    conn: pyodbc.Connection,
    summaries_2025: dict[str, pd.DataFrame],
    imvalutot_2025: pd.DataFrame,
) -> dict[str, pd.DataFrame]:
    inv_2024 = load_jacob_inventory_by_loc(
        conn,
        dt.date(2024, 12, 31),
        apply_2024_zero_row_exception=True,
    )
    sum_2024 = build_jacob_summaries(inv_2024)

    jacob_validation = pd.DataFrame(
        [
            {
                "Check": "2024 Jacob workbook detail row count",
                "Expected": JACOB_2024_DETAIL_ROWS,
                "Recreated": len(inv_2024),
                "Difference": len(inv_2024) - JACOB_2024_DETAIL_ROWS,
            },
            {
                "Check": "2024 Jacob workbook item row count",
                "Expected": JACOB_2024_ITEM_ROWS,
                "Recreated": len(sum_2024["by_item"]),
                "Difference": len(sum_2024["by_item"]) - JACOB_2024_ITEM_ROWS,
            },
            {
                "Check": "2024 Jacob workbook inventory value",
                "Expected": JACOB_2024_TOTAL,
                "Recreated": float(sum_2024["by_item"]["BILL_COST"].sum()),
                "Difference": float(sum_2024["by_item"]["BILL_COST"].sum()) - JACOB_2024_TOTAL,
            },
            {
                "Check": "2024 current DB zero rows omitted by Jacob workbook",
                "Expected": "RQ2142RS at locations 1,2,3,4",
                "Recreated": "Excluded only for 2024 validation; zero value impact",
                "Difference": 0,
            },
            {
                "Check": "2025 Jacob workbook detail row count",
                "Expected": "n/a",
                "Recreated": len(summaries_2025["by_loc"]),
                "Difference": "n/a",
            },
            {
                "Check": "2025 Jacob workbook item row count",
                "Expected": "n/a",
                "Recreated": len(summaries_2025["by_item"]),
                "Difference": "n/a",
            },
            {
                "Check": "2025 Jacob workbook inventory value",
                "Expected": "n/a",
                "Recreated": float(summaries_2025["by_item"]["BILL_COST"].sum()),
                "Difference": "n/a",
            },
        ]
    )

    pdf_2024 = parse_reference_imvalutot_pdf(REFERENCE_PDF)
    imvalutot_2024 = load_imvalutot_family(conn, dt.date(2024, 12, 31))
    if not pdf_2024.empty:
        pdf_compare = pdf_2024.merge(
            imvalutot_2024.rename(columns={"TOTAL": "CURRENT_IBDEBVL_TOTAL"}),
            on=["FAMILY_CODE", "FAMILY"],
            how="outer",
        )
        pdf_compare["PDF_TOTAL"] = pd.to_numeric(pdf_compare["PDF_TOTAL"], errors="coerce").fillna(0.0)
        pdf_compare["CURRENT_IBDEBVL_TOTAL"] = pd.to_numeric(
            pdf_compare["CURRENT_IBDEBVL_TOTAL"], errors="coerce"
        ).fillna(0.0)
        pdf_compare["DIFF_CURRENT_MINUS_PDF"] = pdf_compare["CURRENT_IBDEBVL_TOTAL"] - pdf_compare["PDF_TOTAL"]
    else:
        pdf_compare = pd.DataFrame(
            [
                {
                    "FAMILY_CODE": "TOTAL",
                    "FAMILY": "Reference PDF could not be parsed",
                    "PDF_TOTAL": IMVALUTOT_2024_PDF_TOTAL,
                    "CURRENT_IBDEBVL_TOTAL": float(imvalutot_2024["TOTAL"].sum()),
                    "DIFF_CURRENT_MINUS_PDF": float(imvalutot_2024["TOTAL"].sum()) - IMVALUTOT_2024_PDF_TOTAL,
                }
            ]
        )

    qprint_sections = parse_imvalutot_pdf_sections(COO_QPRINT_2025) if COO_QPRINT_2025.exists() else []
    qprint_summary = pd.DataFrame()
    qprint_compare = pd.DataFrame()
    if qprint_sections:
        qprint_summary = pd.DataFrame(
            [
                {
                    "SECTION": idx,
                    "HEADER": section["header"],
                    "INTERPRETATION": "Non-temp locations (loc <> 90)" if idx == 1 else "All locations; printed total is cumulative",
                    "FAMILY_ROWS": len(section["rows"]),
                    "CALCULATED_TOTAL_FROM_FAMILY_LINES": section["calculated_total"],
                    "PRINTED_TOTAL_LINE": section["reported_total"],
                    "PRINTED_MINUS_CALCULATED": section["reported_total"] - section["calculated_total"],
                }
                for idx, section in enumerate(qprint_sections, start=1)
            ]
        )

        db_ne90 = load_imvalutot_family(conn, AS_OF_DATE, "AND IB.IBLOC <> 90").rename(
            columns={"TOTAL": "CURRENT_DB_IBDEBVL_TOTAL"}
        )
        db_all = load_imvalutot_family(conn, AS_OF_DATE).rename(columns={"TOTAL": "CURRENT_DB_IBDEBVL_TOTAL"})
        compare_parts = []
        for label, section, db in [
            ("QPRINT section 1 vs current DB loc<>90", qprint_sections[0], db_ne90),
            ("QPRINT section 2 family subtotal vs current DB all locs", qprint_sections[-1], db_all),
        ]:
            q = section["rows"].rename(columns={"TOTAL": "QPRINT_TOTAL"}).copy()
            comp = q.merge(db, on=["FAMILY_CODE", "FAMILY"], how="outer")
            comp["COMPARISON"] = label
            comp["QPRINT_TOTAL"] = pd.to_numeric(comp["QPRINT_TOTAL"], errors="coerce").fillna(0.0)
            comp["CURRENT_DB_IBDEBVL_TOTAL"] = pd.to_numeric(
                comp["CURRENT_DB_IBDEBVL_TOTAL"], errors="coerce"
            ).fillna(0.0)
            comp["DIFF_DB_MINUS_QPRINT"] = comp["CURRENT_DB_IBDEBVL_TOTAL"] - comp["QPRINT_TOTAL"]
            compare_parts.append(comp[["COMPARISON", "FAMILY_CODE", "FAMILY", "QPRINT_TOTAL", "CURRENT_DB_IBDEBVL_TOTAL", "DIFF_DB_MINUS_QPRINT"]])
        qprint_compare = pd.concat(compare_parts, ignore_index=True)

    loc_2025 = summaries_2025["loc_sum"][["LOCATION", "LOCATION_NAME", "INV_QTY", "BILL_INV", "BILL_COST", "AVG_COST"]]
    return {
        "Jacob Validation": jacob_validation,
        "2025 Loc Sum": loc_2025,
        "2024 PDF vs IBDEBVL": pdf_compare,
        "2025 QPRINT Sections": qprint_summary,
        "2025 QPRINT vs IBDEBVL": qprint_compare,
        "2025 IMVALUTOT": imvalutot_2025,
    }


def write_diagnostics_workbook(diagnostics: dict[str, pd.DataFrame], output_path: Path) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, df in diagnostics.items():
            safe_name = sheet_name[:31]
            df.to_excel(writer, sheet_name=safe_name, index=False)
            ws = writer.book[safe_name]
            ws.freeze_panes = "A2"
            for column_cells in ws.columns:
                max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
                ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 12), 60)
                for cell in column_cells:
                    if isinstance(cell.value, (int, float)) and "TOTAL" in str(ws.cell(1, cell.column).value).upper():
                        cell.number_format = "#,##0.00"
                    elif isinstance(cell.value, (int, float)) and str(ws.cell(1, cell.column).value).upper() in {
                        "BILL_COST",
                        "AVG_COST",
                        "CURRENT_IBDEBVL_TOTAL",
                        "CURRENT_DB_IBDEBVL_TOTAL",
                        "QPRINT_TOTAL",
                        "PDF_TOTAL",
                        "DIFF_CURRENT_MINUS_PDF",
                        "DIFF_DB_MINUS_QPRINT",
                    }:
                        cell.number_format = "#,##0.00"


def write_diagnostic_sql(output_path: Path) -> None:
    sql = f"""/*
Inventory report 2025 diagnostic and reproduction queries.

Validated Jacob workbook basis:
  Source: GSFL2K.ITEMBA_EOD
  Date: 2025-12-31 for the new report
  Company: 1
  Locations: all except 90
  INV_QTY: IBDBOH
  BILL_INV: IBDBOH * IMFACT, unless IMFACT = 0
  BILL_COST: BILL_INV * IBLCST

Legacy IMVALUTOT PDF basis:
  Source: GSFL2K.ITEMBA_EOD.IBDEBVL grouped by ITEMMAST/FAMILY.
  The COO-provided QPRINT for 2025 is authoritative for the generated
  Codex IMVALUTOT PDF/source. It contains two sections:
    section 1 family subtotal = $20,048,139.55 (loc <> 90)
    section 2 family subtotal = $20,238,507.18 (all locations)
  The printed section 2 total line is cumulative ($40,286,646.73), not the
  standalone all-location subtotal.
*/

-- 1) Confirm EOD snapshots exist.
SELECT
    MIN(IBEODDATE) AS MIN_DATE,
    MAX(IBEODDATE) AS MAX_DATE,
    SUM(CASE WHEN IBEODDATE = DATE('2024-12-31') THEN 1 ELSE 0 END) AS ROWS_ON_2024_1231,
    SUM(CASE WHEN IBEODDATE = DATE('2025-12-31') THEN 1 ELSE 0 END) AS ROWS_ON_2025_1231
FROM GSFL2K.ITEMBA_EOD
WHERE IBCO = 1;

-- 2) Jacob-style item/location detail for 2025.
SELECT
    IB.IBLOC AS LOCATION,
    TRIM(COALESCE(LC.LCRNAM, '')) AS LOCATION_NAME,
    TRIM(IB.IBITEM) AS ITEM,
    TRIM(IM.IMDESC) AS DESCRIPTION,
    TRIM(FM.FMDESC) AS FAMILY,
    TRIM(IM.IMUM1) AS INV_UNIT,
    SUM(IB.IBDBOH) AS INV_QTY,
    IM.IMFACT AS FACTOR,
    TRIM(IM.IMUM2) AS BILL_UNIT,
    CASE WHEN IM.IMFACT = 0 THEN SUM(IB.IBDBOH) ELSE SUM(IB.IBDBOH * IM.IMFACT) END AS BILL_INV,
    CASE WHEN IM.IMFACT = 0 THEN SUM(IB.IBDBOH * IB.IBLCST) ELSE SUM(IB.IBDBOH * IM.IMFACT * IB.IBLCST) END AS BILL_COST,
    CASE
        WHEN (CASE WHEN IM.IMFACT = 0 THEN SUM(IB.IBDBOH) ELSE SUM(IB.IBDBOH * IM.IMFACT) END) = 0 THEN 0
        ELSE
            (CASE WHEN IM.IMFACT = 0 THEN SUM(IB.IBDBOH * IB.IBLCST) ELSE SUM(IB.IBDBOH * IM.IMFACT * IB.IBLCST) END)
            /
            (CASE WHEN IM.IMFACT = 0 THEN SUM(IB.IBDBOH) ELSE SUM(IB.IBDBOH * IM.IMFACT) END)
    END AS AVG_COST
FROM GSFL2K.ITEMBA_EOD IB
JOIN GSFL2K.ITEMMAST IM
  ON TRIM(IB.IBITEM) = TRIM(IM.IMITEM)
JOIN GSFL2K.FAMILY FM
  ON TRIM(IM.IMFMCD) = TRIM(FM.FMFMCD)
LEFT JOIN GSFL2K.LOCATION LC
  ON LC.LCCO = IB.IBCO
 AND LC.LCLOC = IB.IBLOC
WHERE IB.IBEODDATE = DATE('2025-12-31')
  AND IB.IBCO = 1
  AND IB.IBLOC <> 90
GROUP BY
    IB.IBLOC,
    TRIM(COALESCE(LC.LCRNAM, '')),
    TRIM(IB.IBITEM),
    TRIM(IM.IMDESC),
    TRIM(FM.FMDESC),
    IM.IMFACT,
    TRIM(IM.IMUM1),
    TRIM(IM.IMUM2)
ORDER BY ITEM, LOCATION;

-- 3) Jacob-style location totals for 2025.
WITH INV_LOC AS (
    SELECT
        IB.IBLOC AS LOCATION,
        TRIM(COALESCE(LC.LCRNAM, '')) AS LOCATION_NAME,
        CASE WHEN IM.IMFACT = 0 THEN SUM(IB.IBDBOH * IB.IBLCST) ELSE SUM(IB.IBDBOH * IM.IMFACT * IB.IBLCST) END AS BILL_COST
    FROM GSFL2K.ITEMBA_EOD IB
    JOIN GSFL2K.ITEMMAST IM
      ON TRIM(IB.IBITEM) = TRIM(IM.IMITEM)
    LEFT JOIN GSFL2K.LOCATION LC
      ON LC.LCCO = IB.IBCO
     AND LC.LCLOC = IB.IBLOC
    WHERE IB.IBEODDATE = DATE('2025-12-31')
      AND IB.IBCO = 1
      AND IB.IBLOC <> 90
    GROUP BY IB.IBLOC, TRIM(COALESCE(LC.LCRNAM, '')), TRIM(IB.IBITEM), IM.IMFACT
)
SELECT LOCATION, LOCATION_NAME, SUM(BILL_COST) AS INVENTORY_VALUE
FROM INV_LOC
GROUP BY LOCATION, LOCATION_NAME
HAVING SUM(BILL_COST) <> 0
ORDER BY LOCATION;

-- 4) 2024-only investigation: rows now present in Gartman but absent from Jacob's 2024 workbook.
SELECT
    IB.IBLOC AS LOCATION,
    TRIM(IB.IBITEM) AS ITEM,
    TRIM(IM.IMDESC) AS DESCRIPTION,
    SUM(IB.IBDBOH) AS INV_QTY,
    SUM(CASE WHEN IM.IMFACT = 0 THEN IB.IBDBOH * IB.IBLCST ELSE IB.IBDBOH * IM.IMFACT * IB.IBLCST END) AS BILL_COST
FROM GSFL2K.ITEMBA_EOD IB
JOIN GSFL2K.ITEMMAST IM
  ON TRIM(IB.IBITEM) = TRIM(IM.IMITEM)
WHERE IB.IBEODDATE = DATE('2024-12-31')
  AND IB.IBCO = 1
  AND IB.IBLOC IN (1, 2, 3, 4)
  AND TRIM(IB.IBITEM) = 'RQ2142RS'
GROUP BY IB.IBLOC, TRIM(IB.IBITEM), TRIM(IM.IMDESC)
ORDER BY IB.IBLOC;

-- 5) Legacy IMVALUTOT-style all-location family totals from current DB.
--    This should be compared to COO QPRINT section 2 family subtotal.
SELECT
    TRIM(IM.IMFMCD) AS FAMILY_CODE,
    TRIM(FM.FMDESC) AS FAMILY,
    SUM(IB.IBDEBVL) AS TOTAL
FROM GSFL2K.ITEMBA_EOD IB
JOIN GSFL2K.ITEMMAST IM
  ON TRIM(IB.IBITEM) = TRIM(IM.IMITEM)
JOIN GSFL2K.FAMILY FM
  ON TRIM(IM.IMFMCD) = TRIM(FM.FMFMCD)
WHERE IB.IBEODDATE = DATE('2025-12-31')
  AND IB.IBCO = 1
GROUP BY TRIM(IM.IMFMCD), TRIM(FM.FMDESC)
HAVING SUM(IB.IBDEBVL) <> 0
ORDER BY TRIM(IM.IMFMCD);
"""
    output_path.write_text(sql, encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate 2025 inventory reports from validated Gartman basis.")
    parser.add_argument("--skip-pdf", action="store_true", help="Skip Excel COM PDF export for IMVALUTOT.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    print("Connecting to Gartman...")
    conn = connect()
    try:
        print("Pulling Jacob-style inventory detail for 2025...")
        inv_loc_2025 = load_jacob_inventory_by_loc(conn, AS_OF_DATE)
        summaries_2025 = build_jacob_summaries(inv_loc_2025)
        print(
            f"  detail rows={len(summaries_2025['by_loc']):,}; "
            f"item rows={len(summaries_2025['by_item']):,}; "
            f"value={money(summaries_2025['by_item']['BILL_COST'].sum())}"
        )

        print(f"Writing Jacob workbook: {OUTPUT_XLSX}")
        write_jacob_workbook(summaries_2025, OUTPUT_XLSX)
        recalc_workbook_with_excel(OUTPUT_XLSX)

        print("Pulling legacy IMVALUTOT-style family totals for 2025...")
        db_imvalutot_2025 = load_imvalutot_family(conn, AS_OF_DATE)
        imvalutot_2025, qprint_summary = parse_coo_qprint_2025(COO_QPRINT_2025) if COO_QPRINT_2025.exists() else (pd.DataFrame(), pd.DataFrame())
        if not imvalutot_2025.empty:
            print(
                f"  using COO QPRINT section {imvalutot_2025.attrs.get('section')}; "
                f"family rows={len(imvalutot_2025):,}; value={money(imvalutot_2025['TOTAL'].sum())}"
            )
            print(f"  current DB IBDEBVL all-location value={money(db_imvalutot_2025['TOTAL'].sum())}")
        else:
            imvalutot_2025 = db_imvalutot_2025
            print(f"  QPRINT unavailable; using current DB IBDEBVL. family rows={len(imvalutot_2025):,}; value={money(imvalutot_2025['TOTAL'].sum())}")

        print(f"Writing IMVALUTOT source workbook: {OUTPUT_IMVALUTOT_XLSX}")
        write_imvalutot_source_workbook(imvalutot_2025, OUTPUT_IMVALUTOT_XLSX)
        if not args.skip_pdf:
            print(f"Exporting IMVALUTOT PDF: {OUTPUT_IMVALUTOT_PDF}")
            export_pdf_with_excel(OUTPUT_IMVALUTOT_XLSX, OUTPUT_IMVALUTOT_PDF)

        print("Building diagnostics...")
        diagnostics = build_diagnostics(conn, summaries_2025, imvalutot_2025)
        write_diagnostics_workbook(diagnostics, OUTPUT_DIAGNOSTICS)
        write_diagnostic_sql(OUTPUT_SQL)
    finally:
        conn.close()

    print("Done.")
    print(f"Excel report: {OUTPUT_XLSX}")
    print(f"IMVALUTOT PDF: {OUTPUT_IMVALUTOT_PDF}")
    print(f"Diagnostics: {OUTPUT_DIAGNOSTICS}")
    print(f"SQL: {OUTPUT_SQL}")


if __name__ == "__main__":
    main()
