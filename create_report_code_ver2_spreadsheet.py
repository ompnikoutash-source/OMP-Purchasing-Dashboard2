from __future__ import annotations

import re
import sys
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PDF_PATH = Path(r"H:\2025\MISC Reports\For Tony\report code 2025 ver 2.pdf")
OUTPUT_PATH = PDF_PATH.with_suffix(".xlsx")
FALLBACK_OUTPUT_PATH = PDF_PATH.with_name(f"{PDF_PATH.stem} corrected.xlsx")

HEADERS = [
    "item number",
    "description",
    "loc",
    "serial/lot",
    "conversion",
    "bin",
    "on hand",
    "sold",
    "on order",
    "unit cost",
    "extended",
]

BOUNDS = {
    "left": (0.0, 306.0),
    "loc": (306.0, 343.0),
    "serial": (343.0, 451.8),
    "bin": (451.8, 484.2),
    "on_hand": (484.2, 532.8),
    "sold": (532.8, 581.4),
    "on_order": (581.4, 630.0),
    "unit_cost": (630.0, 684.0),
    "cost_type": (684.0, 694.8),
    "extended": (694.8, 1000.0),
}

LOC_RE = re.compile(r"^\d{3}/\d{2}$")
HEADER_FRAGMENT = "I_T_E_M__N_U_M_B_E_R"
NUMERIC_RE = re.compile(r"^(?:-?(?:\d[\d,]*\.\d+|\.\d+|\d+)|(?:\d[\d,]*\.\d+|\.\d+|\d+)-)$")
NUMERIC_COLUMN_CENTERS = {
    "on_hand": 505.8,
    "sold": 554.4,
    "on_order": 603.0,
    "unit_cost": 654.3,
    "extended": 721.8,
}


def classify_word(word: dict) -> str:
    x0 = round(float(word["x0"]), 1)
    xmid = round((float(word["x0"]) + float(word["x1"])) / 2, 1)
    text = str(word["text"])

    if LOC_RE.match(text):
        return "loc"

    if xmid < BOUNDS["serial"][1]:
        if x0 < BOUNDS["left"][1]:
            return "left"
        if x0 < BOUNDS["serial"][0]:
            return "loc"
        return "serial"

    if text in {"A", "B", "C", "D", "N"}:
        return "cost_type"

    if NUMERIC_RE.match(text):
        return min(
            NUMERIC_COLUMN_CENTERS,
            key=lambda name: abs(xmid - NUMERIC_COLUMN_CENTERS[name]),
        )

    for name, (low, high) in BOUNDS.items():
        if low <= x0 < high:
            return name
    return "other"


def group_lines(words: list[dict], tolerance: float = 2.0) -> list[list[dict]]:
    lines: list[list[dict]] = []
    line_tops: list[float] = []
    for word in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if not lines or abs(float(word["top"]) - line_tops[-1]) > tolerance:
            lines.append([word])
            line_tops.append(float(word["top"]))
        else:
            lines[-1].append(word)
    return [sorted(line, key=lambda w: w["x0"]) for line in lines]


def join_text(words: list[dict]) -> str:
    return " ".join(word["text"] for word in sorted(words, key=lambda w: w["x0"])).strip()


def parse_decimal(text: str) -> Decimal | None:
    if not text:
        return None
    clean = text.replace(",", "")
    is_trailing_negative = clean.endswith("-")
    if is_trailing_negative:
        clean = clean[:-1]
    if clean.startswith("."):
        clean = f"0{clean}"
    if clean.startswith("-."):
        clean = clean.replace("-.", "-0.", 1)
    try:
        value = Decimal(clean)
        return -value if is_trailing_negative else value
    except InvalidOperation:
        return None


def parse_serial(words: list[dict]) -> str:
    tokens = [word["text"] for word in sorted(words, key=lambda w: w["x0"])]
    if not tokens:
        return ""

    if len(tokens) == 2 and tokens[0].replace("-", "").isdigit() and tokens[1].isdigit():
        return f"{tokens[0]}-{tokens[1]}"

    return " ".join(tokens)


def extract_rows(pdf_path: Path) -> tuple[list[dict], Decimal | None]:
    rows: list[dict] = []
    current_item = ""
    current_desc = ""
    last_detail: dict | None = None
    grand_total: Decimal | None = None

    with pdfplumber.open(str(pdf_path)) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            words = page.extract_words(
                x_tolerance=1,
                y_tolerance=2,
                keep_blank_chars=False,
                use_text_flow=False,
            )

            for line_words in group_lines(words):
                text = join_text(line_words)
                if not text:
                    continue
                if text.startswith("IM150 S T O C K S T A T U S R E P O R T"):
                    continue
                if text == "TONY STOCK_STS":
                    continue
                if HEADER_FRAGMENT in text:
                    continue
                if text.startswith("** GRAND TOTAL"):
                    numbers = re.findall(r"\d[\d,]*\.\d{2}", text)
                    if numbers:
                        grand_total = parse_decimal(numbers[-1])
                    continue

                columns: dict[str, list[dict]] = defaultdict(list)
                for word in line_words:
                    columns[classify_word(word)].append(word)

                has_loc = any(LOC_RE.match(word["text"]) for word in columns["loc"])
                has_qty_or_value = any(
                    columns[name]
                    for name in ("on_hand", "sold", "on_order", "unit_cost", "extended")
                )

                if has_loc:
                    row = {
                        "item number": current_item,
                        "description": current_desc,
                        "loc": join_text([word for word in columns["loc"] if LOC_RE.match(word["text"])]),
                        "serial/lot": parse_serial(columns["serial"]),
                        "conversion": "",
                        "bin": join_text(columns["bin"]),
                        "on hand": parse_decimal(join_text(columns["on_hand"])),
                        "sold": parse_decimal(join_text(columns["sold"])),
                        "on order": parse_decimal(join_text(columns["on_order"])),
                        "unit cost": parse_decimal(join_text(columns["unit_cost"])),
                        "extended": parse_decimal(join_text(columns["extended"])),
                        "_page": page_number,
                        "_source_text": text,
                    }
                    rows.append(row)
                    last_detail = row
                    continue

                # A conversion row is the lone value printed under the serial/lot column.
                if (
                    last_detail is not None
                    and columns["serial"]
                    and not columns["left"]
                    and not columns["bin"]
                    and not has_qty_or_value
                    and not columns["cost_type"]
                ):
                    last_detail["conversion"] = parse_serial(columns["serial"])
                    continue

                if not has_qty_or_value and not columns["cost_type"]:
                    header_text = join_text(
                        columns["left"]
                        + columns["loc"]
                        + columns["serial"]
                        + columns["bin"]
                        + columns["other"]
                    )
                    if header_text:
                        parts = header_text.split(maxsplit=1)
                        current_item = parts[0]
                        current_desc = parts[1] if len(parts) > 1 else ""
                    continue

                # Remaining numeric-only lines are subtotal/total noise that the user asked to omit.
                continue

    return rows, grand_total


def decimal_to_excel(value: Decimal | None):
    return None if value is None else float(value)


def autosize_columns(worksheet) -> None:
    widths: dict[int, int] = {}
    for row in worksheet.iter_rows():
        for cell in row:
            value = "" if cell.value is None else str(cell.value)
            widths[cell.column] = max(widths.get(cell.column, 0), len(value))

    for column_index, width in widths.items():
        letter = get_column_letter(column_index)
        worksheet.column_dimensions[letter].width = min(max(width + 2, 10), 60)


def build_workbook(rows: list[dict], output_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "report"
    worksheet.append(HEADERS)

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row in rows:
        worksheet.append(
            [
                row["item number"],
                row["description"],
                row["loc"],
                row["serial/lot"],
                row["conversion"],
                row["bin"],
                decimal_to_excel(row["on hand"]),
                decimal_to_excel(row["sold"]),
                decimal_to_excel(row["on order"]),
                decimal_to_excel(row["unit cost"]),
                decimal_to_excel(row["extended"]),
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for row in worksheet.iter_rows(min_row=2):
        for cell in row[:6]:
            cell.number_format = "@"
        row[6].number_format = "0.00"
        row[7].number_format = "0.00"
        row[8].number_format = "0.00"
        row[9].number_format = "0.00000"
        row[10].number_format = "0.00"

    autosize_columns(worksheet)

    workbook.save(output_path)


def main() -> None:
    output_path = Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT_PATH
    rows, grand_total = extract_rows(PDF_PATH)
    try:
        build_workbook(rows, output_path)
    except PermissionError:
        if output_path != OUTPUT_PATH:
            raise
        output_path = FALLBACK_OUTPUT_PATH
        build_workbook(rows, output_path)

    parsed_total = sum((row["extended"] or Decimal("0")) for row in rows)
    print(f"Created: {output_path}")
    print(f"Rows: {len(rows)}")
    print(f"Parsed extended total: {parsed_total}")
    if grand_total is not None:
        print(f"PDF grand total: {grand_total}")
        print(f"Difference: {parsed_total - grand_total}")


if __name__ == "__main__":
    main()
