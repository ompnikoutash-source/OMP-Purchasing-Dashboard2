import openpyxl

wb = openpyxl.load_workbook("Forecast SKU List.xlsx")
ws = wb.active
print(f"Sheet name: {ws.title}")
print(f"Max row: {ws.max_row}")

# Read first 10 rows
print("\nFirst 10 entries from column A:")
for i in range(1, min(11, ws.max_row + 1)):
    val = ws[f"A{i}"].value
    print(f"Row {i}: {val}")

# Search for GFALPH SKUs
print("\nSearching for specific GFALPH SKUs...")
target_skus = ["GFALPH401L-10", "GFALPH401R-10", "GFALPH403L-10", "GFALPH403R-10", "GFALPH405L-10", "GFALPH405R-10"]
found = []
for row in ws.iter_rows(values_only=True):
    cell_val = str(row[0] or "").strip().upper()
    if any(sku in cell_val for sku in target_skus):
        found.append(cell_val)
        print(f"  FOUND: {cell_val}")

if not found:
    print("  No GFALPH SKUs found in the list")

# Count GFALPH entries
gfalph_count = 0
gfalph_examples = []
for row in ws.iter_rows(values_only=True):
    cell_val = str(row[0] or "").strip().upper()
    if "GFALPH" in cell_val:
        gfalph_count += 1
        if len(gfalph_examples) < 5:
            gfalph_examples.append(cell_val)

print(f"\nTotal GFALPH entries in list: {gfalph_count}")
if gfalph_examples:
    print("Sample GFALPH entries:")
    for ex in gfalph_examples:
        print(f"  {ex}")
