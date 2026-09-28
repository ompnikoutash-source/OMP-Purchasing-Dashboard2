"""
Export Gartman customer accounts within 30 miles of ZIP 93401.

The main workbook sheet contains the requested columns only:
account number, account name, address, city, state, zipcode, sales rep number,
and trailing-365 sales revenue.
"""

from __future__ import annotations

import contextlib
import io
import json
import math
import os
import sqlite3
import subprocess
from datetime import datetime
from pathlib import Path

import pandas as pd

from core.db_connection import connect


ORIGIN_QUERY = "93401, San Luis Obispo, CA"
ORIGIN_FALLBACK = (35.2828, -120.6596)
RADIUS_MILES = 30.0
OUTPUT_DIR = Path(__file__).resolve().parent / "account_radius_lists"
GEOCODE_DB_PATH = Path(__file__).resolve().parent / "geocode_cache.db"
ARCGIS_MIN_SCORE = 80


CUSTOMER_QUERY = """
WITH SALES_365 AS (
    SELECT
        TRIM(H.SHCUST) AS ACCOUNT_NUMBER,
        SUM(COALESCE(L.SLENET, 0)) AS SALES_REVENUE_365D
    FROM GSFL2K.SHHEAD H
    JOIN GSFL2K.SHLINE L
      ON L.SLCO   = H.SHCO
     AND L.SLLOC  = H.SHLOC
     AND L.SLINV# = H.SHINV#
     AND L.SLORD# = H.SHORD#
    WHERE H.SHIDAT >= (CURRENT_DATE - 365 DAYS)
    GROUP BY TRIM(H.SHCUST)
)
SELECT
    TRIM(C.CMCUST) AS ACCOUNT_NUMBER,
    TRIM(C.CMNAME) AS ACCOUNT_NAME,
    TRIM(C.CMADR1) AS ADDRESS_LINE_1,
    TRIM(C.CMADR2) AS ADDRESS_LINE_2,
    TRIM(C.CMADR3) AS CITY_STATE,
    C.CMZIP AS ZIPCODE_RAW,
    TRIM(CHAR(C.CMSLMN)) AS SALES_REP_NUMBER,
    COALESCE(S.SALES_REVENUE_365D, 0) AS SALES_REVENUE_365D
FROM GSFL2K.CUSTMAST C
LEFT JOIN SALES_365 S
  ON S.ACCOUNT_NUMBER = TRIM(C.CMCUST)
WHERE C.CMZIP >= 934000000
  AND C.CMZIP <  935000000
ORDER BY TRIM(C.CMCUST)
"""


def main() -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)

    print("Connecting to Gartman and pulling 934xx account candidates...")
    with contextlib.redirect_stdout(io.StringIO()):
        conn = connect()
    try:
        db_today = pd.read_sql(
            "SELECT CURRENT DATE AS TODAY FROM SYSIBM.SYSDUMMY1",
            conn,
        ).loc[0, "TODAY"]
        df = pd.read_sql(CUSTOMER_QUERY, conn)
    finally:
        conn.close()

    print(f"Fetched {len(df)} candidate account rows.")
    df = _normalize_customer_rows(df)

    cache_conn = _init_geocode_db()
    try:
        origin_coords = _quiet_geocode(ORIGIN_QUERY, cache_conn) or ORIGIN_FALLBACK
        print(
            "Using origin "
            f"{ORIGIN_QUERY}: {origin_coords[0]:.6f}, {origin_coords[1]:.6f}"
        )
        audit_df = _add_distance_columns(df, origin_coords, cache_conn)
    finally:
        cache_conn.close()

    within_df = audit_df[
        audit_df["DISTANCE_MILES"].notna()
        & (audit_df["DISTANCE_MILES"] <= RADIUS_MILES)
    ].copy()

    within_df.sort_values(
        ["SALES_REP_NUMBER", "SALES_REVENUE_365D", "ACCOUNT_NAME"],
        ascending=[True, False, True],
        inplace=True,
    )

    main_df = within_df[
        [
            "ACCOUNT_NUMBER",
            "ACCOUNT_NAME",
            "ADDRESS",
            "CITY",
            "STATE",
            "ZIPCODE",
            "SALES_REP_NUMBER",
            "SALES_REVENUE_365D",
        ]
    ].copy()
    main_df.rename(
        columns={
            "ACCOUNT_NUMBER": "Account Number",
            "ACCOUNT_NAME": "Account Name",
            "ADDRESS": "Address",
            "CITY": "City",
            "STATE": "State",
            "ZIPCODE": "Zipcode",
            "SALES_REP_NUMBER": "Sales Rep Number",
            "SALES_REVENUE_365D": "Sales Revenue Past 365 Days",
        },
        inplace=True,
    )

    rep_summary_df = (
        main_df.groupby("Sales Rep Number", dropna=False, as_index=False)
        .agg(
            Accounts=("Account Number", "count"),
            Sales_Revenue_Past_365_Days=(
                "Sales Revenue Past 365 Days",
                "sum",
            ),
        )
        .sort_values("Sales_Revenue_Past_365_Days", ascending=False)
    )
    rep_summary_df.rename(
        columns={"Sales_Revenue_Past_365_Days": "Sales Revenue Past 365 Days"},
        inplace=True,
    )

    audit_out_df = audit_df.copy()
    audit_out_df["WITHIN_30_MILES"] = audit_out_df["DISTANCE_MILES"] <= RADIUS_MILES
    audit_out_df = audit_out_df[
        [
            "ACCOUNT_NUMBER",
            "ACCOUNT_NAME",
            "ADDRESS",
            "CITY",
            "STATE",
            "ZIPCODE",
            "SALES_REP_NUMBER",
            "SALES_REVENUE_365D",
            "DISTANCE_MILES",
            "GEOCODE_USED",
            "WITHIN_30_MILES",
        ]
    ].sort_values(["WITHIN_30_MILES", "DISTANCE_MILES"], ascending=[False, True])

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx_path = OUTPUT_DIR / f"slo_93401_30mi_accounts_{timestamp}.xlsx"
    csv_path = OUTPUT_DIR / f"slo_93401_30mi_accounts_{timestamp}.csv"

    main_df.to_csv(csv_path, index=False)
    _write_workbook(
        xlsx_path=xlsx_path,
        accounts_df=main_df,
        rep_summary_df=rep_summary_df,
        audit_df=audit_out_df,
        db_today=str(pd.Timestamp(db_today).date()),
        origin_coords=origin_coords,
    )

    print(f"Accounts within {RADIUS_MILES:.0f} miles: {len(main_df)}")
    print(f"Workbook: {xlsx_path}")
    print(f"CSV: {csv_path}")


def _normalize_customer_rows(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in [
        "ACCOUNT_NUMBER",
        "ACCOUNT_NAME",
        "ADDRESS_LINE_1",
        "ADDRESS_LINE_2",
        "CITY_STATE",
        "SALES_REP_NUMBER",
    ]:
        out[col] = out[col].fillna("").astype(str).str.strip()

    parsed = out["CITY_STATE"].apply(_parse_city_state)
    out["CITY"] = [item[0] for item in parsed]
    out["STATE"] = [item[1] for item in parsed]
    out["ZIPCODE"] = out["ZIPCODE_RAW"].apply(_format_zip)
    out["ADDRESS"] = [
        _join_nonblank([row.ADDRESS_LINE_1, row.ADDRESS_LINE_2])
        for row in out.itertuples(index=False)
    ]
    out["FULL_ADDRESS"] = [
        _join_nonblank([row.ADDRESS, row.CITY, row.STATE, row.ZIPCODE])
        for row in out.itertuples(index=False)
    ]
    out["SALES_REVENUE_365D"] = pd.to_numeric(
        out["SALES_REVENUE_365D"],
        errors="coerce",
    ).fillna(0.0)
    return out


def _add_distance_columns(
    df: pd.DataFrame,
    origin_coords: tuple[float, float],
    cache_conn,
) -> pd.DataFrame:
    rows = []
    total = len(df)
    for idx, row in enumerate(df.itertuples(index=False), start=1):
        if idx == 1 or idx % 10 == 0 or idx == total:
            print(f"Geocoding {idx}/{total}...")

        coords = _quiet_geocode(row.FULL_ADDRESS, cache_conn)
        geocode_used = "address"
        if coords is None:
            fallback_query = _join_nonblank([row.CITY, row.STATE, row.ZIPCODE])
            coords = _quiet_geocode(fallback_query, cache_conn)
            geocode_used = "city/state/zip" if coords is not None else "failed"

        data = row._asdict()
        if coords is None:
            data["DISTANCE_MILES"] = math.nan
            data["GEOCODE_USED"] = geocode_used
        else:
            data["DISTANCE_MILES"] = round(_haversine_miles(origin_coords, coords), 2)
            data["GEOCODE_USED"] = geocode_used
        rows.append(data)

    return pd.DataFrame(rows)


def _write_workbook(
    xlsx_path: Path,
    accounts_df: pd.DataFrame,
    rep_summary_df: pd.DataFrame,
    audit_df: pd.DataFrame,
    db_today: str,
    origin_coords: tuple[float, float],
) -> None:
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        accounts_df.to_excel(writer, sheet_name="Accounts", index=False)
        rep_summary_df.to_excel(writer, sheet_name="Rep Summary", index=False)
        audit_df.to_excel(writer, sheet_name="Radius Audit", index=False)

        summary_df = pd.DataFrame(
            [
                ("Gartman DB Date", db_today),
                ("Origin", ORIGIN_QUERY),
                ("Origin Latitude", round(origin_coords[0], 6)),
                ("Origin Longitude", round(origin_coords[1], 6)),
                ("Radius Miles", RADIUS_MILES),
                ("Accounts Within Radius", len(accounts_df)),
            ],
            columns=["Metric", "Value"],
        )
        summary_df.to_excel(writer, sheet_name="Run Summary", index=False)

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            _format_sheet(ws)


def _format_sheet(ws) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for col_cells in ws.columns:
        header = str(col_cells[0].value or "")
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells)
        col_letter = col_cells[0].column_letter
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, len(header) + 2), 48)
        if "Sales Revenue" in header:
            for cell in col_cells[1:]:
                cell.number_format = '$#,##0.00;[Red]-$#,##0.00'
        if header == "DISTANCE_MILES":
            for cell in col_cells[1:]:
                cell.number_format = '0.00'


def _quiet_geocode(full_address: str, cache_conn) -> tuple[float, float] | None:
    return _geocode_address(full_address, cache_conn)


def _init_geocode_db() -> sqlite3.Connection:
    conn = sqlite3.connect(str(GEOCODE_DB_PATH))
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS geocode_cache (
            address_key  TEXT PRIMARY KEY,
            lat          REAL,
            lon          REAL,
            display_name TEXT,
            failed       INTEGER DEFAULT 0,
            queried_at   TEXT DEFAULT (datetime('now'))
        )
        """
    )
    conn.commit()
    return conn


def _geocode_address(
    full_address: str,
    cache_conn: sqlite3.Connection,
) -> tuple[float, float] | None:
    if not str(full_address or "").strip():
        return None

    key = str(full_address).strip().upper()
    row = cache_conn.execute(
        "SELECT lat, lon, failed FROM geocode_cache WHERE address_key = ?",
        (key,),
    ).fetchone()
    if row:
        lat, lon, failed = row
        if not failed and lat is not None and lon is not None:
            return float(lat), float(lon)
        return None

    result = _arcgis_geocode_via_powershell(full_address)
    if result is None:
        cache_conn.execute(
            """
            INSERT OR REPLACE INTO geocode_cache
                (address_key, lat, lon, display_name, failed, queried_at)
            VALUES (?, NULL, NULL, NULL, 1, datetime('now'))
            """,
            (key,),
        )
        cache_conn.commit()
        return None

    lat, lon, display_name = result
    cache_conn.execute(
        """
        INSERT OR REPLACE INTO geocode_cache
            (address_key, lat, lon, display_name, failed, queried_at)
        VALUES (?, ?, ?, ?, 0, datetime('now'))
        """,
        (key, lat, lon, display_name),
    )
    cache_conn.commit()
    return float(lat), float(lon)


def _arcgis_geocode_via_powershell(full_address: str) -> tuple[float, float, str] | None:
    script = r"""
$ErrorActionPreference = 'Stop'
$encoded = [System.Uri]::EscapeDataString($env:GEOCODE_SINGLELINE)
$uri = "https://geocode.arcgis.com/arcgis/rest/services/World/GeocodeServer/findAddressCandidates?SingleLine=$encoded&f=json&outFields=&maxLocations=1&countryCode=USA"
$response = Invoke-RestMethod -Uri $uri -TimeoutSec 20
$response | ConvertTo-Json -Depth 10 -Compress
"""
    env = os.environ.copy()
    env["GEOCODE_SINGLELINE"] = str(full_address)
    try:
        proc = subprocess.run(
            ["powershell", "-NoProfile", "-Command", script],
            capture_output=True,
            env=env,
            text=True,
            timeout=30,
        )
    except Exception:
        return None

    if proc.returncode != 0 or not proc.stdout.strip():
        return None

    try:
        data = json.loads(proc.stdout)
    except json.JSONDecodeError:
        return None

    candidates = data.get("candidates", [])
    if not candidates:
        return None

    candidate = candidates[0]
    if float(candidate.get("score", 0)) < ARCGIS_MIN_SCORE:
        return None

    location = candidate.get("location") or {}
    lat = location.get("y")
    lon = location.get("x")
    if lat is None or lon is None:
        return None

    return float(lat), float(lon), str(candidate.get("address", "ArcGIS"))


def _parse_city_state(value: str) -> tuple[str, str]:
    parts = str(value or "").strip().split()
    if len(parts) >= 2 and len(parts[-1]) == 2 and parts[-1].isalpha():
        return " ".join(parts[:-1]).strip(), parts[-1].upper()
    return " ".join(parts).strip(), ""


def _format_zip(value) -> str:
    if value is None or value == "":
        return ""
    try:
        text = str(int(float(value)))
    except (TypeError, ValueError):
        text = str(value).strip()
    if not text or text == "0":
        return ""
    if len(text) >= 5:
        return text[:5]
    return text.zfill(5)


def _join_nonblank(parts: list[str]) -> str:
    return ", ".join(str(part).strip() for part in parts if str(part).strip())


def _haversine_miles(
    coord_a: tuple[float, float],
    coord_b: tuple[float, float],
) -> float:
    lat1, lon1 = coord_a
    lat2, lon2 = coord_b
    radius_miles = 3958.7613

    phi1 = math.radians(lat1)
    phi2 = math.radians(lat2)
    d_phi = math.radians(lat2 - lat1)
    d_lambda = math.radians(lon2 - lon1)

    h = (
        math.sin(d_phi / 2) ** 2
        + math.cos(phi1) * math.cos(phi2) * math.sin(d_lambda / 2) ** 2
    )
    return 2 * radius_miles * math.atan2(math.sqrt(h), math.sqrt(1 - h))


if __name__ == "__main__":
    main()
