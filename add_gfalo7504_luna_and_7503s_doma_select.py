from __future__ import annotations

from datetime import datetime
from pathlib import Path
import shutil
import sys

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from add_gfalo7506s_strada_select import ensure_columns, header_map, set_row
from refresh_garrison_color_rollups import (
    refresh_buckets,
    refresh_summary,
    set_gfalo7508_bucket_and_dimensions,
)
from update_gfalo7508_from_image import image_metrics


WORKBOOK = Path(r"H:\2025\MISC Reports\For Marketing\Garrison Color Analysis 2.12.2026.xlsx")
PENDING_WORKBOOK = WORKBOOK.with_name(f"{WORKBOOK.stem} pending GFALO7504 GFALO7503S update{WORKBOOK.suffix}")
BACKUP = WORKBOOK.with_name(
    f"{WORKBOOK.stem} backup before GFALO7504 GFALO7503S {datetime.now():%Y%m%d_%H%M%S}{WORKBOOK.suffix}"
)


PRODUCTS = [
    {
        "ProductNumber": "GFALO7504",
        "Product": "Luna",
        "Grade": "Select Character",
        "ImagePath": Path(r"H:\2025\MISC Reports\For Marketing\Overheads for Color Analysis\European Oak Luna 7.5.webp"),
        "HueBucket": 3,
        "BucketName": "Honey Blonde",
        "HueFamily": "Golden/Honey",
        "Undertone": "Neutral",
        "Description": 'EURO OAK 7.5" LUNA',
    },
    {
        "ProductNumber": "GFALO7503S",
        "Product": "Doma Select",
        "Grade": "Select",
        "ImagePath": Path(r"H:\2025\MISC Reports\For Marketing\Overheads for Color Analysis\European Oak Doma Select 7.5.webp"),
        "HueBucket": 5,
        "BucketName": "Golden Oak",
        "HueFamily": "Golden/Honey",
        "Undertone": "Warm",
        "Description": 'EURO OAK 7.5" DOMA SELECT',
    },
]


def next_index(ws) -> int:
    return max(cell.value for cell in ws["A"][1:] if isinstance(cell.value, (int, float))) + 1


def find_product_row(ws, product_number: str) -> int | None:
    headers = header_map(ws)
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, headers["ProductNumber"]).value or "").strip().upper() == product_number:
            return row
    return None


def main():
    shutil.copy2(WORKBOOK, BACKUP)
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb["Products"]
    ensure_columns(ws)
    headers = header_map(ws)

    written = []
    for product in PRODUCTS:
        metrics = image_metrics(product["ImagePath"])
        row_idx = find_product_row(ws, product["ProductNumber"])
        if row_idx is None:
            row_idx = ws.max_row + 1
            index_value = next_index(ws)
        else:
            index_value = ws.cell(row_idx, headers["INDEX"]).value

        values = {
            "INDEX": index_value,
            "Collection": "Allora",
            "Product": product["Product"],
            "Width_in": 7.5,
            "HueBucket": product["HueBucket"],
            "BucketName": product["BucketName"],
            "HueFamily": product["HueFamily"],
            "Undertone": product["Undertone"],
            "Notes": f'Gartman item description: {product["Description"]}. Metrics computed from source image file.',
            "ProductNumber": product["ProductNumber"],
            "Species": "European Oak",
            "Construction": "Engineered (Birch Plywood)",
            "Grade": product["Grade"],
            "Finish": "Matte UV Lacquer",
            "SurfaceTreatment": "Light Wire-Brush",
            "JankaRating": 1360,
            "Thickness": '5/8"',
            "TopLayer": "4mm",
            "Length": "Mostly 7'8 1/2\"",
            "InstallationMethods": "Glue Down, Floating",
            "EdgeDetail": "Micro-Bevel",
            "Warranty": "Residential: Limited Lifetime Structure, Limited 25 Year Finish | Commercial: Limited 25 Year Structure, Limited 10 Year Finish",
            "ImageAnalysisSource": str(product["ImagePath"]),
        }
        values.update(metrics)
        set_row(ws, row_idx, values)
        written.append((product, row_idx, metrics))

    set_gfalo7508_bucket_and_dimensions(wb)
    refresh_buckets(wb)
    refresh_summary(wb)

    save_target = WORKBOOK
    try:
        wb.save(save_target)
    except PermissionError:
        save_target = PENDING_WORKBOOK
        wb.save(save_target)

    print(f"Updated: {save_target}")
    if save_target != WORKBOOK:
        print(f"Original locked: {WORKBOOK}")
    print(f"Backup:  {BACKUP}")
    for product, row_idx, metrics in written:
        print()
        print(f"{product['ProductNumber']} / {product['Product']}")
        print(f"Row:     {row_idx}")
        print(f"Bucket:  {product['HueBucket']} - {product['BucketName']}")
        for key in [
            "MeanHex",
            "MedianHex",
            "L_mean",
            "HueDeg",
            "Chroma",
            "Sat_mean",
            "Value_mean",
            "Contrast_L_std",
            "EdgeDensity",
            "WarmthScore",
            "ImageWidthPx",
            "ImageHeightPx",
        ]:
            print(f"{key}: {metrics[key]}")


if __name__ == "__main__":
    main()
