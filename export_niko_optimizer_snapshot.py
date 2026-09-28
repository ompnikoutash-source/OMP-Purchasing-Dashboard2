from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from core.niko_pricing import (
    DEFAULT_NIKO_VENDOR_ALIASES,
    DEFAULT_STRIP_VENDOR_COLUMNS,
    coerce_positive_float,
    load_niko_pricing_inputs,
    lookup_prices_for_sku,
)
from core.vendor_optimizer import (
    DEFAULT_PALLET_SF,
    MAX_PALLET_UPLIFT_PCT,
    MIN_FULL_PALLET_QTY_SF,
    MIN_PURCHASABLE_QTY_SF,
    TRUCKLOAD_CAPACITY_SF,
    PurchaseRequest,
    VendorQuote,
    apply_minimum_order_policy,
    normalize_item_key,
    optimize_vendor_mix,
)


REPO_DIR = Path(__file__).resolve().parent
DEFAULT_PAYLOAD_PATH = REPO_DIR / "flooringwebappJSON"
DEFAULT_PRICING_PATH = REPO_DIR / "Unfinished Pricing.xlsx"
DEFAULT_STRIP_PATH = REPO_DIR / "StripSKUList.xlsx"
TRUCK_COLOR_PALETTE = [
    "D9F99D",
    "FDE68A",
    "BFDBFE",
    "FBCFE8",
    "C7F9CC",
    "FECACA",
]
HEADER_FILL = PatternFill(fill_type="solid", fgColor="BDD7EE")
SECTION_FILL = PatternFill(fill_type="solid", fgColor="DCE6F1")
WINNER_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _build_sku_consolidation_groups(strip_path: Path) -> Dict[str, List[str]]:
    try:
        df = pd.read_excel(strip_path, sheet_name="Strip SKU Combination", header=0)
    except Exception:
        return {}
    if df.empty:
        return {}

    index_col = df.columns[0]
    sku_col = df.columns[1]
    df[index_col] = pd.to_numeric(df[index_col], errors="coerce")
    df[sku_col] = df[sku_col].astype(str).str.strip().str.upper()

    groups: Dict[str, List[str]] = {}
    for _, grp in df.dropna(subset=[index_col]).groupby(index_col):
        skus = [sku for sku in grp[sku_col].tolist() if sku and sku != "NAN"]
        if len(skus) >= 2:
            groups["/".join(skus)] = skus
    return groups


def _load_strip_sku_details(strip_path: Path) -> Tuple[set[str], Dict[str, Dict[str, Any]]]:
    try:
        df = pd.read_excel(strip_path, sheet_name="Who Produces", header=0)
    except Exception:
        return set(), {}

    details: Dict[str, Dict[str, Any]] = {}
    strip_skus: set[str] = set()
    for _, row in df.iterrows():
        sku = str(row.get("Item Number", "")).strip().upper()
        if not sku or sku == "NAN":
            continue
        strip_skus.add(sku)
        details[sku] = {
            "Description": str(row.get("Description", "")).strip(),
            "Thickness": str(row.get("Thickness", "")).strip(),
            "Width": str(row.get("Width", "")).strip(),
            "Species": str(row.get("Species", "")).strip(),
            "Grade/Cut": str(row.get("Grade/Cut", "")).strip(),
            "Edge": str(row.get("Edge", "")).strip(),
            "Bundles": str(row.get("Bundles", "")).strip(),
        }
    return strip_skus, details


def _consolidate_monthly_projections(monthly_rows: List[Dict[str, Any]], groups: Dict[str, List[str]]) -> List[Dict[str, Any]]:
    if not monthly_rows:
        return []

    sku_month_rows: Dict[Tuple[str, str], List[Dict[str, Any]]] = {}
    all_months: set[str] = set()
    for row in monthly_rows:
        sku = str(row.get("SKU", row.get("sku", ""))).strip().upper()
        month = str(row.get("Month", ""))
        sku_month_rows.setdefault((sku, month), []).append(row)
        if month:
            all_months.add(month)

    consolidated_skus: set[str] = set()
    consolidated_rows: List[Dict[str, Any]] = []
    for consolidated_name, member_skus in groups.items():
        member_skus_upper = [member.upper() for member in member_skus]
        has_members = any((sku, month) in sku_month_rows for sku in member_skus_upper for month in all_months)
        if not has_members:
            continue
        consolidated_skus.update(member_skus_upper)
        for month in all_months:
            member_month_rows: List[Dict[str, Any]] = []
            for member_sku in member_skus_upper:
                member_month_rows.extend(sku_month_rows.get((member_sku, month), []))
            if not member_month_rows:
                continue
            base_row = member_month_rows[0]
            consolidated_rows.append(
                {
                    "SKU": consolidated_name,
                    "sku": consolidated_name,
                    "Month": base_row.get("Month"),
                    "Row_Type": base_row.get("Row_Type", "FCST"),
                    "Forecast": sum(float(r.get("Forecast", 0) or 0) for r in member_month_rows),
                    "Beginning Inventory": sum(float(r.get("Beginning Inventory", 0) or 0) for r in member_month_rows),
                    "Ending Inventory": sum(float(r.get("Ending Inventory", 0) or 0) for r in member_month_rows),
                    "Order Quantity": sum(float(r.get("Order Quantity", 0) or 0) for r in member_month_rows),
                    "vendor_name": base_row.get("vendor_name", ""),
                    "collection": base_row.get("collection", ""),
                    "description": base_row.get("description", ""),
                    "sf_per_pallet": base_row.get("sf_per_pallet"),
                }
            )

    result = []
    for row in monthly_rows:
        sku = str(row.get("SKU", row.get("sku", ""))).strip().upper()
        if sku not in consolidated_skus:
            result.append(row)

    result.extend(consolidated_rows)
    return result


def _build_group_descriptions(
    groups: Dict[str, List[str]],
    strip_details: Dict[str, Dict[str, Any]],
) -> Dict[str, str]:
    descriptions: Dict[str, str] = {}
    for consolidated_name, members in groups.items():
        for member in members:
            desc = str(strip_details.get(member.upper(), {}).get("Description", "")).strip()
            if desc:
                descriptions[consolidated_name] = desc
                break
    return descriptions


def _build_optimizer_rows(
    payload: Dict[str, Any],
    strip_skus: set[str],
    strip_details: Dict[str, Dict[str, Any]],
    groups: Dict[str, List[str]],
    group_descriptions: Dict[str, str],
    pricing_inputs: Dict[str, Any],
) -> Tuple[pd.Timestamp, List[Dict[str, Any]], List[str]]:
    monthly_rows = payload.get("Monthly_Projections", []) or []
    strip_monthly = [
        row for row in monthly_rows
        if str(row.get("SKU", "")).strip().upper() in strip_skus
    ]
    consolidated = _consolidate_monthly_projections(strip_monthly, groups)
    df = pd.DataFrame(consolidated)
    if df.empty:
        return pd.Timestamp.now().to_period("M").to_timestamp(), [], ["No strip monthly projections were found in the saved dashboard payload."]

    run_stamp = payload.get("run_meta", {}).get("run_timestamp_local")
    run_month = pd.Timestamp(run_stamp).to_period("M").to_timestamp()
    df["Month"] = pd.to_datetime(df["Month"], errors="coerce")
    df = df[df["Month"] == run_month]
    df = df[pd.to_numeric(df["Order Quantity"], errors="coerce") > 0]

    if df.empty:
        return run_month, [], ["The saved dashboard payload has no current-month reorder quantities for the strip tab."]

    pricing_map = pricing_inputs.get("prices_by_sku", {}) or {}
    sku_to_consolidated = {
        normalize_item_key(member): group_key
        for group_key, members in groups.items()
        for member in members
    }

    optimizer_rows: List[Dict[str, Any]] = []
    warnings_out: List[str] = []
    for _, row in df.sort_values(["SKU"]).iterrows():
        sku = str(row.get("SKU", "")).strip()
        desc = str(row.get("description", "")).strip()
        if not desc:
            desc = group_descriptions.get(normalize_item_key(sku), "")
        if not desc:
            desc = strip_details.get(sku.upper(), {}).get("Description", "")

        qty_numeric = coerce_positive_float(row.get("Order Quantity"))
        if qty_numeric is None or qty_numeric <= 0:
            continue

        sf_per_pallet = coerce_positive_float(row.get("sf_per_pallet")) or DEFAULT_PALLET_SF
        prices = lookup_prices_for_sku(
            sku,
            pricing_map,
            sku_to_consolidated=sku_to_consolidated,
            consolidation_groups=groups,
        )

        if not prices:
            warnings_out.append(f"No vendor pricing found for {normalize_item_key(sku)}.")

        optimizer_rows.append(
            {
                "sku": sku,
                "description": desc,
                "quantity": float(qty_numeric),
                "sf_per_pallet": float(sf_per_pallet),
                "prices": prices,
            }
        )

    return run_month, optimizer_rows, warnings_out


def _run_optimizer(
    optimizer_rows: List[Dict[str, Any]],
    vendor_names: List[str],
    freight_costs: Dict[str, float],
) -> Tuple[Dict[str, Any], List[str], List[Dict[str, Any]]]:
    raw_requests: List[PurchaseRequest] = []
    missing_quotes: List[str] = []

    for row in optimizer_rows:
        sku = str(row.get("sku", "")).strip()
        qty = float(row.get("quantity") or 0.0)
        if not sku or qty <= 0:
            continue
        if qty < MIN_PURCHASABLE_QTY_SF:
            raw_requests.append(
                PurchaseRequest(
                    sku=sku,
                    qty_sf=qty,
                    vendor_quotes={},
                    description=str(row.get("description", "")).strip(),
                    pallet_sf=float(row.get("sf_per_pallet") or DEFAULT_PALLET_SF),
                )
            )
            continue

        quotes: Dict[str, VendorQuote] = {}
        for vendor in vendor_names:
            price = coerce_positive_float((row.get("prices") or {}).get(vendor))
            if price is None:
                continue
            quotes[vendor] = VendorQuote(
                vendor=vendor,
                material_price=float(price),
                freight_per_truck=float(freight_costs.get(vendor, 0.0) or 0.0),
            )

        if not quotes:
            missing_quotes.append(normalize_item_key(sku))
            continue

        raw_requests.append(
            PurchaseRequest(
                sku=sku,
                qty_sf=qty,
                vendor_quotes=quotes,
                description=str(row.get("description", "")).strip(),
                pallet_sf=float(row.get("sf_per_pallet") or DEFAULT_PALLET_SF),
            )
        )

    requests, omitted_items = apply_minimum_order_policy(
        raw_requests,
        omit_below_sf=MIN_PURCHASABLE_QTY_SF,
        round_up_above_sf=MIN_PURCHASABLE_QTY_SF,
        minimum_order_sf=MIN_FULL_PALLET_QTY_SF,
    )
    result = optimize_vendor_mix(
        requests,
        truck_capacity_sf=TRUCKLOAD_CAPACITY_SF,
        allow_pallet_uplift=True,
        max_uplift_pct=MAX_PALLET_UPLIFT_PCT,
        default_pallet_sf=DEFAULT_PALLET_SF,
    )
    result["omitted_items"] = omitted_items
    return result, sorted(set(missing_quotes)), omitted_items


def _style_table_cell(cell, header: bool = False, fill: PatternFill | None = None) -> None:
    cell.border = THIN_BORDER
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if header:
        cell.font = Font(bold=True, color="1F1F1F")
        cell.fill = fill or HEADER_FILL
    elif fill is not None:
        cell.fill = fill


def _write_summary_sheet(
    ws,
    run_meta: Dict[str, Any],
    run_month: pd.Timestamp,
    pricing_path: Path,
    payload_path: Path,
    pricing_inputs: Dict[str, Any],
    optimizer_rows: List[Dict[str, Any]],
    optimization_result: Dict[str, Any],
    missing_quotes: List[str],
    warnings_out: List[str],
) -> None:
    ws.title = "Summary"
    ws["A1"] = "Niko Optimizer Snapshot"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"Saved dashboard run: {run_meta.get('run_timestamp_local', '--')}"
    ws["A3"] = f"Reorder month: {run_month.strftime('%B %Y')}"
    ws["A4"] = f"Pricing workbook: {pricing_path.name}"
    ws["A5"] = f"Payload source: {payload_path.name}"
    ws["A6"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

    summary_rows = [
        ("Quoted reorder items", sum(1 for row in optimizer_rows if row.get("prices"))),
        ("Total reorder items", len(optimizer_rows)),
        ("Total base reorder SF", sum(float(row.get("quantity", 0.0) or 0.0) for row in optimizer_rows)),
        ("Total added pallet SF", float(optimization_result.get("total_added_qty", 0.0) or 0.0)),
        ("Material total", float(optimization_result.get("total_material_cost", 0.0) or 0.0)),
        ("Freight total", float(optimization_result.get("total_freight_cost", 0.0) or 0.0)),
        ("Landed total", float(optimization_result.get("total_cost", 0.0) or 0.0)),
    ]

    start_row = 8
    ws.cell(start_row, 1, "Metric")
    ws.cell(start_row, 2, "Value")
    _style_table_cell(ws.cell(start_row, 1), header=True)
    _style_table_cell(ws.cell(start_row, 2), header=True)
    for offset, (label, value) in enumerate(summary_rows, start=1):
        row_idx = start_row + offset
        ws.cell(row_idx, 1, label)
        value_cell = ws.cell(row_idx, 2, value)
        _style_table_cell(ws.cell(row_idx, 1))
        _style_table_cell(value_cell)
        if label.endswith("SF"):
            value_cell.number_format = "#,##0.00"
        elif "total" in label.lower():
            value_cell.number_format = "$#,##0.00"

    note_row = start_row + len(summary_rows) + 3
    ws.cell(note_row, 1, "Notes")
    _style_table_cell(ws.cell(note_row, 1), header=True, fill=SECTION_FILL)
    notes = []
    notes.extend(pricing_inputs.get("warnings", []) or [])
    notes.extend(warnings_out)
    if missing_quotes:
        notes.append("Missing vendor quotes: " + ", ".join(missing_quotes))
    if optimization_result.get("errors"):
        notes.extend(optimization_result.get("errors", []))
    if not notes:
        notes.append("No warnings.")

    for idx, note in enumerate(notes, start=1):
        ws.cell(note_row + idx, 1, note)
        _style_table_cell(ws.cell(note_row + idx, 1))

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18


def _write_optimizer_card_sheet(
    ws,
    run_month: pd.Timestamp,
    vendor_names: List[str],
    optimizer_rows: List[Dict[str, Any]],
    optimization_result: Dict[str, Any],
) -> None:
    ws.title = "Optimizer Card"
    ws["A1"] = "Optimizer Card Snapshot"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = (
        f"Truck capacity {TRUCKLOAD_CAPACITY_SF:,.0f} SF | "
        f"Pallet uplift cap {MAX_PALLET_UPLIFT_PCT * 100:.0f}% | "
        f"Reorder month {run_month.strftime('%B %Y')}"
    )

    vendor_summary = optimization_result.get("vendor_summary", {}) or {}
    item_costs = {
        normalize_item_key(sku): value
        for sku, value in (optimization_result.get("item_costs") or {}).items()
    }

    header_row = 4
    ws.cell(header_row, 1, "Cost Breakdown")
    _style_table_cell(ws.cell(header_row, 1), header=True, fill=SECTION_FILL)
    for col_idx, vendor in enumerate(vendor_names, start=2):
        cell = ws.cell(header_row, col_idx, vendor)
        _style_table_cell(cell, header=True)

    cost_metrics = [
        ("Freight / Truck", "freight_per_truck", "$#,##0.00"),
        ("Truckloads", "truckloads", "0"),
        ("Vendor Total SF", "qty_sf", "#,##0"),
        ("Truck Fill %", "truck_fill_pct", "0.0%"),
        ("Material Total", "material_total", "$#,##0.00"),
        ("Freight Total", "freight_total", "$#,##0.00"),
        ("Landed Total", "landed_total", "$#,##0.00"),
        ("Added SF", "added_qty_sf", "#,##0"),
    ]
    for row_offset, (label, key, num_fmt) in enumerate(cost_metrics, start=1):
        row_idx = header_row + row_offset
        label_cell = ws.cell(row_idx, 1, label)
        _style_table_cell(label_cell)
        for col_idx, vendor in enumerate(vendor_names, start=2):
            value = vendor_summary.get(vendor, {}).get(key, 0.0) or 0.0
            value_cell = ws.cell(row_idx, col_idx, value)
            _style_table_cell(value_cell)
            value_cell.number_format = num_fmt

    price_start = header_row + len(cost_metrics) + 3
    ws.cell(price_start, 1, "Pricing")
    _style_table_cell(ws.cell(price_start, 1), header=True, fill=SECTION_FILL)
    for col_idx, vendor in enumerate(vendor_names, start=2):
        _style_table_cell(ws.cell(price_start + 1, col_idx, vendor), header=True)
    _style_table_cell(ws.cell(price_start + 1, 1, "Item"), header=True)

    row_cursor = price_start + 2
    for row in optimizer_rows:
        sku = normalize_item_key(row.get("sku"))
        desc = str(row.get("description", "")).strip()
        item = item_costs.get(sku, {})
        winning_vendor = item.get("vendor", "")
        label = sku if not desc else f"{sku}\n{desc}"
        label_cell = ws.cell(row_cursor, 1, label)
        _style_table_cell(label_cell)
        for col_idx, vendor in enumerate(vendor_names, start=2):
            price = (row.get("prices") or {}).get(vendor)
            cell = ws.cell(row_cursor, col_idx, price if price is not None else "")
            fill = WINNER_FILL if vendor == winning_vendor and price is not None else None
            _style_table_cell(cell, fill=fill)
            if price is not None:
                cell.number_format = "$#,##0.00"
        row_cursor += 1

    qty_start = row_cursor + 2
    ws.cell(qty_start, 1, "Reorder Quantity")
    _style_table_cell(ws.cell(qty_start, 1), header=True, fill=SECTION_FILL)
    for col_idx, vendor in enumerate(vendor_names, start=2):
        _style_table_cell(ws.cell(qty_start + 1, col_idx, vendor), header=True)
    _style_table_cell(ws.cell(qty_start + 1, 1, "Item"), header=True)

    row_cursor = qty_start + 2
    for row in optimizer_rows:
        sku = normalize_item_key(row.get("sku"))
        desc = str(row.get("description", "")).strip()
        item = item_costs.get(sku, {})
        winning_vendor = item.get("vendor", "")
        label = sku if not desc else f"{sku}\n{desc}"
        label_cell = ws.cell(row_cursor, 1, label)
        _style_table_cell(label_cell)

        for col_idx, vendor in enumerate(vendor_names, start=2):
            if vendor == winning_vendor and item:
                qty = float(item.get("qty", 0.0) or 0.0)
                added = float(item.get("added_qty", 0.0) or 0.0)
                truck_labels = item.get("truck_labels", []) or []
                parts = [f"{qty:,.0f} SF"]
                sub_parts = []
                if added > 0:
                    sub_parts.append(f"+{added:,.0f} SF")
                if truck_labels:
                    sub_parts.append(" / ".join(f"T{truck}" for truck in truck_labels))
                value = parts[0] if not sub_parts else parts[0] + "\n" + " | ".join(sub_parts)
                fill = PatternFill(
                    fill_type="solid",
                    fgColor=TRUCK_COLOR_PALETTE[(int(item.get("primary_truck", 1) or 1) - 1) % len(TRUCK_COLOR_PALETTE)],
                )
                _style_table_cell(ws.cell(row_cursor, col_idx, value), fill=fill)
            else:
                _style_table_cell(ws.cell(row_cursor, col_idx, ""))
        row_cursor += 1

    footer_rows = [
        ("Vendor Total SF", "qty_sf", "#,##0"),
        ("Added SF", "added_qty_sf", "#,##0"),
        ("Landed Total", "landed_total", "$#,##0.00"),
    ]
    for label, key, num_fmt in footer_rows:
        label_cell = ws.cell(row_cursor, 1, label)
        _style_table_cell(label_cell, header=True)
        for col_idx, vendor in enumerate(vendor_names, start=2):
            value = vendor_summary.get(vendor, {}).get(key, 0.0) or 0.0
            cell = ws.cell(row_cursor, col_idx, value if value else "")
            _style_table_cell(cell, header=True)
            if value:
                cell.number_format = num_fmt
        row_cursor += 1

    ws.freeze_panes = "B5"
    ws.column_dimensions["A"].width = 42
    for col_letter in "BCDEFGHIJKLMNOP":
        ws.column_dimensions[col_letter].width = 16


def _write_vendor_summary_sheet(ws, optimization_result: Dict[str, Any]) -> None:
    ws.title = "Vendor Summary"
    vendor_summary = optimization_result.get("vendor_summary", {}) or {}
    columns = [
        "Vendor",
        "Qty SF",
        "Truckloads",
        "Truck Fill %",
        "Freight / Truck",
        "Freight Total",
        "Material Total",
        "Landed Total",
        "Avg Landed Cost / SF",
        "Added SF",
    ]
    for col_idx, label in enumerate(columns, start=1):
        _style_table_cell(ws.cell(1, col_idx, label), header=True)

    for row_idx, vendor in enumerate(sorted(vendor_summary), start=2):
        summary = vendor_summary[vendor]
        values = [
            vendor,
            summary.get("qty_sf", 0.0),
            summary.get("truckloads", 0),
            summary.get("truck_fill_pct", 0.0),
            summary.get("freight_per_truck", 0.0),
            summary.get("freight_total", 0.0),
            summary.get("material_total", 0.0),
            summary.get("landed_total", 0.0),
            summary.get("avg_landed_cost", 0.0),
            summary.get("added_qty_sf", 0.0),
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row_idx, col_idx, value)
            _style_table_cell(cell)
            if col_idx in (2, 10):
                cell.number_format = "#,##0.00"
            elif col_idx == 3:
                cell.number_format = "0"
            elif col_idx == 4:
                cell.number_format = "0.0%"
            elif col_idx >= 5:
                cell.number_format = "$#,##0.00" if col_idx != 10 else "#,##0.00"

    for col_letter, width in zip("ABCDEFGHIJ", [18, 14, 11, 12, 14, 14, 14, 14, 18, 12]):
        ws.column_dimensions[col_letter].width = width


def _write_item_detail_sheet(ws, optimization_result: Dict[str, Any]) -> None:
    ws.title = "Item Detail"
    columns = [
        "SKU",
        "Description",
        "Winning Vendor",
        "Base Qty SF",
        "Added Qty SF",
        "Final Qty SF",
        "Pallet SF",
        "Material Price / SF",
        "Freight Share / SF",
        "Landed Cost / SF",
        "Line Total",
        "Truck Labels",
    ]
    for col_idx, label in enumerate(columns, start=1):
        _style_table_cell(ws.cell(1, col_idx, label), header=True)

    item_costs = optimization_result.get("item_costs", {}) or {}
    for row_idx, sku in enumerate(sorted(item_costs), start=2):
        item = item_costs[sku]
        values = [
            sku,
            item.get("description", ""),
            item.get("vendor", ""),
            item.get("base_qty", 0.0),
            item.get("added_qty", 0.0),
            item.get("qty", 0.0),
            item.get("pallet_sf", 0.0),
            item.get("price", 0.0),
            item.get("freight_share", 0.0),
            item.get("landed_cost", 0.0),
            item.get("line_total", 0.0),
            " / ".join(f"T{truck}" for truck in (item.get("truck_labels", []) or [])),
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row_idx, col_idx, value)
            _style_table_cell(cell)
            if col_idx in (4, 5, 6, 7):
                cell.number_format = "#,##0.00"
            elif col_idx in (8, 9, 10, 11):
                cell.number_format = "$#,##0.0000" if col_idx in (8, 9, 10) else "$#,##0.00"

    widths = [28, 36, 18, 12, 12, 12, 10, 16, 16, 16, 14, 12]
    for col_letter, width in zip("ABCDEFGHIJKL", widths):
        ws.column_dimensions[col_letter].width = width


def _write_inputs_sheet(
    ws,
    vendor_names: List[str],
    optimizer_rows: List[Dict[str, Any]],
    pricing_inputs: Dict[str, Any],
) -> None:
    ws.title = "Inputs"
    headers = ["SKU", "Description", "Qty Needed SF", "SF per Pallet"] + vendor_names
    for col_idx, label in enumerate(headers, start=1):
        _style_table_cell(ws.cell(1, col_idx, label), header=True)

    for row_idx, row in enumerate(optimizer_rows, start=2):
        values = [
            row.get("sku", ""),
            row.get("description", ""),
            row.get("quantity", 0.0),
            row.get("sf_per_pallet", 0.0),
        ] + [(row.get("prices") or {}).get(vendor, "") for vendor in vendor_names]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row_idx, col_idx, value)
            _style_table_cell(cell)
            if col_idx in (3, 4):
                cell.number_format = "#,##0.00"
            elif col_idx >= 5 and value != "":
                cell.number_format = "$#,##0.00"

    freight_start = len(optimizer_rows) + 4
    _style_table_cell(ws.cell(freight_start, 1, "Freight / Truck"), header=True, fill=SECTION_FILL)
    for col_idx, vendor in enumerate(vendor_names, start=2):
        _style_table_cell(ws.cell(freight_start, col_idx, vendor), header=True)
        value = pricing_inputs.get("freight_costs", {}).get(vendor, 0.0) or 0.0
        cell = ws.cell(freight_start + 1, col_idx, value if value else "")
        _style_table_cell(cell)
        if value:
            cell.number_format = "$#,##0.00"
    ws.column_dimensions["A"].width = 34
    for col_letter in "BCDEFGHIJKLMNOP":
        ws.column_dimensions[col_letter].width = 15


def _write_missing_quotes_sheet(ws, missing_quotes: List[str]) -> None:
    ws.title = "Missing Quotes"
    _style_table_cell(ws.cell(1, 1, "SKU"), header=True)
    if not missing_quotes:
        _style_table_cell(ws.cell(2, 1, "No missing quotes."), fill=None)
    else:
        for row_idx, sku in enumerate(missing_quotes, start=2):
            _style_table_cell(ws.cell(row_idx, 1, sku))
    ws.column_dimensions["A"].width = 40


def build_snapshot_workbook(
    payload_path: Path,
    pricing_path: Path,
    strip_path: Path,
    output_path: Path,
) -> Dict[str, Any]:
    payload = json.loads(payload_path.read_text(encoding="utf-8"))
    strip_skus, strip_details = _load_strip_sku_details(strip_path)
    groups = _build_sku_consolidation_groups(strip_path)
    group_descriptions = _build_group_descriptions(groups, strip_details)

    pricing_inputs = load_niko_pricing_inputs(
        pricing_path,
        default_vendor_names=DEFAULT_STRIP_VENDOR_COLUMNS,
        vendor_aliases=DEFAULT_NIKO_VENDOR_ALIASES,
    )
    run_month, optimizer_rows, warnings_out = _build_optimizer_rows(
        payload,
        strip_skus,
        strip_details,
        groups,
        group_descriptions,
        pricing_inputs,
    )

    vendor_names = pricing_inputs.get("vendor_names") or list(DEFAULT_STRIP_VENDOR_COLUMNS)
    optimization_result, missing_quotes, omitted_items = _run_optimizer(
        optimizer_rows,
        vendor_names,
        pricing_inputs.get("freight_costs", {}) or {},
    )

    wb = Workbook()
    _write_summary_sheet(
        wb.active,
        payload.get("run_meta", {}) or {},
        run_month,
        pricing_path,
        payload_path,
        pricing_inputs,
        optimizer_rows,
        optimization_result,
        missing_quotes,
        warnings_out,
    )
    _write_optimizer_card_sheet(wb.create_sheet(), run_month, vendor_names, optimizer_rows, optimization_result)
    _write_vendor_summary_sheet(wb.create_sheet(), optimization_result)
    _write_item_detail_sheet(wb.create_sheet(), optimization_result)
    _write_inputs_sheet(wb.create_sheet(), vendor_names, optimizer_rows, pricing_inputs)
    _write_missing_quotes_sheet(wb.create_sheet(), missing_quotes)
    wb.save(output_path)

    return {
        "output_path": str(output_path),
        "run_month": run_month.strftime("%Y-%m-%d"),
        "quoted_items": sum(1 for row in optimizer_rows if row.get("prices")),
        "total_items": len(optimizer_rows),
        "missing_quotes": missing_quotes,
        "omitted_items": omitted_items,
        "total_cost": float(optimization_result.get("total_cost", 0.0) or 0.0),
        "total_material_cost": float(optimization_result.get("total_material_cost", 0.0) or 0.0),
        "total_freight_cost": float(optimization_result.get("total_freight_cost", 0.0) or 0.0),
        "total_added_qty": float(optimization_result.get("total_added_qty", 0.0) or 0.0),
        "vendor_summary": optimization_result.get("vendor_summary", {}) or {},
        "warnings": (pricing_inputs.get("warnings", []) or []) + warnings_out + (optimization_result.get("errors", []) or []),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Export a standalone Niko optimizer workbook from the saved dashboard payload.")
    parser.add_argument("--payload", type=Path, default=DEFAULT_PAYLOAD_PATH)
    parser.add_argument("--pricing", type=Path, default=DEFAULT_PRICING_PATH)
    parser.add_argument("--strip", type=Path, default=DEFAULT_STRIP_PATH)
    parser.add_argument("--output", type=Path, default=None)
    args = parser.parse_args()

    if args.output is None:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        args.output = REPO_DIR / f"niko_optimizer_snapshot_{stamp}.xlsx"

    result = build_snapshot_workbook(
        payload_path=args.payload,
        pricing_path=args.pricing,
        strip_path=args.strip,
        output_path=args.output,
    )
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
