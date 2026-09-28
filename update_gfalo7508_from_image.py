from __future__ import annotations

from copy import copy
from datetime import datetime
from pathlib import Path
import math
import shutil

import numpy as np
import openpyxl
from openpyxl.styles import PatternFill
from PIL import Image


WORKBOOK = Path(r"H:\2025\MISC Reports\For Marketing\Garrison Color Analysis 2.12.2026.xlsx")
IMAGE_PATH = Path(r"H:\2025\MISC Reports\For Marketing\Overheads for Color Analysis\European Oak Volto 7.5.webp")
BACKUP = WORKBOOK.with_name(
    f"{WORKBOOK.stem} backup before exact GFALO7508 pixels {datetime.now():%Y%m%d_%H%M%S}{WORKBOOK.suffix}"
)


def header_map(ws):
    return {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}


def copy_cell_style(src, dst):
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)


def rgb_fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=f"FF{hex_color.lstrip('#').upper()}")


def srgb_to_lab(rgb255: np.ndarray) -> np.ndarray:
    rgb = rgb255.astype(np.float64) / 255.0
    linear = np.where(rgb <= 0.04045, rgb / 12.92, ((rgb + 0.055) / 1.055) ** 2.4)
    matrix = np.array(
        [
            [0.4124564, 0.3575761, 0.1804375],
            [0.2126729, 0.7151522, 0.0721750],
            [0.0193339, 0.1191920, 0.9503041],
        ]
    )
    xyz = linear @ matrix.T
    white = np.array([0.95047, 1.00000, 1.08883])
    xyz = xyz / white
    delta = 6 / 29
    f = np.where(xyz > delta**3, np.cbrt(xyz), xyz / (3 * delta**2) + 4 / 29)
    lab = np.empty_like(f)
    lab[:, 0] = 116 * f[:, 1] - 16
    lab[:, 1] = 500 * (f[:, 0] - f[:, 1])
    lab[:, 2] = 200 * (f[:, 1] - f[:, 2])
    return lab


def rgb_to_hsv_stats(rgb255: np.ndarray) -> tuple[float, float]:
    rgb = rgb255.astype(np.float64) / 255.0
    maxc = rgb.max(axis=1)
    minc = rgb.min(axis=1)
    delta = maxc - minc
    saturation = np.where(maxc == 0, 0, delta / maxc)
    return float(saturation.mean()), float(maxc.mean())


def edge_density(gray: np.ndarray) -> float:
    # Lightweight Sobel-style gradient fraction. The old workbook does not
    # store the source algorithm, so this gives a reproducible texture measure.
    g = gray.astype(np.float64)
    dx = np.zeros_like(g)
    dy = np.zeros_like(g)
    dx[:, 1:-1] = g[:, 2:] - g[:, :-2]
    dy[1:-1, :] = g[2:, :] - g[:-2, :]
    mag = np.hypot(dx, dy)
    return float((mag > 30).mean())


def image_metrics(path: Path) -> dict[str, float | str]:
    img = Image.open(path).convert("RGB")
    arr = np.asarray(img)
    rgb = arr.reshape(-1, 3)
    mean_rgb = rgb.mean(axis=0)
    median_rgb = np.median(rgb, axis=0)
    lab = srgb_to_lab(rgb)
    mean_lab = lab.mean(axis=0)
    hue = (math.degrees(math.atan2(float(mean_lab[2]), float(mean_lab[1]))) + 360) % 360
    chroma = math.hypot(float(mean_lab[1]), float(mean_lab[2]))
    sat, val = rgb_to_hsv_stats(rgb)
    gray = np.dot(arr[..., :3], [0.299, 0.587, 0.114])
    contrast = float(srgb_to_lab(rgb)[:, 0].std())
    warm_score = float((mean_rgb[0] - mean_rgb[2]) / 255.0)

    return {
        "MeanHex": "#{:02X}{:02X}{:02X}".format(*np.rint(mean_rgb).astype(int)),
        "MedianHex": "#{:02X}{:02X}{:02X}".format(*np.rint(median_rgb).astype(int)),
        "MeanRGB_R": float(mean_rgb[0]),
        "MeanRGB_G": float(mean_rgb[1]),
        "MeanRGB_B": float(mean_rgb[2]),
        "L_mean": float(mean_lab[0]),
        "a_mean": float(mean_lab[1]),
        "b_mean": float(mean_lab[2]),
        "HueDeg": hue,
        "Chroma": chroma,
        "Sat_mean": sat,
        "Value_mean": val,
        "Contrast_L_std": contrast,
        "EdgeDensity": edge_density(gray),
        "WarmthScore": warm_score,
        "ImageWidthPx": img.width,
        "ImageHeightPx": img.height,
    }


def bucket_for_metrics(metrics: dict[str, float | str]) -> tuple[int, str, str]:
    hue = float(metrics["HueDeg"])
    light = float(metrics["L_mean"])
    chroma = float(metrics["Chroma"])
    if light >= 82 and chroma < 4:
        return 2, "Neutral Whitewash", "Cool/Neutral"
    if light >= 80 and hue >= 75:
        return 3, "Honey Blonde", "Golden/Honey"
    if light >= 70:
        return 4, "Warm Blonde", "Golden/Honey"
    if light >= 60:
        return 5, "Golden Oak", "Golden/Honey"
    if light >= 50:
        return 6, "Amber", "Amber/Caramel"
    if light >= 40:
        return 7, "Caramel", "Amber/Caramel"
    if light >= 30:
        return 8, "Walnut Brown", "Brown"
    if chroma < 16:
        return 9, "Deep Espresso", "Brown"
    return 10, "Smoky Espresso", "Brown"


def main():
    metrics = image_metrics(IMAGE_PATH)
    bucket, bucket_name, hue_family = bucket_for_metrics(metrics)

    shutil.copy2(WORKBOOK, BACKUP)
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb["Products"]
    headers = header_map(ws)

    row_idx = None
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, headers["ProductNumber"]).value or "").strip().upper() == "GFALO7508":
            row_idx = r
            break
    if row_idx is None:
        raise RuntimeError("GFALO7508 row not found. Add the product before exact pixel refresh.")

    updates = {
        "HueBucket": bucket,
        "BucketName": bucket_name,
        "HueFamily": hue_family,
        "ImageAnalysisSource": str(IMAGE_PATH),
        "Notes": 'Gartman item description: EURO OAK 7.5" VOLTO. Metrics computed from source WebP file.',
    }
    updates.update(metrics)

    for header, value in updates.items():
        if header in headers:
            ws.cell(row_idx, headers[header]).value = value
    ws.cell(row_idx, headers["Swatch"]).fill = rgb_fill(str(metrics["MeanHex"]))

    wb.save(WORKBOOK)

    print(f"Updated: {WORKBOOK}")
    print(f"Backup:  {BACKUP}")
    print(f"Image:   {IMAGE_PATH}")
    print(f"Row:     {row_idx}")
    print(f"Bucket:  {bucket} - {bucket_name}")
    for key in [
        "MeanHex",
        "MedianHex",
        "MeanRGB_R",
        "MeanRGB_G",
        "MeanRGB_B",
        "L_mean",
        "a_mean",
        "b_mean",
        "HueDeg",
        "Chroma",
        "Sat_mean",
        "Value_mean",
        "Contrast_L_std",
        "EdgeDensity",
        "WarmthScore",
        "ImageWidthPx",
        "ImageHeightPx",
    ]:
        print(f"{key}: {metrics[key]}")


if __name__ == "__main__":
    main()
