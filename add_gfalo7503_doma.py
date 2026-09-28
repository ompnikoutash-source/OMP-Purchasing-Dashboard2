from __future__ import annotations

from datetime import datetime
from pathlib import Path
import shutil
import sys

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from add_gfalo7506s_strada_select import ensure_columns, header_map, set_row
from refresh_garrison_color_rollups import (
    refresh_buckets,
    refresh_summary,
    set_gfalo7508_bucket_and_dimensions,
)
from update_gfalo7508_from_image import image_metrics


WORKBOOK = Path(r"H:\2025\MISC Reports\For Marketing\Garrison Color Analysis 2.12.2026.xlsx")
IMAGE_PATH = Path(r"H:\2025\MISC Reports\For Marketing\Overheads for Color Analysis\European Oak Doma 7.5.jpg")
PENDING_WORKBOOK = WORKBOOK.with_name(f"{WORKBOOK.stem} pending GFALO7503 update{WORKBOOK.suffix}")
BACKUP = WORKBOOK.with_name(
    f"{WORKBOOK.stem} backup before GFALO7503 {datetime.now():%Y%m%d_%H%M%S}{WORKBOOK.suffix}"
)


def main():
    metrics = image_metrics(IMAGE_PATH)
    shutil.copy2(WORKBOOK, BACKUP)
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb["Products"]
    ensure_columns(ws)
    headers = header_map(ws)

    product_number = "GFALO7503"
    row_idx = None
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, headers["ProductNumber"]).value or "").strip().upper() == product_number:
            row_idx = row
            break
    if row_idx is None:
        row_idx = ws.max_row + 1
        next_index = max(cell.value for cell in ws["A"][1:] if isinstance(cell.value, (int, float))) + 1
    else:
        next_index = ws.cell(row_idx, headers["INDEX"]).value

    values = {
        "INDEX": next_index,
        "Collection": "Allora",
        "Product": "Doma",
        "Width_in": 7.5,
        "HueBucket": 5,
        "BucketName": "Golden Oak",
        "HueFamily": "Golden/Honey",
        "Undertone": "Warm",
        "Notes": 'Gartman item description: EURO OAK 7.5" DOMA. Metrics computed from source JPG file.',
        "ProductNumber": product_number,
        "Species": "European Oak",
        "Construction": "Engineered (Birch Plywood)",
        "Grade": "Select Character",
        "Finish": "Matte UV Lacquer",
        "SurfaceTreatment": "Light Wire-Brush",
        "JankaRating": 1360,
        "Thickness": '5/8"',
        "TopLayer": "4mm",
        "Length": "Mostly 7'8 1/2\"",
        "InstallationMethods": "Glue Down, Floating",
        "EdgeDetail": "Micro-Bevel",
        "Warranty": "Residential: Limited Lifetime Structure, Limited 25 Year Finish | Commercial: Limited 25 Year Structure, Limited 10 Year Finish",
        "ImageAnalysisSource": str(IMAGE_PATH),
    }
    values.update(metrics)

    set_row(ws, row_idx, values)
    set_gfalo7508_bucket_and_dimensions(wb)
    refresh_buckets(wb)
    refresh_summary(wb)

    save_target = WORKBOOK
    try:
        wb.save(save_target)
    except PermissionError:
        save_target = PENDING_WORKBOOK
        wb.save(save_target)

    print(f"Updated: {save_target}")
    if save_target != WORKBOOK:
        print(f"Original locked: {WORKBOOK}")
    print(f"Backup:  {BACKUP}")
    print(f"Row:     {row_idx}")
    print("Bucket:  5 - Golden Oak")
    for key in [
        "MeanHex",
        "MedianHex",
        "L_mean",
        "HueDeg",
        "Chroma",
        "Sat_mean",
        "Value_mean",
        "Contrast_L_std",
        "EdgeDensity",
        "WarmthScore",
        "ImageWidthPx",
        "ImageHeightPx",
    ]:
        print(f"{key}: {metrics[key]}")


if __name__ == "__main__":
    main()
