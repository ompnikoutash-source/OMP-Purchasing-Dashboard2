import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from gartman_connection import get_connection


BASE_DIR = Path(__file__).resolve().parent
SQL_PATH = BASE_DIR / "nonstock_sample_orders_since_2024_source_query.sql"
DEFAULT_START_DATE = date(2024, 1, 1)

NOTE_LINE_CSV = BASE_DIR / "nonstock_sample_orders_since_2024_catalog_note_lines.csv"
PARSED_CSV = BASE_DIR / "nonstock_sample_orders_since_2024_catalog_parsed.csv"
REVIEW_CSV = BASE_DIR / "nonstock_sample_orders_since_2024_catalog_review.csv"
WORKBOOK_PATH = BASE_DIR / "nonstock_sample_orders_since_2024_catalog_review.xlsx"
RUN_SUMMARY_PATH = BASE_DIR / "nonstock_sample_orders_since_2024_catalog_run_summary.txt"
PRIOR_SUMMARY_PATH = BASE_DIR / "final_order_notes_results_summary.csv"
CATALOG_JSON_PATH = BASE_DIR / "catalog_intelligence" / "data" / "en_dashboard_data.json"
CATALOG_ALIAS_CSV = BASE_DIR / "nonstock_sample_orders_catalog_aliases.csv"


SUMMARY_HEADERS = [
    "ORDER_NUMBER",
    "INVOICE_NUMBER",
    "ORDER_DATE",
    "ACCOUNT_NUMBER",
    "SAMPLE_TYPE",
    "SAMPLE_TYPE_SOURCE",
    "SAMPLE_QTY",
    "SAMPLE_QTY_SOURCE",
    "SAMPLE_QTY_RULE",
    "CATALOG_EACH_ITEM_COUNT",
    "CATALOG_EACH_ALIASES_FOUND",
    "CATALOG_EACH_ITEMS_FOUND",
    "PRIOR_SUMMARY_TYPE",
    "PRIOR_SUMMARY_QTY",
    "PRIOR_SUMMARY_WORKBOOK_ROW",
    "FREIGHT_DOLLAR_AMOUNT",
    "FREIGHT_AMOUNT_SOURCE",
    "CUSTOMER_NO_CHARGE_NOTE",
    "DOLLAR_AMOUNTS_FOUND",
    "REVIEW_NEEDED",
    "REVIEW_FLAGS",
    "MATCH_TERMS_FOUND",
    "GARTMAN_ORDER_DATE",
    "INVOICE_DATE",
    "COMPANY_NUMBER",
    "LOCATION_NUMBER",
    "RELEASE_NUMBER",
    "ORDER_INVOICE_KEY",
    "NONSTOCK_LINE_COUNT",
    "NONSTOCK_QTY_UNITS",
    "NONSTOCK_NET_AMOUNT",
    "NOTE_LINE_COUNT",
    "FULL_ORDER_NOTES",
]

WORD_NUMBERS = {
    "A": 1,
    "AN": 1,
    "ONE": 1,
    "TWO": 2,
    "THREE": 3,
    "FOUR": 4,
    "FIVE": 5,
    "SIX": 6,
    "SEVEN": 7,
    "EIGHT": 8,
    "NINE": 9,
    "TEN": 10,
    "ELEVEN": 11,
    "TWELVE": 12,
    "THIRTEEN": 13,
    "FOURTEEN": 14,
    "FIFTEEN": 15,
    "SIXTEEN": 16,
    "SEVENTEEN": 17,
    "EIGHTEEN": 18,
    "NINETEEN": 19,
    "TWENTY": 20,
}

MATCH_TERM_PATTERNS = {
    "pcs": re.compile(r"\bP\.?\s*C\.?\s*S\.?\b"),
    "pieces": re.compile(r"\bPIECES?\b"),
    "sample": re.compile(r"\bSAMPLES?\b"),
    "cut": re.compile(r"\bCUTS?\b"),
    "panel": re.compile(r"\bPANELS?\b"),
    "plank": re.compile(r"\bPLANKS?\b"),
}

UNIT_WORDS = (
    "SAMPLE",
    "SAMPLES",
    "CHIP",
    "CHIPS",
    "PCS",
    "PC",
    "PIECE",
    "PIECES",
    "CUT",
    "CUTS",
    "PANEL",
    "PANELS",
    "PLANK",
    "PLANKS",
    "BOARD",
    "BOARDS",
    "SET",
    "SETS",
)

CATALOG_STOP_TOKENS = {
    "A",
    "AL",
    "ALUMINUM",
    "AMERICAN",
    "BABY",
    "BACKBOARD",
    "BACKBOARDS",
    "BASESHOE",
    "BEV",
    "BF",
    "BIRCH",
    "BOARD",
    "BOARDS",
    "BY",
    "CALC",
    "CHAR",
    "CHARACTER",
    "CHIP",
    "CHIPS",
    "CUT",
    "CUTS",
    "DIST",
    "DISTRESS",
    "DISTRESSED",
    "ENG",
    "ENGINEERED",
    "EUR",
    "EURO",
    "EUROPEAN",
    "FLOORING",
    "GARRISON",
    "HB",
    "HERRINGBONE",
    "HICKORY",
    "HOLD",
    "IN",
    "L",
    "LACQUER",
    "LABEL",
    "MAPLE",
    "MM",
    "MOULD",
    "MOULDING",
    "MOULDINGS",
    "NOSING",
    "NOSINGS",
    "OAK",
    "OIL",
    "OXIDE",
    "PANEL",
    "PANELS",
    "PART",
    "PC",
    "PCS",
    "PLANK",
    "PLANKS",
    "R",
    "RED",
    "REDUCER",
    "REDUCERS",
    "SAMPLE",
    "SAMPLES",
    "SF",
    "SMOOTH",
    "SOLID",
    "SQ",
    "SQUARE",
    "STAIN",
    "T",
    "THRESHOLD",
    "TRESHOLD",
    "TMOLD",
    "TMOULD",
    "TMOULDING",
    "USA",
    "UV",
    "WATERFALL",
    "WHITE",
    "WOOD",
    "X",
}

CATALOG_STOP_PHRASES = {
    "ALLORA",
    "ALLORA BY GARRISON",
    "BEVERLY HILLS",
    "CAROLINA CLASSIC",
    "COMPETITION BUSTER",
    "CONTRACTORS CHOICE",
    "CRYSTAL VALLEY",
    "CRYSTAL VALLEY USA",
    "EXOTICS",
    "GII DISTRESSED",
    "GII SMOOTH",
    "GOLD LABEL",
    "LEGENDS",
    "NEWPORT",
    "PRIVATE SELECTION",
    "SHEOGA",
    "VINEYARD",
}


def parse_args():
    parser = argparse.ArgumentParser(
        description="Pull and parse NONSTOCK sample-note orders by order/invoice."
    )
    parser.add_argument("--start-date", default=DEFAULT_START_DATE.isoformat())
    parser.add_argument("--end-date", default=date.today().isoformat())
    parser.add_argument(
        "--chunk-days",
        type=int,
        default=31,
        help="Date-window size for the ODBC query. Smaller chunks avoid SQL0101.",
    )
    parser.add_argument(
        "--reuse-note-lines",
        default="",
        help="Optional existing raw note-line CSV to parse instead of querying Gartman.",
    )
    parser.add_argument("--no-xlsx", action="store_true")
    return parser.parse_args()


def parse_date(value):
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise SystemExit(f"Invalid date {value!r}; expected YYYY-MM-DD") from exc


def date_windows(start_date, end_date, chunk_days):
    current = start_date
    while current <= end_date:
        chunk_end = min(end_date, current + timedelta(days=chunk_days - 1))
        yield current, chunk_end
        current = chunk_end + timedelta(days=1)


def clean_cell(value):
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format_decimal(value)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value).strip()


def format_decimal(value):
    if value is None or value == "":
        return ""
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return str(value)
    return f"{number.quantize(Decimal('0.01'))}"


def normalize_numeric_label(value):
    value = str(value or "").strip()
    if not value:
        return ""
    value = value.replace(",", "")
    try:
        number = Decimal(value)
    except InvalidOperation:
        return value
    if number == number.to_integral_value():
        return str(int(number))
    return str(number.normalize())


def normalize_note(text):
    text = (text or "").upper()
    text = text.replace("\u00a0", " ")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    text = text.replace("\u2018", "'").replace("\u2019", "'")
    replacements = {
        "SMAPLE": "SAMPLE",
        "SAMLPLE": "SAMPLE",
        "SMPLE": "SAMPLE",
        "SAMPE": "SAMPLE",
        "SHEEL": "SHELL",
        "SAMPLES CHIPS": "SAMPLE CHIPS",
        "CHIPES": "CHIPS",
        "PEICES": "PIECES",
    }
    for bad, good in replacements.items():
        text = re.sub(rf"\b{bad}\b", good, text)
    text = re.sub(r"\bP\s*/\s*C\s*/\s*S\b", "PCS", text)
    text = re.sub(r"\bP\s+C\s+S\b", "PCS", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_numbers(text):
    unit_pattern = "|".join(UNIT_WORDS)

    def repl(match):
        return str(WORD_NUMBERS[match.group(1)])

    return re.sub(
        rf"\b({'|'.join(WORD_NUMBERS)})\b(?=\s+(?:{unit_pattern})\b)",
        repl,
        text,
    )


def match_terms(normalized):
    return [
        term for term, pattern in MATCH_TERM_PATTERNS.items() if pattern.search(normalized)
    ]


def detect_type(normalized):
    hits = []
    has_set = bool(
        re.search(
            r"\bSAMPLE\s+SETS?\b|\bSETS?\s+OF\s+SAMPLES?\b|\bSAMPLES?\s+SETS?\b",
            normalized,
        )
    )
    has_panel = bool(
        re.search(
            r"\b(?:SAMPLE\s+)?(?:PANELS?|PLANKS?|BOARDS?)\b|"
            r"\b(?:PANELS?|PLANKS?|BOARDS?)\s+SAMPLES?\b",
            normalized,
        )
    )
    has_chip_explicit = bool(
        re.search(
            r"\b(?:SAMPLE\s+)?CHIPS?\b|\bP\.?\s*C\.?\s*S\.?\b|"
            r"\bPIECES?\b|\bCUTS?\b",
            normalized,
        )
    )
    has_generic_sample = bool(re.search(r"\bSAMPLES?\b", normalized))

    if has_set:
        hits.append("sample set")
    if has_panel:
        hits.append("sample panel")
    if has_chip_explicit or (has_generic_sample and not has_set and not has_panel):
        hits.append("sample chip")

    flags = []
    if not hits:
        flags.append("NO_SAMPLE_TYPE_DETECTED")
    elif len(hits) > 1:
        flags.append("MULTIPLE_SAMPLE_TYPES")

    return "; ".join(hits), flags


def normalize_prior_type(value):
    text = normalize_note(value)
    if not text:
        return ""
    hits = []
    if "SET" in text:
        hits.append("sample set")
    if "PANEL" in text or "PLANK" in text or "BOARD" in text:
        hits.append("sample panel")
    if "CHIP" in text or "PCS" in text or "PIECE" in text or "CUT" in text:
        hits.append("sample chip")
    return "; ".join(dict.fromkeys(hits))


def width_value(value):
    text = normalize_note(value)
    match = re.search(r"\b\d+(?:\.\d+)?\b", text)
    return match.group(0) if match else ""


def catalog_tokens(value):
    text = normalize_note(value)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\([^)]*\)", " ", text)
    text = text.replace("&", " ")
    text = re.sub(r"[^A-Z0-9.]+", " ", text)
    return re.findall(r"[A-Z]+|\d+(?:\.\d+)?", text)


def alias_is_allowed(alias):
    alias = re.sub(r"\s+", " ", alias or "").strip()
    if not alias or alias in CATALOG_STOP_PHRASES:
        return False
    tokens = alias.split()
    if not tokens or not any(re.search(r"[A-Z]", token) for token in tokens):
        return False
    if all(token in CATALOG_STOP_TOKENS or re.fullmatch(r"\d+(?:\.\d+)?", token) for token in tokens):
        return False
    if len(tokens) == 1 and (len(tokens[0]) < 4 or tokens[0] in CATALOG_STOP_TOKENS):
        return False
    return True


def alias_candidates_from_catalog_text(value, width=""):
    tokens = catalog_tokens(value)
    filtered = []
    for token in tokens:
        if token in CATALOG_STOP_TOKENS:
            continue
        if re.fullmatch(r"\d+(?:\.\d+)?", token):
            continue
        filtered.append(token)

    aliases = []
    if filtered:
        aliases.append(" ".join(filtered))
        if len(filtered) >= 2:
            aliases.append(" ".join(filtered[-2:]))
        if len(filtered) >= 3:
            aliases.append(" ".join(filtered[-3:]))
        if filtered[-1] == "SELECT" and len(filtered) >= 2:
            aliases.append(" ".join(filtered[:-1]))

    clean_width = width_value(width)
    if clean_width:
        width_aliases = []
        for alias in aliases:
            width_aliases.append(f"{alias} {clean_width}")
            width_aliases.append(f"{clean_width} {alias}")
        aliases.extend(width_aliases)

    return [alias for alias in dict.fromkeys(aliases) if alias_is_allowed(alias)]


def catalog_display(record):
    item = str(record.get("item") or record.get("ITEM") or "").strip()
    collection = str(record.get("collection") or record.get("COLLECTION") or "").strip()
    description = str(record.get("description") or record.get("DESCRIPTION") or "").strip()
    parts = []
    if collection:
        parts.append(collection)
    if description:
        parts.append(description)
    if item:
        parts.append(f"[{item}]")
    return " - ".join(parts)


def local_catalog_records():
    if not CATALOG_JSON_PATH.exists():
        return []

    data = json.loads(CATALOG_JSON_PATH.read_text(encoding="utf-8"))
    records_by_item = {}
    for section in ("catalogItems", "itemActions"):
        for record in data.get(section, []) or []:
            item = str(record.get("item") or record.get("itemKey") or "").strip()
            if item and item not in records_by_item:
                records_by_item[item] = record
    return list(records_by_item.values())


def itemmast_catalog_records():
    records = []
    try:
        conn = get_connection()
        try:
            cursor = conn.cursor()
            cursor.execute(
                """
                SELECT
                    TRIM(IMITEM) AS ITEM,
                    TRIM(IMDESC) AS DESCRIPTION,
                    TRIM(IMCOLR) AS COLLECTION,
                    IMDIV AS DIVISION
                FROM GSFL2K.ITEMMAST
                WHERE COALESCE(IMDELT, '') <> 'D'
                  AND IMDIV IN (1, 5, 7)
                  AND TRIM(COALESCE(IMDESC, '')) <> ''
                """
            )
            headers = [column[0].upper() for column in cursor.description]
            for row in cursor.fetchall():
                records.append(dict(zip(headers, row)))
        finally:
            conn.close()
    except Exception as exc:
        print(f"Warning: could not load ITEMMAST catalog aliases: {exc}")
    return records


def build_catalog_aliases():
    alias_map = {}
    for record in [*local_catalog_records(), *itemmast_catalog_records()]:
        fields = [
            record.get("productName"),
            record.get("description"),
            record.get("DESCRIPTION"),
        ]
        width = record.get("width") or ""
        display = catalog_display(record)
        if not display:
            continue

        for field in fields:
            if not field:
                continue
            for alias in alias_candidates_from_catalog_text(field, width):
                alias_map.setdefault(alias, {})
                alias_map[alias][display] = True

    aliases = []
    for alias, items in alias_map.items():
        pattern = r"(?<![A-Z0-9])" + r"\s+".join(
            re.escape(part) for part in alias.split()
        ) + r"(?![A-Z0-9])"
        aliases.append(
            {
                "alias": alias,
                "items": list(items.keys()),
                "tokens": set(alias.split()),
                "pattern": re.compile(pattern),
            }
        )
    aliases.sort(key=lambda entry: (-len(entry["alias"]), entry["alias"]))
    return aliases


def write_catalog_alias_csv(catalog_aliases):
    rows = [
        {
            "CATALOG_ALIAS": entry["alias"],
            "CATALOG_ITEMS": "; ".join(entry["items"]),
            "CATALOG_ITEM_COUNT": len(entry["items"]),
        }
        for entry in catalog_aliases
    ]
    write_csv(
        CATALOG_ALIAS_CSV,
        ["CATALOG_ALIAS", "CATALOG_ITEMS", "CATALOG_ITEM_COUNT"],
        rows,
    )


def find_catalog_each_matches(normalized, start_index, catalog_aliases):
    tail = normalized[start_index : start_index + 750]
    tail = re.sub(
        r"\b(?:UPS|FEDEX|FED EX|SHIP|SHIPPING|FREIGHT|DELIVERY|DLV|NO CHARGE|"
        r"PAID|VISA|AMEX|CHECK|CASH|CARD|TR#|TRACKING|QUOTE)\b.*$",
        "",
        tail,
    )
    tail_tokens = set(catalog_tokens(tail))
    matches = []
    occupied = []

    for entry in catalog_aliases:
        if not entry["tokens"].issubset(tail_tokens):
            continue
        alias = entry["alias"]
        for match in entry["pattern"].finditer(tail):
            start = start_index + match.start()
            end = start_index + match.end()
            if any(start < used_end and end > used_start for used_start, used_end in occupied):
                continue
            occupied.append((start, end))
            matches.append(
                {
                    "alias": alias,
                    "items": entry["items"],
                    "start": start,
                    "end": end,
                }
            )

    matches.sort(key=lambda item: item["start"])
    aliases_found = [match["alias"] for match in matches]
    items_found = []
    for match in matches:
        item_label = match["items"][0] if match["items"] else match["alias"]
        if len(match["items"]) > 1:
            item_label = f"{item_label} (+{len(match['items']) - 1} catalog variants)"
        items_found.append(item_label)

    return {
        "count": len(matches),
        "aliases": aliases_found,
        "items": items_found,
    }


def quantity_candidates(normalized):
    text = normalize_numbers(normalized)
    candidates = []

    patterns = [
        ("explicit_total", r"\bTOTAL\s+OF\s+(\d{1,4})\b"),
        ("number_before_sample_set", r"\b(\d{1,4})\s+SAMPLE\s+SETS?\b"),
        ("number_before_set_of_samples", r"\b(\d{1,4})\s+SETS?\s+OF\s+SAMPLES?\b"),
        ("number_before_sample_panel", r"\b(\d{1,4})\s+SAMPLE\s+(?:PANELS?|PLANKS?|BOARDS?)\b"),
        ("number_before_panel", r"\b(\d{1,4})\s+(?:PANELS?|PLANKS?|BOARDS?)\b"),
        ("number_before_sample_chip", r"\b(\d{1,4})\s+SAMPLE\s+CHIPS?\b"),
        ("number_before_chip", r"\b(\d{1,4})\s+CHIPS?\b"),
        ("number_before_pcs", r"\b(\d{1,4})\s+P\.?\s*C\.?\s*S\.?\b"),
        ("number_before_piece", r"\b(\d{1,4})\s+PIECES?\b"),
        ("number_before_cut", r"\b(\d{1,4})\s+CUTS?\b"),
        ("number_before_generic_sample", r"\b(\d{1,4})\s+SAMPLES?\b"),
        ("singular_sample_of", r"\b(?:1|A|AN)\s+SAMPLE\s+OF\b"),
        ("singular_cut_of", r"\b(?:1|A|AN)\s+CUT\s+OF\b"),
        ("singular_panel_of", r"\b(?:1|A|AN)\s+(?:PANEL|PLANK|BOARD)\s+OF\b"),
    ]

    for rule, pattern in patterns:
        for match in re.finditer(pattern, text):
            value = 1 if rule.startswith("singular_") else int(match.group(1))
            if value < 1 or value > 500:
                continue
            context_start = max(0, match.start() - 35)
            context_end = min(len(text), match.end() + 90)
            context = text[context_start:context_end]
            if re.search(r"\b(?:PACKAGE|PKG|BOX|BXS|CTNS?|CARTONS?|PALLETS?)\b", context):
                if not re.search(r"\b(?:SAMPLE|CHIP|PCS|PIECE|CUT|PANEL|PLANK|BOARD|SET)\b", context):
                    continue
            candidates.append(
                {
                    "qty": value,
                    "rule": rule,
                    "index": match.start(),
                    "context": context,
                }
            )

    return candidates


def estimate_listed_item_count(text, start_index):
    tail = text[start_index : start_index + 550]
    tail = re.sub(r"\([^)]*\)", " ", tail)
    tail = re.sub(
        r"\b(?:UPS|FEDEX|SHIP|SHIPPING|FREIGHT|DELIVERY|DLV|NO CHARGE|"
        r"PAID|VISA|AMEX|CHECK|CASH|CARD|TR#|TRACKING|QUOTE)\b.*$",
        "",
        tail,
    )
    parts = re.split(r"\s+\|\s+|,\s+|\s+AND\s+|(?<!\d)\s+-+\s*(?=[A-Z])", tail)
    cleaned = []
    for part in parts:
        value = re.sub(r"\s+", " ", part).strip(" :-*.,")
        if len(value) < 3:
            continue
        if re.search(
            r"\b(?:PACKAGE|PKG|BOX|BXS|CTNS?|SF|SQFT|LF|BF|S/F|LOT|PRICE)\b",
            value,
        ):
            continue
        if re.fullmatch(r"[\d\s./\"-]+", value):
            continue
        cleaned.append(value)
    return len(cleaned)


def parse_quantity(raw_text, sample_type, catalog_aliases):
    normalized = normalize_note(raw_text)
    candidates = quantity_candidates(normalized)
    flags = []
    catalog_count = ""
    catalog_aliases_found = ""
    catalog_items_found = ""

    if not sample_type:
        return "", "", "", catalog_count, catalog_aliases_found, catalog_items_found, "NO_SAMPLE_TYPE_FOR_QTY"

    explicit_totals = [c for c in candidates if c["rule"] == "explicit_total"]
    if explicit_totals:
        unique_totals = sorted({c["qty"] for c in explicit_totals})
        if len(unique_totals) == 1:
            return (
                str(unique_totals[0]),
                "parsed_from_order_notes",
                "explicit_total",
                catalog_count,
                catalog_aliases_found,
                catalog_items_found,
                "",
            )
        return (
            "",
            "needs_review",
            "explicit_total",
            catalog_count,
            catalog_aliases_found,
            catalog_items_found,
            "MULTIPLE_EXPLICIT_TOTALS_REVIEW_REQUIRED",
        )

    adjusted = []
    for candidate in candidates:
        qty = candidate["qty"]
        rule = candidate["rule"]
        context_after = normalized[candidate["index"] : candidate["index"] + 260]
        each_pattern = re.search(
            r"\b(?:OF\s+EACH|EACH\s+COLOR|EA(?:CH)?\s+COLOR|EACH)\b",
            context_after,
        )
        if each_pattern:
            each_abs_start = candidate["index"] + each_pattern.end()
            catalog_match = (
                find_catalog_each_matches(normalized, each_abs_start, catalog_aliases)
                if catalog_aliases
                else {"count": 0, "aliases": [], "items": []}
            )
            if catalog_match["count"] > 1:
                adjusted.append(
                    {
                        **candidate,
                        "qty": qty * catalog_match["count"],
                        "rule": f"{rule}_times_{catalog_match['count']}_catalog_items_after_each",
                        "catalog_count": catalog_match["count"],
                        "catalog_aliases": catalog_match["aliases"],
                        "catalog_items": catalog_match["items"],
                    }
                )
                continue

            item_count = estimate_listed_item_count(normalized, each_abs_start)
            if item_count > 1:
                adjusted.append(
                    {
                        **candidate,
                        "qty": qty * item_count,
                        "rule": f"{rule}_times_{item_count}_listed_items",
                        "catalog_count": catalog_match["count"],
                        "catalog_aliases": catalog_match["aliases"],
                        "catalog_items": catalog_match["items"],
                    }
                )
                flags.append("EACH_PATTERN_MULTIPLIED_BY_LISTED_ITEMS")
                continue
            flags.append("EACH_PATTERN_ITEM_COUNT_NOT_CONFIDENT")
        adjusted.append(candidate)

    if not adjusted:
        return "", "needs_review", "", catalog_count, catalog_aliases_found, catalog_items_found, "NO_SAMPLE_QUANTITY_FOUND"

    seen = {}
    for candidate in sorted(adjusted, key=lambda item: (item["index"], item["qty"])):
        seen[(candidate["qty"], candidate["rule"])] = candidate
    unique = list(seen.values())
    unique_qtys = sorted({item["qty"] for item in unique})
    rules = "; ".join(dict.fromkeys(item["rule"] for item in unique))
    best_catalog_count = max((item.get("catalog_count") or 0 for item in unique), default=0)
    if best_catalog_count:
        catalog_count = str(best_catalog_count)
        aliases = []
        items = []
        for item in unique:
            if (item.get("catalog_count") or 0) == best_catalog_count:
                aliases.extend(item.get("catalog_aliases") or [])
                items.extend(item.get("catalog_items") or [])
        catalog_aliases_found = "; ".join(dict.fromkeys(aliases))
        catalog_items_found = "; ".join(dict.fromkeys(items))

    if len(unique_qtys) == 1:
        return (
            str(unique_qtys[0]),
            "parsed_from_order_notes",
            rules,
            catalog_count,
            catalog_aliases_found,
            catalog_items_found,
            "; ".join(flags),
        )

    flags.append("MULTIPLE_QUANTITY_CANDIDATES_REVIEW_REQUIRED")
    return (
        "",
        "needs_review",
        rules,
        catalog_count,
        catalog_aliases_found,
        catalog_items_found,
        "; ".join(dict.fromkeys(flags)),
    )


def parse_dollars(raw_text):
    normalized = normalize_note(raw_text)
    amount_matches = list(re.finditer(r"\$\s*\d+(?:,\d{3})*(?:\.\d{1,2})?", raw_text or ""))
    amounts = []
    strict = []
    flags = []

    no_charge = bool(
        re.search(
            r"\b(?:NO|MO)\s+(?:UPS\s+|FREIGHT\s+|SHIPPING\s+|DELIVERY\s+|DLV\s+)?"
            r"(?:CHARGE|CHRAGE|CARGE|CHG)\b",
            normalized,
        )
        or re.search(r"\bN\s*/\s*C\b", normalized)
        or re.search(
            r"\b(?:UPS|FREIGHT|SHIPPING|DELIVERY|DLV|SAMPLE|SAMPLES).{0,45}"
            r"\b(?:NO|MO)\s+(?:CHARGE|CHRAGE|CARGE|CHG)\b",
            normalized,
        )
    )

    for match in amount_matches:
        amount = match.group(0)
        numeric_text = re.sub(r"[^\d.]", "", amount)
        try:
            numeric = Decimal(numeric_text).quantize(Decimal("0.01"))
        except InvalidOperation:
            continue
        amount_text = f"{numeric}"
        amounts.append(amount_text)

        start = max(0, match.start() - 80)
        end = min(len(raw_text or ""), match.end() + 100)
        context = normalize_note((raw_text or "")[start:end])
        has_freight_words = bool(
            re.search(
                r"\b(?:FREIGHT|SHIP|SHIPPING|UPS|FEDEX|FED EX|DELIVERY|DLV|GROUND|POSTAGE)\b",
                context,
            )
        )
        has_payment_words = bool(
            re.search(
                r"\b(?:PAID|PIAD|PAOID|VISA|AMEX|MASTER|M/C|MC|CASH|CHECK|CHK|"
                r"CK#|CARD|CREDIT|REFUND|DEPOSIT|BALANCE|ORDER#|MATERIAL|MAT\s+FRM|"
                r"PRICE|REBATE|RESTOCK|WIRE)\b",
                context,
            )
        )
        if has_freight_words and not has_payment_words:
            strict.append((amount_text, context))

    unique_amounts = list(dict.fromkeys(amounts))
    unique_strict = list(dict.fromkeys(amount for amount, _context in strict))
    source = ""
    selected = ""

    if len(unique_strict) == 1:
        selected = unique_strict[0]
        source = "strict_freight_shipping_context"
    elif len(unique_strict) > 1:
        source = "multiple_shipping_amounts_needs_review"
        flags.append("MULTIPLE_SHIPPING_DOLLAR_CANDIDATES")
    elif len(unique_amounts) == 1:
        selected = unique_amounts[0]
        source = "single_dollar_amount_in_order_notes"
        flags.append("VERIFY_AMOUNT_IS_FREIGHT_OR_SAMPLE_COST")
    elif len(unique_amounts) > 1:
        source = "multiple_dollar_amounts_needs_review"
        flags.append("MULTIPLE_DOLLAR_AMOUNTS_REVIEW_REQUIRED")

    return {
        "amount": selected,
        "source": source,
        "no_charge": "Y" if no_charge else "",
        "amounts_found": "; ".join(unique_amounts),
        "flags": flags,
    }


def order_key(row):
    return (
        clean_cell(row["COMPANY_NUMBER"]),
        clean_cell(row["LOCATION_NUMBER"]),
        clean_cell(row["RELEASE_NUMBER"]),
        clean_cell(row["ORDER_NUMBER"]),
        clean_cell(row["INVOICE_NUMBER"]),
    )


def display_key(row):
    return f"{clean_cell(row['ORDER_NUMBER'])}|{clean_cell(row['INVOICE_NUMBER'])}"


def load_prior_summary_labels(path):
    labels = {}
    if not path.exists():
        return labels

    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            order_number = str(row.get("ORDER_NUMBER", "")).strip()
            invoice_number = str(row.get("INVOICE_NUMBER", "")).strip()
            if not order_number or not invoice_number:
                continue

            key = f"{order_number}|{invoice_number}"
            entry = labels.setdefault(
                key,
                {
                    "types": [],
                    "qtys": [],
                    "workbook_rows": [],
                },
            )
            prior_type = str(row.get("TYPE", "")).strip()
            prior_qty = normalize_numeric_label(row.get("QTY", ""))
            workbook_row = str(row.get("WORKBOOK_ROW", "")).strip()

            if prior_type:
                entry["types"].append(prior_type)
            if prior_qty:
                entry["qtys"].append(prior_qty)
            if workbook_row:
                entry["workbook_rows"].append(workbook_row)

    collapsed = {}
    for key, entry in labels.items():
        collapsed[key] = {
            "type": "; ".join(dict.fromkeys(entry["types"])),
            "sample_type": normalize_prior_type("; ".join(entry["types"])),
            "qty": "; ".join(dict.fromkeys(entry["qtys"])),
            "workbook_row": "; ".join(dict.fromkeys(entry["workbook_rows"])),
        }
    return collapsed


def fetch_note_lines(start_date, end_date, chunk_days):
    sql = SQL_PATH.read_text(encoding="utf-8")
    conn = get_connection()
    all_rows = []
    try:
        cursor = conn.cursor()
        for window_start, window_end in date_windows(start_date, end_date, chunk_days):
            print(f"Pulling {window_start} through {window_end}...")
            cursor.execute(sql, window_start, window_end)
            headers = [column[0].upper() for column in cursor.description]
            fetched = cursor.fetchall()
            for db_row in fetched:
                all_rows.append(dict(zip(headers, db_row)))
            print(f"  note lines: {len(fetched)}")
    finally:
        conn.close()
    return all_rows


def read_note_lines_csv(path):
    with Path(path).open("r", newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def aggregate_orders(note_lines):
    grouped = {}
    for row in note_lines:
        key = order_key(row)
        if key not in grouped:
            grouped[key] = {
                "base": row,
                "notes": [],
            }
        note_text = clean_cell(row.get("NOTE_TEXT"))
        if note_text:
            grouped[key]["notes"].append(
                (
                    clean_cell(row.get("NOTE_SEQUENCE")),
                    clean_cell(row.get("NOTE_TEXT_SEQUENCE")),
                    clean_cell(row.get("NOTE_RRN")),
                    note_text,
                )
            )
    return grouped


def parse_order(group, prior_labels, catalog_aliases):
    base = group["base"]
    notes = [
        note_text
        for _seq, _text_seq, _rrn, note_text in sorted(
            group["notes"], key=lambda item: (item[0], item[1], item[2])
        )
    ]
    full_notes = " | ".join(notes)
    normalized = normalize_note(full_notes)
    terms = match_terms(normalized)
    sample_type, type_flags = detect_type(normalized)
    sample_type_source = "parsed_from_order_notes" if sample_type else ""
    prior = prior_labels.get(display_key(base), {})
    prior_sample_type = prior.get("sample_type", "")
    prior_type_raw = prior.get("type", "")
    prior_qty = prior.get("qty", "")
    prior_workbook_row = prior.get("workbook_row", "")

    if prior_sample_type and not sample_type:
        sample_type = prior_sample_type
        sample_type_source = "prior_summary_type"
        type_flags = [flag for flag in type_flags if flag != "NO_SAMPLE_TYPE_DETECTED"]
    elif prior_sample_type and sample_type:
        parsed_types = set(sample_type.split("; "))
        prior_types = set(prior_sample_type.split("; "))
        if not prior_types.issubset(parsed_types):
            type_flags.append("PRIOR_TYPE_DIFFERS_FROM_PARSED_TYPE")

    (
        qty,
        qty_source,
        qty_rule,
        catalog_count,
        catalog_aliases_found,
        catalog_items_found,
        qty_flags,
    ) = parse_quantity(full_notes, sample_type, catalog_aliases)
    parsed_qty_before_prior = qty
    if prior_qty:
        unique_prior_qtys = [value for value in prior_qty.split("; ") if value]
        if len(unique_prior_qtys) == 1:
            if parsed_qty_before_prior and normalize_numeric_label(parsed_qty_before_prior) != unique_prior_qtys[0]:
                qty_flags = "; ".join(
                    flag
                    for flag in [qty_flags, "PRIOR_QTY_DIFFERS_FROM_PARSED_QTY"]
                    if flag
                )
            qty = unique_prior_qtys[0]
            qty_source = "prior_summary_manual_qty"
        else:
            qty_flags = "; ".join(
                flag
                for flag in [qty_flags, "MULTIPLE_PRIOR_QTY_VALUES_REVIEW_REQUIRED"]
                if flag
            )

    dollar_info = parse_dollars(full_notes)

    flags = []
    flags.extend(type_flags)
    if qty_flags:
        flags.extend(flag.strip() for flag in qty_flags.split(";") if flag.strip())
    flags.extend(dollar_info["flags"])
    if dollar_info["no_charge"] == "Y" and dollar_info["amounts_found"]:
        flags.append("NO_CUSTOMER_CHARGE_LANGUAGE_PRESENT_WITH_DOLLAR_AMOUNT")
    if not terms:
        flags.append("NO_EXACT_NOTE_KEYWORD_AFTER_NORMALIZE")

    review_needed = "Y" if flags or qty_source == "needs_review" else ""

    return {
        "ORDER_NUMBER": clean_cell(base["ORDER_NUMBER"]),
        "INVOICE_NUMBER": clean_cell(base["INVOICE_NUMBER"]),
        "ORDER_DATE": clean_cell(base["ORDER_DATE"]),
        "ACCOUNT_NUMBER": clean_cell(base["ACCOUNT_NUMBER"]),
        "SAMPLE_TYPE": sample_type,
        "SAMPLE_TYPE_SOURCE": sample_type_source,
        "SAMPLE_QTY": qty,
        "SAMPLE_QTY_SOURCE": qty_source,
        "SAMPLE_QTY_RULE": qty_rule,
        "CATALOG_EACH_ITEM_COUNT": catalog_count,
        "CATALOG_EACH_ALIASES_FOUND": catalog_aliases_found,
        "CATALOG_EACH_ITEMS_FOUND": catalog_items_found,
        "PRIOR_SUMMARY_TYPE": prior_type_raw,
        "PRIOR_SUMMARY_QTY": prior_qty,
        "PRIOR_SUMMARY_WORKBOOK_ROW": prior_workbook_row,
        "FREIGHT_DOLLAR_AMOUNT": dollar_info["amount"],
        "FREIGHT_AMOUNT_SOURCE": dollar_info["source"],
        "CUSTOMER_NO_CHARGE_NOTE": dollar_info["no_charge"],
        "DOLLAR_AMOUNTS_FOUND": dollar_info["amounts_found"],
        "REVIEW_NEEDED": review_needed,
        "REVIEW_FLAGS": "; ".join(dict.fromkeys(flags)),
        "MATCH_TERMS_FOUND": "; ".join(terms),
        "GARTMAN_ORDER_DATE": clean_cell(base["GARTMAN_ORDER_DATE"]),
        "INVOICE_DATE": clean_cell(base["INVOICE_DATE"]),
        "COMPANY_NUMBER": clean_cell(base["COMPANY_NUMBER"]),
        "LOCATION_NUMBER": clean_cell(base["LOCATION_NUMBER"]),
        "RELEASE_NUMBER": clean_cell(base["RELEASE_NUMBER"]),
        "ORDER_INVOICE_KEY": display_key(base),
        "NONSTOCK_LINE_COUNT": clean_cell(base["NONSTOCK_LINE_COUNT"]),
        "NONSTOCK_QTY_UNITS": clean_cell(base["NONSTOCK_QTY_UNITS"]),
        "NONSTOCK_NET_AMOUNT": clean_cell(base["NONSTOCK_NET_AMOUNT"]),
        "NOTE_LINE_COUNT": str(len(notes)),
        "FULL_ORDER_NOTES": full_notes,
    }


def write_csv(path, headers, rows):
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def note_line_output_rows(note_lines):
    rows = []
    for row in note_lines:
        output = {key: clean_cell(value) for key, value in row.items()}
        output["ORDER_INVOICE_KEY"] = display_key(row)
        rows.append(output)
    return rows


def write_workbook(path, parsed_rows, review_rows, note_line_rows, stats):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print("openpyxl is not installed; skipping xlsx output.")
        return

    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    def populate(sheet, headers, rows):
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(width, 60)

    summary_sheet = workbook.active
    summary_sheet.title = "Parsed Orders"
    populate(summary_sheet, SUMMARY_HEADERS, parsed_rows)

    review_sheet = workbook.create_sheet("Needs Review")
    populate(review_sheet, SUMMARY_HEADERS, review_rows)

    note_headers = list(note_line_rows[0].keys()) if note_line_rows else []
    notes_sheet = workbook.create_sheet("Raw Note Lines")
    if note_headers:
        populate(notes_sheet, note_headers, note_line_rows)

    stats_sheet = workbook.create_sheet("Run Summary")
    stats_sheet.append(["Metric", "Value"])
    for key, value in stats.items():
        stats_sheet.append([key, value])
    for cell in stats_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    stats_sheet.freeze_panes = "A2"
    stats_sheet.column_dimensions["A"].width = 42
    stats_sheet.column_dimensions["B"].width = 80

    workbook.save(path)


def build_stats(parsed_rows, note_lines, start_date, end_date):
    type_counts = Counter(row["SAMPLE_TYPE"] or "(blank)" for row in parsed_rows)
    invoice_sets_by_order = defaultdict(set)
    for row in parsed_rows:
        invoice_sets_by_order[row["ORDER_NUMBER"]].add(row["INVOICE_NUMBER"])

    duplicate_order_numbers = sorted(
        order for order, invoices in invoice_sets_by_order.items() if len(invoices) > 1
    )
    order_454480_invoices = sorted(
        row["INVOICE_NUMBER"] for row in parsed_rows if row["ORDER_NUMBER"] == "454480"
    )

    return {
        "Start date filter": start_date.isoformat(),
        "End date filter": end_date.isoformat(),
        "Date field used for filter": "SHHEAD.SHIDAT / INVOICE_DATE",
        "Raw note line rows": len(note_lines),
        "Parsed order/invoice rows": len(parsed_rows),
        "Rows needing review": sum(1 for row in parsed_rows if row["REVIEW_NEEDED"] == "Y"),
        "Rows with filled quantity": sum(1 for row in parsed_rows if row["SAMPLE_QTY"]),
        "Rows using prior manual quantity": sum(
            1 for row in parsed_rows if row["SAMPLE_QTY_SOURCE"] == "prior_summary_manual_qty"
        ),
        "Rows with prior summary type": sum(
            1 for row in parsed_rows if row["PRIOR_SUMMARY_TYPE"]
        ),
        "Rows using catalog EACH multiplier": sum(
            1 for row in parsed_rows if row["CATALOG_EACH_ITEM_COUNT"]
        ),
        "Rows with freight dollar amount": sum(
            1 for row in parsed_rows if row["FREIGHT_DOLLAR_AMOUNT"]
        ),
        "Rows with customer no-charge note": sum(
            1 for row in parsed_rows if row["CUSTOMER_NO_CHARGE_NOTE"] == "Y"
        ),
        "Duplicate order numbers with multiple invoices in output": len(
            duplicate_order_numbers
        ),
        "Order 454480 invoices in output": ", ".join(order_454480_invoices),
        "Type counts": "; ".join(f"{key}: {value}" for key, value in type_counts.items()),
    }


def write_run_summary(path, stats):
    lines = [f"{key}: {value}" for key, value in stats.items()]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    args = parse_args()
    start_date = parse_date(args.start_date)
    end_date = parse_date(args.end_date)
    if start_date > end_date:
        raise SystemExit("start-date must be on or before end-date")
    if args.chunk_days < 1:
        raise SystemExit("chunk-days must be at least 1")

    if args.reuse_note_lines:
        print(f"Reading existing note lines from {args.reuse_note_lines}...")
        note_lines = read_note_lines_csv(args.reuse_note_lines)
    else:
        note_lines = fetch_note_lines(start_date, end_date, args.chunk_days)

    if not note_lines:
        raise SystemExit("No note lines found for the requested date range.")

    print("Loading product catalog aliases...")
    catalog_aliases = build_catalog_aliases()
    write_catalog_alias_csv(catalog_aliases)
    print(f"  catalog aliases: {len(catalog_aliases)}")

    note_line_rows = note_line_output_rows(note_lines)
    note_headers = list(note_line_rows[0].keys())
    write_csv(NOTE_LINE_CSV, note_headers, note_line_rows)

    grouped = aggregate_orders(note_lines)
    prior_labels = load_prior_summary_labels(PRIOR_SUMMARY_PATH)
    parsed_rows = [
        parse_order(group, prior_labels, catalog_aliases) for group in grouped.values()
    ]
    parsed_rows.sort(
        key=lambda row: (
            row["ORDER_DATE"],
            int(row["ORDER_NUMBER"]) if row["ORDER_NUMBER"].isdigit() else row["ORDER_NUMBER"],
            row["INVOICE_NUMBER"],
            row["RELEASE_NUMBER"],
        )
    )
    review_rows = [row for row in parsed_rows if row["REVIEW_NEEDED"] == "Y"]

    write_csv(PARSED_CSV, SUMMARY_HEADERS, parsed_rows)
    write_csv(REVIEW_CSV, SUMMARY_HEADERS, review_rows)

    stats = build_stats(parsed_rows, note_lines, start_date, end_date)
    write_run_summary(RUN_SUMMARY_PATH, stats)
    if not args.no_xlsx:
        write_workbook(WORKBOOK_PATH, parsed_rows, review_rows, note_line_rows, stats)

    print()
    for key, value in stats.items():
        print(f"{key}: {value}")
    print()
    print(f"Created: {PARSED_CSV}")
    print(f"Created: {REVIEW_CSV}")
    print(f"Created: {NOTE_LINE_CSV}")
    print(f"Created: {CATALOG_ALIAS_CSV}")
    if not args.no_xlsx:
        print(f"Created: {WORKBOOK_PATH}")
    print(f"Created: {RUN_SUMMARY_PATH}")


if __name__ == "__main__":
    main()
