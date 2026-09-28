from __future__ import annotations

from copy import copy
from datetime import datetime
from pathlib import Path
import shutil
import sys

import openpyxl
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from update_gfalo7508_from_image import image_metrics


WORKBOOK = Path(r"H:\2025\MISC Reports\For Marketing\Garrison Color Analysis 2.12.2026.xlsx")
IMAGE_PATH = Path(r"H:\2025\MISC Reports\For Marketing\Overheads for Color Analysis\European Oak Ventasso 7.5.webp")
BACKUP = WORKBOOK.with_name(
    f"{WORKBOOK.stem} backup before GFALO7507 {datetime.now():%Y%m%d_%H%M%S}{WORKBOOK.suffix}"
)


EXTRA_HEADERS = [
    "ProductNumber",
    "Species",
    "Construction",
    "Grade",
    "Finish",
    "SurfaceTreatment",
    "JankaRating",
    "Thickness",
    "TopLayer",
    "Length",
    "InstallationMethods",
    "EdgeDetail",
    "Warranty",
    "ImageAnalysisSource",
    "ImageWidthPx",
    "ImageHeightPx",
]


def header_map(ws):
    return {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}


def copy_cell_style(src, dst):
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)


def rgb_fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=f"FF{hex_color.lstrip('#').upper()}")


def ensure_columns(ws):
    widths = {
        "ProductNumber": 16,
        "Species": 18,
        "Construction": 26,
        "Grade": 20,
        "Finish": 18,
        "SurfaceTreatment": 22,
        "JankaRating": 14,
        "Thickness": 14,
        "TopLayer": 14,
        "Length": 20,
        "InstallationMethods": 28,
        "EdgeDetail": 18,
        "Warranty": 70,
        "ImageAnalysisSource": 70,
        "ImageWidthPx": 14,
        "ImageHeightPx": 15,
    }
    headers = header_map(ws)
    for name in EXTRA_HEADERS:
        if name in headers:
            continue
        col = ws.max_column + 1
        copy_cell_style(ws.cell(1, col - 1), ws.cell(1, col))
        ws.cell(1, col).value = name
        ws.column_dimensions[ws.cell(1, col).column_letter].width = widths.get(name, 18)
        for row in range(2, ws.max_row + 1):
            copy_cell_style(ws.cell(row, col - 1), ws.cell(row, col))
        headers = header_map(ws)


def set_row(ws, row_idx, values):
    headers = header_map(ws)
    template_row = max(2, row_idx - 1)
    for col_idx in range(1, ws.max_column + 1):
        copy_cell_style(ws.cell(template_row, col_idx), ws.cell(row_idx, col_idx))
    for header, value in values.items():
        if header in headers:
            ws.cell(row_idx, headers[header]).value = value
    ws.cell(row_idx, headers["Swatch"]).fill = rgb_fill(values["MeanHex"])


def main():
    metrics = image_metrics(IMAGE_PATH)
    shutil.copy2(WORKBOOK, BACKUP)
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb["Products"]
    ensure_columns(ws)
    headers = header_map(ws)

    product_number = "GFALO7507"
    row_idx = None
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, headers["ProductNumber"]).value or "").strip().upper() == product_number:
            row_idx = row
            break
    if row_idx is None:
        row_idx = ws.max_row + 1

    next_index = max(
        cell.value for cell in ws["A"][1:] if isinstance(cell.value, (int, float))
    ) + (0 if row_idx <= ws.max_row and ws.cell(row_idx, headers["ProductNumber"]).value == product_number else 1)
    if row_idx <= ws.max_row and ws.cell(row_idx, headers["ProductNumber"]).value == product_number:
        next_index = ws.cell(row_idx, headers["INDEX"]).value

    values = {
        "INDEX": next_index,
        "Collection": "Allora",
        "Product": "Ventasso",
        "Width_in": 7.5,
        "HueBucket": 10,
        "BucketName": "Smoky Espresso",
        "HueFamily": "Brown",
        "Undertone": "Warm",
        "Notes": 'Gartman item description: EURO OAK 7.5" VENTASSO. Metrics computed from source WebP file.',
        "ProductNumber": product_number,
        "Species": "European Oak",
        "Construction": "Engineered (Birch Plywood)",
        "Grade": "Select Character",
        "Finish": "Matte UV Lacquer",
        "SurfaceTreatment": "Light Wire-Brush",
        "JankaRating": 1360,
        "Thickness": '5/8"',
        "TopLayer": "4mm",
        "Length": "Mostly 7'8 1/2\"",
        "InstallationMethods": "Glue Down, Floating",
        "EdgeDetail": "Micro-Bevel",
        "Warranty": "Residential: Limited Lifetime Structure, Limited 25 Year Finish | Commercial: Limited 25 Year Structure, Limited 10 Year Finish",
        "ImageAnalysisSource": str(IMAGE_PATH),
    }
    values.update(metrics)

    set_row(ws, row_idx, values)
    wb.save(WORKBOOK)

    print(f"Updated: {WORKBOOK}")
    print(f"Backup:  {BACKUP}")
    print(f"Row:     {row_idx}")
    print("Bucket:  10 - Smoky Espresso")
    for key in [
        "MeanHex",
        "MedianHex",
        "MeanRGB_R",
        "MeanRGB_G",
        "MeanRGB_B",
        "L_mean",
        "a_mean",
        "b_mean",
        "HueDeg",
        "Chroma",
        "Sat_mean",
        "Value_mean",
        "Contrast_L_std",
        "EdgeDensity",
        "WarmthScore",
        "ImageWidthPx",
        "ImageHeightPx",
    ]:
        print(f"{key}: {metrics[key]}")


if __name__ == "__main__":
    main()
