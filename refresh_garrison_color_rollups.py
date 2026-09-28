from __future__ import annotations

from copy import copy
from datetime import datetime
from pathlib import Path
import shutil

import openpyxl
from openpyxl.styles import PatternFill


WORKBOOK = Path(r"H:\2025\MISC Reports\For Marketing\Garrison Color Analysis 2.12.2026.xlsx")
BACKUP = WORKBOOK.with_name(
    f"{WORKBOOK.stem} backup before rollup refresh {datetime.now():%Y%m%d_%H%M%S}{WORKBOOK.suffix}"
)


def header_map(ws):
    return {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}


def copy_cell_style(src, dst):
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)


def rgb_fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=f"FF{hex_color.lstrip('#').upper()}")


def ensure_product_column(ws, name: str, width: float = 18) -> None:
    headers = header_map(ws)
    if name in headers:
        return
    new_col = ws.max_column + 1
    copy_cell_style(ws.cell(1, new_col - 1), ws.cell(1, new_col))
    ws.cell(1, new_col).value = name
    ws.column_dimensions[ws.cell(1, new_col).column_letter].width = width
    for r in range(2, ws.max_row + 1):
        copy_cell_style(ws.cell(r, new_col - 1), ws.cell(r, new_col))


def set_gfalo7508_bucket_and_dimensions(wb):
    ws = wb["Products"]
    ensure_product_column(ws, "ImageWidthPx", 14)
    ensure_product_column(ws, "ImageHeightPx", 15)
    headers = header_map(ws)
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, headers["ProductNumber"]).value or "").strip().upper() == "GFALO7508":
            ws.cell(r, headers["HueBucket"]).value = 6
            ws.cell(r, headers["BucketName"]).value = "Amber"
            ws.cell(r, headers["HueFamily"]).value = "Amber/Caramel"
            ws.cell(r, headers["ImageWidthPx"]).value = 768
            ws.cell(r, headers["ImageHeightPx"]).value = 755
            return
    raise RuntimeError("GFALO7508 row not found")


def products(ws):
    headers = header_map(ws)
    for r in range(2, ws.max_row + 1):
        bucket = ws.cell(r, headers["HueBucket"]).value
        if bucket is None:
            continue
        yield {
            "row": r,
            "bucket": int(bucket),
            "collection": ws.cell(r, headers["Collection"]).value,
            "L_mean": float(ws.cell(r, headers["L_mean"]).value or 0),
            "HueDeg": float(ws.cell(r, headers["HueDeg"]).value or 0),
            "Chroma": float(ws.cell(r, headers["Chroma"]).value or 0),
            "Sat_mean": float(ws.cell(r, headers["Sat_mean"]).value or 0),
            "MeanHex": ws.cell(r, headers["MeanHex"]).value,
        }


def refresh_buckets(wb):
    product_rows = list(products(wb["Products"]))
    ws = wb["Buckets"]
    headers = header_map(ws)

    for r in range(2, ws.max_row + 1):
        bucket = ws.cell(r, headers["HueBucket"]).value
        rows = [p for p in product_rows if p["bucket"] == bucket]
        if not rows:
            continue
        count = len(rows)
        ws.cell(r, headers["Count"]).value = count
        ws.cell(r, headers["Avg_L"]).value = sum(p["L_mean"] for p in rows) / count
        ws.cell(r, headers["Avg_Hue"]).value = sum(p["HueDeg"] for p in rows) / count
        ws.cell(r, headers["Avg_Chroma"]).value = sum(p["Chroma"] for p in rows) / count
        ws.cell(r, headers["Avg_Sat"]).value = sum(p["Sat_mean"] for p in rows) / count

        # Preserve the established representative color unless this is the
        # bucket that received the new image; give Amber a refreshed swatch.
        if int(bucket) == 6:
            ws.cell(r, headers["RepresentativeHex"]).value = "#AA8052"
            ws.cell(r, headers["Swatch"]).fill = rgb_fill("#AA8052")


def refresh_summary(wb):
    product_rows = list(products(wb["Products"]))
    ws = wb["Summary"]
    headers = header_map(ws)
    collections = sorted({p["collection"] for p in product_rows if p["collection"]})

    for collection in collections:
        if collection not in headers:
            new_col = ws.max_column + 1
            copy_cell_style(ws.cell(1, new_col - 1), ws.cell(1, new_col))
            ws.cell(1, new_col).value = collection
            ws.column_dimensions[ws.cell(1, new_col).column_letter].width = max(12, len(collection) + 2)
            for r in range(2, ws.max_row + 1):
                copy_cell_style(ws.cell(r, new_col - 1), ws.cell(r, new_col))
            headers = header_map(ws)

    for r in range(2, ws.max_row + 1):
        bucket = ws.cell(r, headers["HueBucket"]).value
        for collection in collections:
            ws.cell(r, headers[collection]).value = sum(
                1 for p in product_rows if p["bucket"] == bucket and p["collection"] == collection
            )


def main():
    shutil.copy2(WORKBOOK, BACKUP)
    wb = openpyxl.load_workbook(WORKBOOK)
    set_gfalo7508_bucket_and_dimensions(wb)
    refresh_buckets(wb)
    refresh_summary(wb)
    wb.save(WORKBOOK)
    print(f"Updated: {WORKBOOK}")
    print(f"Backup:  {BACKUP}")


if __name__ == "__main__":
    main()
