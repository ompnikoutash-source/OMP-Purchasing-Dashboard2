from __future__ import annotations

import argparse
import datetime as dt
import math
import warnings
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter

from gartman_connection import get_connection


ARCHIVE_PATH = Path(r"H:\2025\MISC Reports\For Marketing\Blast Archive.xlsx")
OUTPUT_PATH = Path(r"H:\2025\MISC Reports\For Marketing\Blast Impact Analysis.xlsx")

WINDOW_DAYS = [30, 90, 180, 365]
INVALID_ITEM_MARKERS = {"", "NAN", "NONE", "N/A", "NA", "SKIP", "SKIP FOR NOW"}
FAMILY_CODE_ALIASES = {
    # The FAMILY table uses SF = SOLID FLOORING; the marketing sheet uses SO.
    "SO": "SF",
}
COLLECTION_ALIASES = {
    "GARRISON II SMOOTH": "GII SMOOTH",
    "GARRISON II DISTRESSED": "GII DISTRESSED",
}
INTERNAL_CUSTOMER_FILTERS = ["TRANSFER", "OMP000", "INV000", "OLD001"]


def clean_text(value: Any) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def norm_key(value: Any) -> str:
    return clean_text(value).upper()


def index_key(value: Any) -> str:
    if pd.isna(value):
        return ""
    try:
        number = float(value)
        if math.isfinite(number) and number.is_integer():
            return str(int(number))
        return f"{number:g}"
    except (TypeError, ValueError):
        return clean_text(value)


def sql_date(value: pd.Timestamp) -> str:
    return pd.Timestamp(value).strftime("%Y-%m-%d")


def pct_change(current: float | None, baseline: float | None) -> float:
    if current is None or baseline is None:
        return np.nan
    if pd.isna(current) or pd.isna(baseline) or abs(float(baseline)) < 1e-12:
        return np.nan
    return (float(current) / float(baseline) - 1.0) * 100.0


def read_archive(path: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    timeline = pd.read_excel(path, sheet_name="Timeline", dtype=object)
    item_index = pd.read_excel(path, sheet_name="Item Index", dtype=object)
    comparison = pd.read_excel(path, sheet_name="Comparison List", dtype=object)

    for df in (timeline, item_index, comparison):
        df.columns = [clean_text(c) for c in df.columns]
        df["INDEX_KEY"] = df["INDEX"].map(index_key)
        df["INDEX_SORT"] = pd.to_numeric(df["INDEX"], errors="coerce")

    timeline["Date Start"] = pd.to_datetime(timeline["Date Start"], errors="coerce")
    timeline["Date End"] = pd.to_datetime(timeline["Date End"], errors="coerce")
    timeline = timeline.dropna(subset=["INDEX_KEY", "Date Start", "Date End"])
    item_index = item_index[item_index["INDEX_KEY"] != ""].copy()
    comparison = comparison[comparison["INDEX_KEY"] != ""].copy()

    return timeline, item_index, comparison


def fetch_master(conn) -> pd.DataFrame:
    query = """
    SELECT
        TRIM(IM.IMITEM) AS ITEM_NUMBER,
        TRIM(IM.IMDESC) AS DESCRIPTION,
        RTRIM(CHAR(IM.IMDIV)) AS DIVISION,
        TRIM(IM.IMFMCD) AS FAMILY_CODE,
        TRIM(FM.FMDESC) AS FAMILY_NAME,
        TRIM(IX.IMCOLLECT) AS COLLECTION,
        TRIM(CHAR(IM.IMVEND)) AS VENDOR_NUMBER,
        TRIM(VM.VMNAME) AS VENDOR_NAME,
        TRIM(IM.IMUM2) AS SALES_UOM
    FROM GSFL2K.ITEMMAST IM
    LEFT JOIN GSFL2K.ITEMXTRA IX
      ON IX.IMXITM = IM.IMITEM
    LEFT JOIN GSFL2K.FAMILY FM
      ON TRIM(FM.FMFMCD) = TRIM(IM.IMFMCD)
    LEFT JOIN GSFL2K.VENDMAST VM
      ON VM.VMVEND = IM.IMVEND
    ORDER BY TRIM(IM.IMITEM)
    """
    master = pd.read_sql(query, conn)
    master.columns = [c.upper() for c in master.columns]
    for col in [
        "ITEM_NUMBER",
        "DESCRIPTION",
        "DIVISION",
        "FAMILY_CODE",
        "FAMILY_NAME",
        "COLLECTION",
        "VENDOR_NUMBER",
        "VENDOR_NAME",
        "SALES_UOM",
    ]:
        master[col] = master[col].map(clean_text)
    master["ITEM_KEY"] = master["ITEM_NUMBER"].str.upper()
    master["DIVISION_KEY"] = master["DIVISION"].map(norm_key)
    master["FAMILY_KEY"] = master["FAMILY_CODE"].map(norm_key)
    master["COLLECTION_KEY"] = master["COLLECTION"].map(norm_key)
    master["VENDOR_KEY"] = master["VENDOR_NUMBER"].map(norm_key)
    return master.drop_duplicates(subset=["ITEM_KEY"], keep="first").copy()


def fetch_sales_lines(conn, start_date: pd.Timestamp, end_date: pd.Timestamp) -> pd.DataFrame:
    query = f"""
    SELECT
        H.SHIDAT AS SALES_DATE,
        TRIM(L.SLITEM) AS ITEM_NUMBER,
        DECIMAL(COALESCE(L.SLBLUS, 0), 18, 4) AS UNITS_SOLD,
        DECIMAL(COALESCE(L.SLENET, 0), 18, 2) AS REVENUE,
        RTRIM(CHAR(H.SHCO))
            || '|' || RTRIM(CHAR(H.SHLOC))
            || '|' || RTRIM(CHAR(H.SHORD#))
            || '|' || RTRIM(CHAR(H.SHINV#)) AS ORDER_KEY,
        TRIM(H.SHCUST) AS CUSTOMER_NUMBER
    FROM GSFL2K.SHHEAD H
    JOIN GSFL2K.SHLINE L
      ON L.SLCO = H.SHCO
     AND L.SLLOC = H.SHLOC
     AND L.SLORD# = H.SHORD#
     AND L.SLINV# = H.SHINV#
    WHERE H.SHIDAT >= DATE('{sql_date(start_date)}')
      AND H.SHIDAT <= DATE('{sql_date(end_date)}')
      AND TRIM(L.SLITEM) <> ''
      AND UPPER(TRIM(L.SLITEM)) NOT LIKE 'PROMO%'
      AND (COALESCE(L.SLBLUS, 0) <> 0 OR COALESCE(L.SLENET, 0) <> 0)
      AND UPPER(TRIM(H.SHCUST)) NOT LIKE '%TRANSFER%'
      AND UPPER(TRIM(H.SHCUST)) NOT LIKE '%OMP000%'
      AND UPPER(TRIM(H.SHCUST)) NOT LIKE '%INV000%'
      AND UPPER(TRIM(H.SHCUST)) NOT LIKE '%OLD001%'
    """
    sales = pd.read_sql(query, conn)
    sales.columns = [c.upper() for c in sales.columns]
    sales["SALES_DATE"] = pd.to_datetime(sales["SALES_DATE"], errors="coerce").dt.normalize()
    sales["ITEM_NUMBER"] = sales["ITEM_NUMBER"].map(clean_text)
    sales["ITEM_KEY"] = sales["ITEM_NUMBER"].str.upper()
    sales["UNITS_SOLD"] = pd.to_numeric(sales["UNITS_SOLD"], errors="coerce").fillna(0.0)
    sales["REVENUE"] = pd.to_numeric(sales["REVENUE"], errors="coerce").fillna(0.0)
    sales["ORDER_KEY"] = sales["ORDER_KEY"].map(clean_text)
    return sales.dropna(subset=["SALES_DATE"]).copy()


def get_db_date_range(conn) -> tuple[pd.Timestamp, pd.Timestamp, pd.Timestamp]:
    query = """
    SELECT
        CURRENT DATE AS DB_TODAY,
        MIN(H.SHIDAT) AS MIN_DATE,
        MAX(H.SHIDAT) AS MAX_DATE
    FROM GSFL2K.SHHEAD H
    JOIN GSFL2K.SHLINE L
      ON L.SLCO = H.SHCO
     AND L.SLLOC = H.SHLOC
     AND L.SLORD# = H.SHORD#
     AND L.SLINV# = H.SHINV#
    """
    row = pd.read_sql(query, conn).iloc[0]
    return (
        pd.Timestamp(row["DB_TODAY"]).normalize(),
        pd.Timestamp(row["MIN_DATE"]).normalize(),
        pd.Timestamp(row["MAX_DATE"]).normalize(),
    )


def build_campaigns(timeline: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for idx, group in timeline.sort_values(["INDEX_SORT", "Chrono"]).groupby("INDEX_KEY", sort=False):
        starts = group["Date Start"].dropna()
        ends = group["Date End"].dropna()
        start = starts.min().normalize()
        end = ends.max().normalize()
        windows = (
            group[["Date Start", "Date End"]]
            .drop_duplicates()
            .sort_values(["Date Start", "Date End"])
        )
        cards = " | ".join(clean_text(v) for v in group["Card Name"].dropna().tolist())
        filter_types = " | ".join(sorted({clean_text(v) for v in group["Item Filter Type"].dropna() if clean_text(v)}))
        filters = " | ".join(sorted({clean_text(v) for v in group["Item Filter"].dropna() if clean_text(v)}))
        breadth = " | ".join(sorted({clean_text(v) for v in group["Breadth"].dropna() if clean_text(v)}))
        discount_types = " | ".join(sorted({clean_text(v) for v in group["Discount Type"].dropna() if clean_text(v)}))
        discount_amounts = " | ".join(sorted({clean_text(v) for v in group["Discount Amount/Unit"].dropna() if clean_text(v)}))
        duration_days = int((end - start).days + 1)

        warnings_for_index: list[str] = []
        if len(windows) > 1:
            warnings_for_index.append("multiple timeline rows merged to outer date span")
        if duration_days > 180:
            warnings_for_index.append("effective window exceeds 180 days; check dates")
        if start.year < 2026:
            warnings_for_index.append("effective window starts before 2026; check dates")
        if group["Breadth"].isna().any():
            warnings_for_index.append("missing breadth on at least one send")
        if group["Item Filter Type"].isna().any():
            warnings_for_index.append("missing item filter type on at least one send")

        rows.append(
            {
                "INDEX": idx,
                "Index Sort": group["INDEX_SORT"].iloc[0],
                "Send Count": len(group),
                "Effective Start": start,
                "Effective End": end,
                "Effective Days": duration_days,
                "Single-Day Effective Period": duration_days == 1,
                "Distinct Timeline Windows": len(windows),
                "Breadth": breadth,
                "Timeline Item Filter Type": filter_types,
                "Timeline Item Filter": filters,
                "Discount Type": discount_types,
                "Discount Amount/Unit": discount_amounts,
                "Card Names": cards,
                "Date Warning": "; ".join(warnings_for_index),
            }
        )
    return pd.DataFrame(rows).sort_values("Index Sort").reset_index(drop=True)


def item_numbers_for_divisions(master: pd.DataFrame, divisions: set[str]) -> set[str]:
    div_keys = {norm_key(v) for v in divisions if norm_key(v)}
    return set(master.loc[master["DIVISION_KEY"].isin(div_keys), "ITEM_KEY"])


def clean_item_number(value: Any) -> str:
    item = norm_key(value)
    return "" if item in INVALID_ITEM_MARKERS else item


def has_total_revenue_comparison(index: str, comparison: pd.DataFrame) -> bool:
    rows = comparison[comparison["INDEX_KEY"] == index]
    return any(norm_key(v) == "TOTAL REVENUE" for v in rows["Filter"].tolist())


def build_target_set(
    index: str,
    item_index: pd.DataFrame,
    comparison: pd.DataFrame,
    master: pd.DataFrame,
) -> dict[str, Any]:
    rows = item_index[item_index["INDEX_KEY"] == index]
    explicit_items = {
        clean_item_number(v)
        for v in rows.get("Item Numbers", pd.Series(dtype=object)).tolist()
    }
    explicit_items = {v for v in explicit_items if v}
    raw_items = [norm_key(v) for v in rows.get("Item Numbers", pd.Series(dtype=object)).tolist()]
    invalid_markers = sorted({v for v in raw_items if v in INVALID_ITEM_MARKERS and v})

    divs = {
        norm_key(v)
        for v in rows.get("Division Number", pd.Series(dtype=object)).tolist()
        if norm_key(v)
    }
    division_items = item_numbers_for_divisions(master, divs) if divs else set()
    item_set = explicit_items | division_items

    if item_set:
        source_bits = []
        if explicit_items:
            source_bits.append(f"{len(explicit_items)} explicit item(s)")
        if divs:
            source_bits.append(f"division(s) {', '.join(sorted(divs))}")
        return {
            "kind": "ITEM_SET",
            "items": item_set,
            "source": "; ".join(source_bits),
            "explicit_items": explicit_items,
            "division_numbers": divs,
            "invalid_markers": invalid_markers,
        }

    if rows.empty and has_total_revenue_comparison(index, comparison):
        return {
            "kind": "TOTAL_REVENUE",
            "items": None,
            "source": "TOTAL REVENUE comparison row, no item-index rows",
            "explicit_items": set(),
            "division_numbers": set(),
            "invalid_markers": [],
        }

    return {
        "kind": "EMPTY",
        "items": set(),
        "source": "no valid promoted items/divisions",
        "explicit_items": explicit_items,
        "division_numbers": divs,
        "invalid_markers": invalid_markers,
    }


def resolve_filter_items(filter_name: str, specific: Any, master: pd.DataFrame) -> tuple[set[str] | None, str, str]:
    filter_key = norm_key(filter_name)
    spec_key = norm_key(specific)
    note = ""

    if filter_key == "TOTAL REVENUE":
        return None, "TOTAL_REVENUE", "all non-promo, non-internal invoice lines"

    if filter_key == "COLLECTION":
        mapped = COLLECTION_ALIASES.get(spec_key, spec_key)
        if mapped != spec_key:
            note = f"collection alias {specific} mapped to {mapped}"
        spec_key = mapped
        mask = master["COLLECTION_KEY"] == spec_key
    elif filter_key == "FAMILY":
        mapped = FAMILY_CODE_ALIASES.get(spec_key, spec_key)
        if mapped != spec_key:
            note = f"family alias {spec_key} mapped to {mapped}"
        mask = master["FAMILY_KEY"] == mapped
    elif filter_key == "DIVISION":
        mask = master["DIVISION_KEY"] == spec_key
    elif filter_key in {"VENDOR", "VENDOR NUMBER"}:
        mask = master["VENDOR_KEY"] == spec_key
    else:
        return set(), "UNKNOWN", f"unsupported comparison filter: {filter_name}"

    return set(master.loc[mask, "ITEM_KEY"]), "ITEM_SET", note


def build_comparison_set(
    index: str,
    comparison: pd.DataFrame,
    master: pd.DataFrame,
    target: dict[str, Any],
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    rows = comparison[comparison["INDEX_KEY"] == index]
    filter_audit: list[dict[str, Any]] = []
    if rows.empty:
        return {
            "kind": "EMPTY",
            "items": set(),
            "source": "no comparison rows",
            "filters": "",
        }, filter_audit

    union_items: set[str] = set()
    total_revenue = False
    source_bits: list[str] = []
    target_items = target["items"] if target["kind"] == "ITEM_SET" else set()

    for _, row in rows.iterrows():
        filter_name = clean_text(row["Filter"])
        specific = clean_text(row["Specific"])
        matched, kind, note = resolve_filter_items(filter_name, specific, master)
        before_count = np.nan if matched is None else len(matched)

        if kind == "TOTAL_REVENUE":
            total_revenue = True
            after_count = np.nan
            source_bits.append("TOTAL REVENUE")
        elif kind == "ITEM_SET":
            after_items = matched - target_items
            union_items |= after_items
            after_count = len(after_items)
            source_bits.append(f"{filter_name}: {specific}")
        else:
            after_count = 0
            source_bits.append(f"{filter_name}: {specific}")

        filter_audit.append(
            {
                "INDEX": index,
                "Filter": filter_name,
                "Specific": specific,
                "Matched Items Before Exclusion": before_count,
                "Matched Items After Promoted Exclusion": after_count,
                "Excluded Promoted Items": "" if matched is None else len((matched or set()) & target_items),
                "Note": note,
            }
        )

    if total_revenue:
        if target["kind"] == "TOTAL_REVENUE":
            return {
                "kind": "EMPTY",
                "items": set(),
                "source": "comparison suppressed because target is TOTAL REVENUE",
                "filters": "; ".join(source_bits),
            }, filter_audit
        return {
            "kind": "TOTAL_REVENUE",
            "items": None,
            "source": "TOTAL REVENUE",
            "filters": "; ".join(source_bits),
        }, filter_audit

    if union_items:
        return {
            "kind": "ITEM_SET",
            "items": union_items,
            "source": "comparison filters minus promoted items",
            "filters": "; ".join(source_bits),
        }, filter_audit

    return {
        "kind": "EMPTY",
        "items": set(),
        "source": "comparison filters matched no items after promoted exclusions",
        "filters": "; ".join(source_bits),
    }, filter_audit


def period_bounds(
    phase: str,
    window_days: int,
    campaign_start: pd.Timestamp,
    campaign_end: pd.Timestamp,
    data_min: pd.Timestamp,
    data_max: pd.Timestamp,
) -> tuple[pd.Timestamp | None, pd.Timestamp | None, int, bool]:
    if phase == "Before":
        wanted_start = campaign_start - pd.Timedelta(days=window_days)
        wanted_end = campaign_start - pd.Timedelta(days=1)
        start = max(wanted_start, data_min)
        end = min(wanted_end, data_max)
        requested_days = window_days
    elif phase == "After":
        wanted_start = campaign_end + pd.Timedelta(days=1)
        wanted_end = campaign_end + pd.Timedelta(days=window_days)
        start = max(wanted_start, data_min)
        end = min(wanted_end, data_max)
        requested_days = window_days
    elif phase == "During":
        start = max(campaign_start, data_min)
        end = min(campaign_end, data_max)
        requested_days = int((campaign_end - campaign_start).days + 1)
    else:
        raise ValueError(f"Unknown phase: {phase}")

    if pd.isna(start) or pd.isna(end) or start > end:
        return None, None, 0, False

    actual_days = int((end - start).days + 1)
    return start, end, actual_days, actual_days == requested_days


def summarize_period(
    sales: pd.DataFrame,
    cohort_kind: str,
    item_set: set[str] | None,
    start: pd.Timestamp | None,
    end: pd.Timestamp | None,
    actual_days: int,
    window_days: int,
) -> dict[str, Any]:
    if start is None or end is None or actual_days <= 0 or cohort_kind == "EMPTY":
        units = revenue = orders = 0.0
    else:
        mask = (sales["SALES_DATE"] >= start) & (sales["SALES_DATE"] <= end)
        if cohort_kind == "ITEM_SET":
            mask &= sales["ITEM_KEY"].isin(item_set or set())
        matched = sales.loc[mask]
        units = float(matched["UNITS_SOLD"].sum())
        revenue = float(matched["REVENUE"].sum())
        orders = float(matched["ORDER_KEY"].nunique())

    avg_units = units / actual_days if actual_days else np.nan
    avg_revenue = revenue / actual_days if actual_days else np.nan
    avg_orders = orders / actual_days if actual_days else np.nan
    return {
        "Units": units,
        "Revenue": revenue,
        "Orders": orders,
        "Avg Daily Units": avg_units,
        "Avg Daily Revenue": avg_revenue,
        "Avg Daily Orders": avg_orders,
        "Run-Rate Units for Window": avg_units * window_days if actual_days else np.nan,
        "Run-Rate Revenue for Window": avg_revenue * window_days if actual_days else np.nan,
        "Run-Rate Orders for Window": avg_orders * window_days if actual_days else np.nan,
    }


def build_outputs(
    campaigns: pd.DataFrame,
    item_index: pd.DataFrame,
    comparison: pd.DataFrame,
    master: pd.DataFrame,
    sales: pd.DataFrame,
    db_today: pd.Timestamp,
    db_min: pd.Timestamp,
    db_max: pd.Timestamp,
) -> dict[str, pd.DataFrame]:
    period_rows: list[dict[str, Any]] = []
    definition_rows: list[dict[str, Any]] = []
    filter_audit_rows: list[dict[str, Any]] = []

    master_item_set = set(master["ITEM_KEY"])

    for _, campaign in campaigns.iterrows():
        idx = campaign["INDEX"]
        campaign_start = pd.Timestamp(campaign["Effective Start"]).normalize()
        campaign_end = pd.Timestamp(campaign["Effective End"]).normalize()
        target = build_target_set(idx, item_index, comparison, master)
        comp, audit = build_comparison_set(idx, comparison, master, target)
        filter_audit_rows.extend(audit)

        explicit_items = target.get("explicit_items", set())
        unmatched_explicit = sorted(explicit_items - master_item_set)
        cohorts = [
            ("Promoted", target["kind"], target["items"], target["source"]),
            ("Comparison", comp["kind"], comp["items"], comp["source"]),
        ]

        if campaign_end > db_max:
            date_note = f"effective/after windows capped at data max {sql_date(db_max)}"
        else:
            date_note = ""

        definition_rows.append(
            {
                "INDEX": idx,
                "Effective Start": campaign_start,
                "Effective End": campaign_end,
                "Effective Days": campaign["Effective Days"],
                "Send Count": campaign["Send Count"],
                "Target Set Type": target["kind"],
                "Target Set Source": target["source"],
                "Target Item Count": np.nan if target["items"] is None else len(target["items"]),
                "Explicit Promoted Item Count": len(explicit_items),
                "Division Promoted Filters": ", ".join(sorted(target.get("division_numbers", set()))),
                "Unmatched Explicit Promoted Items": ", ".join(unmatched_explicit),
                "Ignored Item Index Markers": ", ".join(target.get("invalid_markers", [])),
                "Comparison Set Type": comp["kind"],
                "Comparison Filters": comp.get("filters", ""),
                "Comparison Item Count": np.nan if comp["items"] is None else len(comp["items"]),
                "Date Warning": campaign["Date Warning"],
                "Data Availability Note": date_note,
                "Card Names": campaign["Card Names"],
            }
        )

        for window_days in WINDOW_DAYS:
            phases = ["Before", "After"]
            if not campaign["Single-Day Effective Period"]:
                phases.insert(1, "During")

            for cohort_name, cohort_kind, item_set, source in cohorts:
                if cohort_kind == "EMPTY":
                    continue
                for phase in phases:
                    start, end, actual_days, full_window = period_bounds(
                        phase,
                        window_days,
                        campaign_start,
                        campaign_end,
                        db_min,
                        db_max,
                    )
                    metrics = summarize_period(
                        sales,
                        cohort_kind,
                        item_set,
                        start,
                        end,
                        actual_days,
                        window_days,
                    )
                    period_rows.append(
                        {
                            "INDEX": idx,
                            "Cohort": cohort_name,
                            "Set Type": cohort_kind,
                            "Set Source": source,
                            "Window Days": window_days,
                            "Phase": phase,
                            "Period Start": start,
                            "Period End": end,
                            "Observed Days": actual_days,
                            "Full Requested Window Available": full_window,
                            **metrics,
                        }
                    )

    period_metrics = pd.DataFrame(period_rows)
    set_definitions = pd.DataFrame(definition_rows)
    filter_audit = pd.DataFrame(filter_audit_rows)
    impact = build_impact_long(period_metrics)
    summary = build_index_summary(campaigns, set_definitions, impact, db_today, db_min, db_max)
    highlights = build_30d_revenue_highlights(impact, campaigns, set_definitions)
    assumptions = build_assumptions(db_today, db_min, db_max)

    return {
        "Index Summary": summary,
        "30D Revenue Highlights": highlights,
        "Impact Long": impact,
        "Period Metrics Long": period_metrics,
        "Set Definitions": set_definitions,
        "Filter Match Audit": filter_audit,
        "Campaign Timeline": campaigns,
        "Assumptions": assumptions,
    }


def build_impact_long(period_metrics: pd.DataFrame) -> pd.DataFrame:
    if period_metrics.empty:
        return pd.DataFrame()

    rows: list[dict[str, Any]] = []
    value_cols = {
        "Units": "Avg Daily Units",
        "Revenue": "Avg Daily Revenue",
        "Orders": "Avg Daily Orders",
    }

    grouped = period_metrics.groupby(["INDEX", "Window Days"], sort=False)
    for (idx, window_days), group in grouped:
        for metric, col in value_cols.items():
            def value(cohort: str, phase: str) -> float:
                found = group[(group["Cohort"] == cohort) & (group["Phase"] == phase)]
                if found.empty:
                    return np.nan
                return float(found.iloc[0][col])

            target_before = value("Promoted", "Before")
            target_during = value("Promoted", "During")
            target_after = value("Promoted", "After")
            comp_before = value("Comparison", "Before")
            comp_during = value("Comparison", "During")
            comp_after = value("Comparison", "After")

            target_during_pct = pct_change(target_during, target_before)
            target_after_pct = pct_change(target_after, target_before)
            comp_during_pct = pct_change(comp_during, comp_before)
            comp_after_pct = pct_change(comp_after, comp_before)

            rows.append(
                {
                    "INDEX": idx,
                    "Window Days": window_days,
                    "Metric": metric,
                    "Promoted Before Avg/Day": target_before,
                    "Promoted During Avg/Day": target_during,
                    "Promoted After Avg/Day": target_after,
                    "Comparison Before Avg/Day": comp_before,
                    "Comparison During Avg/Day": comp_during,
                    "Comparison After Avg/Day": comp_after,
                    "Promoted During vs Before %": target_during_pct,
                    "Promoted After vs Before %": target_after_pct,
                    "Comparison During vs Before %": comp_during_pct,
                    "Comparison After vs Before %": comp_after_pct,
                    "Net During Lift vs Comparison (pp)": (
                        target_during_pct - comp_during_pct
                        if pd.notna(target_during_pct) and pd.notna(comp_during_pct)
                        else np.nan
                    ),
                    "Net After Lift vs Comparison (pp)": (
                        target_after_pct - comp_after_pct
                        if pd.notna(target_after_pct) and pd.notna(comp_after_pct)
                        else np.nan
                    ),
                    "Baseline Note": baseline_note(
                        target_before,
                        target_after,
                        comp_before,
                        comp_after,
                        target_during,
                        comp_during,
                    ),
                }
            )
    return pd.DataFrame(rows)


def baseline_note(*values: float) -> str:
    notes = []
    target_before = values[0]
    target_after = values[1]
    comp_before = values[2]
    comp_after = values[3]
    target_during = values[4]
    comp_during = values[5]
    if pd.isna(target_during) and pd.isna(comp_during):
        notes.append("during phase omitted/unavailable")
    elif pd.isna(target_during):
        notes.append("promoted during unavailable")
    elif pd.isna(comp_during):
        notes.append("comparison during unavailable")
    if pd.isna(target_after):
        notes.append("promoted after unavailable")
    if pd.isna(comp_after):
        notes.append("comparison after unavailable")
    if pd.notna(target_before) and abs(float(target_before)) < 1e-12:
        notes.append("promoted before baseline is zero")
    if pd.notna(comp_before) and abs(float(comp_before)) < 1e-12:
        notes.append("comparison before baseline is zero")
    return "; ".join(notes)


def build_30d_revenue_highlights(
    impact: pd.DataFrame,
    campaigns: pd.DataFrame,
    set_definitions: pd.DataFrame,
) -> pd.DataFrame:
    if impact.empty:
        return pd.DataFrame()

    revenue = impact[(impact["Window Days"] == 30) & (impact["Metric"] == "Revenue")].copy()
    if revenue.empty:
        return pd.DataFrame()

    for col in [
        "Promoted During vs Before %",
        "Promoted After vs Before %",
        "Comparison During vs Before %",
        "Comparison After vs Before %",
        "Net During Lift vs Comparison (pp)",
        "Net After Lift vs Comparison (pp)",
    ]:
        revenue[col] = pd.to_numeric(revenue[col], errors="coerce")

    context = campaigns[["INDEX", "Effective Start", "Effective End", "Card Names"]].merge(
        set_definitions[["INDEX", "Target Set Source", "Comparison Filters"]],
        on="INDEX",
        how="left",
    )

    def labeled_rows(label: str, frame: pd.DataFrame, sort_col: str, ascending: bool) -> pd.DataFrame:
        available = frame.dropna(subset=[sort_col]).sort_values(sort_col, ascending=ascending).head(10)
        if available.empty:
            return pd.DataFrame()
        out = available.merge(context, on="INDEX", how="left")
        out.insert(0, "Highlight", label)
        return out[
            [
                "Highlight",
                "INDEX",
                "Effective Start",
                "Effective End",
                "Promoted During vs Before %",
                "Comparison During vs Before %",
                "Net During Lift vs Comparison (pp)",
                "Promoted After vs Before %",
                "Comparison After vs Before %",
                "Net After Lift vs Comparison (pp)",
                "Target Set Source",
                "Comparison Filters",
                "Baseline Note",
                "Card Names",
            ]
        ]

    pieces = [
        labeled_rows(
            "Top Net During Revenue Lift",
            revenue,
            "Net During Lift vs Comparison (pp)",
            False,
        ),
        labeled_rows(
            "Top Promoted After Revenue Lift",
            revenue,
            "Promoted After vs Before %",
            False,
        ),
        labeled_rows(
            "Lowest Promoted After Revenue Lift",
            revenue,
            "Promoted After vs Before %",
            True,
        ),
    ]
    pieces = [p for p in pieces if not p.empty]
    return pd.concat(pieces, ignore_index=True) if pieces else pd.DataFrame()


def build_index_summary(
    campaigns: pd.DataFrame,
    set_definitions: pd.DataFrame,
    impact: pd.DataFrame,
    db_today: pd.Timestamp,
    db_min: pd.Timestamp,
    db_max: pd.Timestamp,
) -> pd.DataFrame:
    summary = campaigns.merge(
        set_definitions[
            [
                "INDEX",
                "Target Set Type",
                "Target Set Source",
                "Target Item Count",
                "Comparison Set Type",
                "Comparison Filters",
                "Comparison Item Count",
                "Unmatched Explicit Promoted Items",
                "Ignored Item Index Markers",
                "Data Availability Note",
            ]
        ],
        on="INDEX",
        how="left",
    )

    impact_30 = impact[(impact["Window Days"] == 30)].copy()
    for metric in ["Units", "Revenue", "Orders"]:
        metric_rows = impact_30[impact_30["Metric"] == metric][
            [
                "INDEX",
                "Promoted During vs Before %",
                "Promoted After vs Before %",
                "Net During Lift vs Comparison (pp)",
                "Net After Lift vs Comparison (pp)",
            ]
        ].rename(
            columns={
                "Promoted During vs Before %": f"{metric} During vs Before 30D %",
                "Promoted After vs Before %": f"{metric} After vs Before 30D %",
                "Net During Lift vs Comparison (pp)": f"{metric} Net During Lift 30D pp",
                "Net After Lift vs Comparison (pp)": f"{metric} Net After Lift 30D pp",
            }
        )
        summary = summary.merge(metric_rows, on="INDEX", how="left")

    summary.insert(1, "Generated At", dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    summary.insert(2, "DB Today", db_today)
    summary.insert(3, "Extracted Sales Data Min", db_min)
    summary.insert(4, "Extracted Sales Data Max", db_max)
    return summary.sort_values("Index Sort").reset_index(drop=True)


def build_assumptions(db_today: pd.Timestamp, db_min: pd.Timestamp, db_max: pd.Timestamp) -> pd.DataFrame:
    rows = [
        ("Database today", sql_date(db_today)),
        ("Extracted sales data range", f"{sql_date(db_min)} through {sql_date(db_max)}"),
        ("Campaign grain", "One output row per INDEX. Multiple Timeline rows are merged to the outer effective date span."),
        ("One-day sends", "If Effective Start = Effective End, the During phase is omitted."),
        ("Units", "Uses SHLINE.SLBLUS because the forecasting reports use it as the customer demand quantity."),
        ("Revenue", "Uses SHLINE.SLENET line net revenue."),
        ("Orders", "Distinct company/location/order/invoice combinations with at least one line in the cohort."),
        ("Sales exclusions", "Excludes item numbers beginning PROMO and customer numbers containing TRANSFER, OMP000, INV000, or OLD001."),
        ("Comparison exclusions", "Comparison item sets subtract promoted items expanded from Item Index before metrics are calculated."),
        ("TOTAL REVENUE", "Used for non-product campaigns with no Item Index rows; otherwise appears as a comparison benchmark."),
        ("Family alias", "Comparison List family SO is mapped to database family code SF, SOLID FLOORING."),
        ("Collection aliases", "Garrison II Smooth/Distressed are mapped to GII SMOOTH/GII DISTRESSED."),
        ("Partial windows", "Before/after windows are capped to available sales history. Check Full Requested Window Available in Period Metrics Long."),
    ]
    return pd.DataFrame(rows, columns=["Topic", "Assumption / Note"])


def write_output(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        for sheet_name, df in sheets.items():
            safe = df.copy()
            for col in safe.columns:
                if safe[col].dtype == "object":
                    safe[col] = safe[col].map(
                        lambda v: ", ".join(sorted(v)) if isinstance(v, set) else v
                    )
            safe.to_excel(writer, sheet_name=sheet_name[:31], index=False)

        workbook = writer.book
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column_cells in sheet.columns:
                letter = get_column_letter(column_cells[0].column)
                max_len = 0
                for cell in column_cells[:200]:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                sheet.column_dimensions[letter].width = min(max(max_len + 2, 10), 60)

            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    header = sheet.cell(row=1, column=cell.column).value
                    if header and ("%" in str(header) or str(header).endswith("pp")):
                        cell.number_format = "0.0"
                    elif header and any(word in str(header) for word in ["Revenue"]):
                        cell.number_format = '$#,##0.00'
                    elif header and (
                        "Date" in str(header)
                        or header
                        in {
                            "Effective Start",
                            "Effective End",
                            "Period Start",
                            "Period End",
                            "DB Today",
                            "Extracted Sales Data Min",
                            "Extracted Sales Data Max",
                        }
                    ):
                        cell.number_format = "yyyy-mm-dd"


def main() -> None:
    parser = argparse.ArgumentParser(description="Analyze marketing blast sales impact.")
    parser.add_argument("--archive", type=Path, default=ARCHIVE_PATH)
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH)
    args = parser.parse_args()

    warnings.filterwarnings(
        "ignore",
        message="pandas only supports SQLAlchemy connectable",
        category=UserWarning,
    )

    print(f"Reading archive: {args.archive}")
    timeline, item_index, comparison = read_archive(args.archive)
    campaigns = build_campaigns(timeline)
    earliest_start = campaigns["Effective Start"].min() - pd.Timedelta(days=max(WINDOW_DAYS))
    latest_end = campaigns["Effective End"].max() + pd.Timedelta(days=max(WINDOW_DAYS))

    print("Connecting to Gartman...")
    conn = get_connection()
    try:
        db_today, db_min, db_max = get_db_date_range(conn)
        query_start = max(pd.Timestamp(earliest_start).normalize(), db_min)
        query_end = min(pd.Timestamp(latest_end).normalize(), db_max)
        print(f"Fetching item master...")
        master = fetch_master(conn)
        print(f"  item master rows: {len(master):,}")
        print(f"Fetching sales lines {sql_date(query_start)} through {sql_date(query_end)}...")
        sales = fetch_sales_lines(conn, query_start, query_end)
        print(f"  sales line rows: {len(sales):,}")
    finally:
        conn.close()

    print("Calculating impact metrics...")
    outputs = build_outputs(campaigns, item_index, comparison, master, sales, db_today, query_start, query_end)
    print(f"Writing workbook: {args.output}")
    write_output(args.output, outputs)
    print("Done.")


if __name__ == "__main__":
    main()
